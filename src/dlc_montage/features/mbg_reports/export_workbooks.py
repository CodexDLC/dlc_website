"""Рендеринг Excel-книг MBG в фирменном стиле DLC Montage."""

import re
from collections import defaultdict
from datetime import date, time, timedelta
from decimal import Decimal
from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

# ==============================================================================
# Корпоративный стиль DLC Montage (Brand Design & Palette)
# ==============================================================================
COLOR_NAVY_900 = "0A2540"  # Deep Navy: шапки таблиц и главные заголовки
COLOR_NAVY_700 = "1A3B5C"  # Navy Secondary: подзаголовки и метки
COLOR_NAVY_100 = "E6EAF0"  # Тонкие разделители и границы ячеек
COLOR_GREEN_600 = "00A86B"  # Clean Energy Green: акцент итоговых сумм
COLOR_GREEN_50 = "E6FBF2"  # Мягкий мятный фон для строк итогов (Gesamt / Summe)
COLOR_GRAY_900 = "111827"  # Основной цвет текста (Dark Slate)
COLOR_GRAY_600 = "4B5563"  # Второстепенный текст и контакты
COLOR_GRAY_100 = "F8FAFC"  # Мягкий фон для чередующихся строк (зебра)
COLOR_WHITE = "FFFFFF"

FONT_FAMILY = "Segoe UI"
GERMAN_WEEKDAYS = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]

# Данные компании для шапки отчётов
COMPANY_NAME = "DLC Montage"
COMPANY_CONTACT = "Dmitri Labodin"
COMPANY_PHONE = "Tel: +49 176 62203619"
COMPANY_EMAIL = "dmitri.labodin@googlemail.com"


def text_cell(cell, value):
    """Пользовательский текст никогда не интерпретируется как формула Excel."""
    cell.value = str(value)
    cell.data_type = "s"


def _apply_company_header_block(sheet, start_col=5, end_col=7, max_rows=3):
    """Разместить контактный блок компании справа вверху."""
    sheet.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    text_cell(sheet.cell(1, start_col), COMPANY_NAME)
    sheet.cell(1, start_col).font = Font(name=FONT_FAMILY, size=13, bold=True, color=COLOR_NAVY_900)
    sheet.cell(1, start_col).alignment = Alignment(horizontal="right", vertical="center")

    if max_rows == 3:
        sheet.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
        text_cell(sheet.cell(2, start_col), COMPANY_CONTACT)
        sheet.cell(2, start_col).font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_600)
        sheet.cell(2, start_col).alignment = Alignment(horizontal="right", vertical="center")

        sheet.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
        text_cell(sheet.cell(3, start_col), f"{COMPANY_PHONE}  ·  {COMPANY_EMAIL}")
        sheet.cell(3, start_col).font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_600)
        sheet.cell(3, start_col).alignment = Alignment(horizontal="right", vertical="center")
    else:
        company_lines = [
            (2, COMPANY_CONTACT),
            (3, COMPANY_PHONE),
            (4, COMPANY_EMAIL),
        ]
        for row_idx, text in company_lines:
            sheet.merge_cells(start_row=row_idx, start_column=start_col, end_row=row_idx, end_column=end_col)
            text_cell(sheet.cell(row_idx, start_col), text)
            sheet.cell(row_idx, start_col).font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_600)
            sheet.cell(row_idx, start_col).alignment = Alignment(horizontal="right", vertical="center")


def _site_sheet_title(site_name, site_id=None):
    """Сформировать валидное имя листа Excel для объекта."""
    title = re.sub(r"[\\/*?:\[\]]", "-", str(site_name)).strip()[:31] or "Standort"
    return title


def build_workbook(snapshots, year, week):
    """Собрать общую недельную книгу AC: сводка и лист каждого Standort в фирменном стиле."""
    grouped = defaultdict(list)
    for snapshot in snapshots:
        grouped[snapshot["site_id"]].append(snapshot)
    try:
        monday = date.fromisocalendar(year, week, 1)
        sunday = monday + timedelta(days=6)
        period_str = f"{monday.strftime('%d.%m.%Y')}-{sunday.strftime('%d.%m.%Y')}"
    except Exception:
        period_str = ""
    return build_weekly_workbook(
        {
            "year": year,
            "week": week,
            "week_period": period_str,
            "grouped": grouped,
        }
    )


def build_weekly_workbook(snapshot):
    """Собрать почасовой Stundennachweis: сводка и лист каждого Standort."""
    book = Workbook()
    grouped = snapshot["grouped"]
    document_title = snapshot.get("document_title", "Stundennachweis")
    period_label = snapshot.get("period_label", "Kalenderwoche")
    period_value = snapshot.get("period_value")
    if period_value is None:
        period_value = f"{snapshot['week']:02} / {snapshot['year']} ({snapshot.get('week_period', '')})"
    total_label = snapshot.get("total_label", "Summe Woche")

    # ---------------------------------------------------------
    # Лист 1: Сводная таблица (Übersicht)
    # ---------------------------------------------------------
    summary = book.active
    summary.title = "Übersicht"
    summary.sheet_view.showGridLines = True

    # Заголовок слева
    summary.merge_cells("A1:D1")
    text_cell(summary["A1"], document_title)
    summary["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)

    summary.cell(2, 1, f"{period_label}:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
    summary.merge_cells("B2:D2")
    text_cell(summary["B2"], period_value)
    summary["B2"].font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

    # Контакты компании справа (строки 1-3)
    _apply_company_header_block(summary, start_col=5, end_col=7, max_rows=3)

    summary_headers = ["Anlage", "Standort", "P/Nr", "Adresse", "Arbeitstage", "Mitarbeiter", "Stunden"]
    for column, label in enumerate(summary_headers, 1):
        cell = summary.cell(4, column, label)
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[4].height = 28

    sorted_sites = sorted(grouped.items())
    for line, (site_id, days) in enumerate(sorted_sites, 5):
        first = days[0]
        employee_names = sorted(
            {
                p["name"]
                for day in days
                for entry in day.get("entries", [])
                for p in entry.get("participants", [])
                if p.get("name")
            }
        )
        site_sheet_name = _site_sheet_title(first.get("site_name", "Standort"), site_id)
        site_entry_count = sum(len(day.get("entries", [])) for day in days)
        site_total_line = 7 + site_entry_count

        for column, value in (
            (1, first.get("group_name", "—")),
            (2, first.get("site_name", "—")),
            (3, first.get("project_number", "—")),
            (4, first.get("address", "—")),
            (6, ", ".join(employee_names)),
        ):
            text_cell(summary.cell(line, column), value)
            summary.cell(line, column).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        summary.cell(line, 5, len({day["date"] for day in days})).font = Font(
            name=FONT_FAMILY, size=10, color=COLOR_GRAY_900
        )
        summary.cell(line, 5).alignment = Alignment(horizontal="center", vertical="top")

        # Формула ссылки на итог листа Standort
        summary.cell(line, 7, f"='{site_sheet_name}'!F{site_total_line}").number_format = "0.00"
        summary.cell(line, 7).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
        summary.cell(line, 7).alignment = Alignment(horizontal="right", vertical="top")

        row_fill = PatternFill("solid", fgColor=COLOR_WHITE if line % 2 != 0 else COLOR_GRAY_100)
        for cell in summary[line]:
            if cell.column not in (5, 7):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
                top=Side(style="thin", color=COLOR_NAVY_100),
                bottom=Side(style="thin", color=COLOR_NAVY_100),
            )
            cell.fill = row_fill

        summary.row_dimensions[line].height = max(28, 16 * max(1, len(employee_names)))

    total_line = 5 + len(grouped)
    text_cell(summary.cell(total_line, 6), "Gesamt")
    summary.cell(total_line, 6).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
    summary.cell(total_line, 6).alignment = Alignment(vertical="center", horizontal="right")

    if len(grouped) > 0:
        summary.cell(total_line, 7, f"=SUM(G5:G{total_line - 1})").number_format = "0.00"
    else:
        summary.cell(total_line, 7, 0.0).number_format = "0.00"
    summary.cell(total_line, 7).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
    summary.cell(total_line, 7).alignment = Alignment(vertical="center", horizontal="right")

    total_fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
    for cell in summary[total_line]:
        cell.fill = total_fill
        cell.border = Border(
            top=Side(style="thin", color=COLOR_GREEN_600),
            bottom=Side(style="double", color=COLOR_NAVY_900),
            left=Side(style="thin", color=COLOR_NAVY_100),
            right=Side(style="thin", color=COLOR_NAVY_100),
        )
    summary.row_dimensions[total_line].height = 28

    for column, width in {"A": 28, "B": 26, "C": 18, "D": 34, "E": 14, "F": 38, "G": 18}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A5"
    summary.print_title_rows = "1:4"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    summary.print_area = f"A1:G{total_line}"
    summary.oddFooter.center.text = "Seite &P / &N"

    # ---------------------------------------------------------
    # Листы объектов (Standort)
    # ---------------------------------------------------------
    for site_id, days in sorted_sites:
        first = days[0]
        sheet_name = _site_sheet_title(first["site_name"], site_id)
        sheet = book.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = True

        # Заголовок и метаданные слева
        sheet.merge_cells("A1:D1")
        text_cell(sheet["A1"], document_title)
        sheet["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)

        meta_info = [
            (2, "Projekt", first["address"]),
            (3, "P/Nr", first["project_number"]),
            (4, period_label, period_value),
        ]
        for row, label, value in meta_info:
            sheet.cell(row, 1, label).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            text_cell(sheet.cell(row, 2), value)
            sheet.cell(row, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        # Контакты компании справа
        _apply_company_header_block(sheet, start_col=5, end_col=7)

        headers = ["Datum", "Name", "Beginn", "Ende", "Weg", "Mont (h)", "Beschreibung der Tätigkeit"]
        for col, label in enumerate(headers, 1):
            cell = sheet.cell(6, col, label)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
            cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
            cell.alignment = Alignment(
                wrap_text=True, vertical="center", horizontal="center" if col in (1, 3, 4, 5, 6) else "left"
            )
        sheet.row_dimensions[6].height = 28

        line = 7
        sorted_days = sorted(days, key=lambda d: d["date"])
        for day_idx, day in enumerate(sorted_days):
            day_date = date.fromisoformat(day["date"])
            weekday_label = f"{GERMAN_WEEKDAYS[day_date.weekday()]}, {day_date.strftime('%d.%m.%Y')}"
            day_start_line = line
            day_fill = PatternFill("solid", fgColor=COLOR_WHITE if day_idx % 2 == 0 else COLOR_GRAY_100)

            for entry in day.get("entries", []):
                description = entry.get("description", "")
                participants = entry.get("participants", [])
                p_count = len(participants) if participants else 1

                if entry.get("actual_person_hours") is not None:
                    person_hours = Decimal(str(entry["actual_person_hours"]))
                elif "minutes" in entry and entry["minutes"] is not None:
                    person_hours = Decimal(entry["minutes"]) * p_count / Decimal(60)
                elif entry.get("start") and entry.get("end"):
                    st_val = entry["start"]
                    en_val = entry["end"]
                    st_time = st_val if isinstance(st_val, time) else time.fromisoformat(str(st_val))
                    en_time = en_val if isinstance(en_val, time) else time.fromisoformat(str(en_val))
                    mins = (en_time.hour * 60 + en_time.minute) - (st_time.hour * 60 + st_time.minute)
                    person_hours = Decimal(mins) * p_count / Decimal(60)
                else:
                    person_hours = Decimal("0.00")

                names = "\n".join(participant.get("name", "") for participant in participants)

                if line == day_start_line:
                    text_cell(sheet.cell(line, 1), weekday_label)
                    sheet.cell(line, 1).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)
                    sheet.cell(line, 1).alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

                text_cell(sheet.cell(line, 2), names)
                sheet.cell(line, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 2).alignment = Alignment(vertical="top", wrap_text=True)

                st_val = entry.get("start")
                if st_val:
                    st_time = st_val if isinstance(st_val, time) else time.fromisoformat(str(st_val))
                    sheet.cell(line, 3, st_time).number_format = "hh:mm"
                sheet.cell(line, 3).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 3).alignment = Alignment(vertical="top", horizontal="center")

                en_val = entry.get("end")
                if en_val:
                    en_time = en_val if isinstance(en_val, time) else time.fromisoformat(str(en_val))
                    sheet.cell(line, 4, en_time).number_format = "hh:mm"
                sheet.cell(line, 4).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 4).alignment = Alignment(vertical="top", horizontal="center")

                # Col 5: Weg
                sheet.cell(line, 5).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 5).alignment = Alignment(vertical="top", horizontal="center")

                sheet.cell(line, 6, float(person_hours)).number_format = "0.00"
                sheet.cell(line, 6).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 6).alignment = Alignment(vertical="top", horizontal="right")

                text_cell(sheet.cell(line, 7), description)
                sheet.cell(line, 7).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                sheet.cell(line, 7).alignment = Alignment(vertical="top", wrap_text=True)

                for col_idx in range(1, 8):
                    cell = sheet.cell(line, col_idx)
                    cell.fill = day_fill
                    cell.border = Border(
                        left=Side(style="thin", color=COLOR_NAVY_100),
                        right=Side(style="thin", color=COLOR_NAVY_100),
                        top=Side(style="thin", color=COLOR_NAVY_100),
                        bottom=Side(style="thin", color=COLOR_NAVY_100),
                    )

                row_lines = max(len(entry["participants"]), len(description) // 48 + 1)
                sheet.row_dimensions[line].height = max(26, 16 * row_lines)
                line += 1

            day_end_line = line - 1
            if day_end_line > day_start_line:
                sheet.merge_cells(start_row=day_start_line, start_column=1, end_row=day_end_line, end_column=1)
                sheet.cell(day_start_line, 1).alignment = Alignment(
                    vertical="center", horizontal="center", wrap_text=True
                )

        text_cell(sheet.cell(line, 2), total_label)
        sheet.cell(line, 2).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
        sheet.cell(line, 2).alignment = Alignment(vertical="center", horizontal="left")

        # Формула суммы часов за неделю по объекту
        if line > 7:
            sheet.cell(line, 6, f"=SUM(F7:F{line - 1})").number_format = "0.00"
        else:
            sheet.cell(line, 6, 0.0).number_format = "0.00"
        sheet.cell(line, 6).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
        sheet.cell(line, 6).alignment = Alignment(vertical="center", horizontal="right")

        total_fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
        for col_idx in range(1, 8):
            c_cell = sheet.cell(line, col_idx)
            c_cell.fill = total_fill
            c_cell.border = Border(
                top=Side(style="thin", color=COLOR_GREEN_600),
                bottom=Side(style="double", color=COLOR_NAVY_900),
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
            )
        sheet.row_dimensions[line].height = 28

        for col, width in {"A": 16, "B": 28, "C": 11, "D": 11, "E": 9, "F": 14, "G": 55}.items():
            sheet.column_dimensions[col].width = width
        sheet.freeze_panes = "A7"
        sheet.print_title_rows = "1:6"
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
        sheet.print_area = f"A1:G{line}"
        sheet.oddFooter.center.text = "Seite &P / &N"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def build_final_hourly_workbook(snapshots, period_start, period_end):
    """Собрать итоговый почасовой отчёт Anlage без ограничения календарной неделей."""
    grouped = defaultdict(list)
    for snapshot in snapshots:
        grouped[snapshot["site_id"]].append(snapshot)
    period_value = f"{period_start:%d.%m.%Y}-{period_end:%d.%m.%Y}"
    return build_weekly_workbook(
        {
            "grouped": grouped,
            "document_title": "Gesamtstundennachweis",
            "period_label": "Zeitraum",
            "period_value": period_value,
            "total_label": "Summe Zeitraum",
        }
    )


def build_performance_workbook(snapshot):
    """Собрать Leistungsnachweis или Zwischennachweis Anlage с динамическими формулами."""
    book = Workbook()
    is_preliminary = snapshot.get("is_preliminary", False)
    incomplete = any(site.get("amount") is None for site in snapshot["sites"])
    total_label = "Teilbetrag" if incomplete else "Gesamt"
    doc_title = "Zwischennachweis" if is_preliminary else "Leistungsnachweis"

    # ---------------------------------------------------------
    # Лист 1: Сводная таблица (Übersicht)
    # ---------------------------------------------------------
    summary = book.active
    summary.title = "Übersicht"
    summary.sheet_view.showGridLines = True

    # Заголовок слева
    summary.merge_cells("A1:C1")
    text_cell(summary["A1"], doc_title)
    summary["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)

    summary.cell(2, 1, "Anlage:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
    summary.merge_cells("B2:C2")
    text_cell(summary["B2"], snapshot["group_name"])
    summary["B2"].font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

    if is_preliminary:
        summary.cell(3, 1, "Stand:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
        summary.merge_cells("B3:C3")
        text_cell(summary["B3"], f"{snapshot.get('export_date', '')} · Vorläufige Leistungsübersicht")
        summary["B3"].font = Font(name=FONT_FAMILY, size=9.5, italic=True, color=COLOR_NAVY_700)

    # Контакты компании справа (строки 1-3)
    company_start_col = 5 if is_preliminary else 4
    company_end_col = 7 if is_preliminary else 6
    _apply_company_header_block(summary, start_col=company_start_col, end_col=company_end_col, max_rows=3)

    header_row = 5 if is_preliminary else 4
    if is_preliminary:
        summary_headers = ["Anlage", "Standort", "P/Nr", "Adresse", "Status", "Positionen", total_label]
    else:
        summary_headers = ["Anlage", "Standort", "P/Nr", "Adresse", "Positionen", total_label]

    for column, label in enumerate(summary_headers, 1):
        cell = summary.cell(header_row, column, label)
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[header_row].height = 28

    start_data_row = header_row + 1
    for line, site_snapshot in enumerate(snapshot["sites"], start_data_row):
        site_sheet_name = _site_sheet_title(site_snapshot["site_name"], site_snapshot["site_id"])
        site_total_line = 9 + len(site_snapshot["rows"])

        text_cell(summary.cell(line, 1), site_snapshot.get("group_name") or snapshot.get("group_name", "—"))
        summary.cell(line, 1).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        text_cell(summary.cell(line, 2), site_snapshot["site_name"])
        summary.cell(line, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        text_cell(summary.cell(line, 3), site_snapshot["project_number"])
        summary.cell(line, 3).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        text_cell(summary.cell(line, 4), site_snapshot["address"])
        summary.cell(line, 4).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        if is_preliminary:
            is_done = site_snapshot.get("status") == "completed"
            status_text = "Abgeschlossen" if is_done else "In Bearbeitung"
            text_cell(summary.cell(line, 5), status_text)
            summary.cell(line, 5).font = Font(
                name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_GREEN_600 if is_done else "D97706"
            )
            summary.cell(line, 5).alignment = Alignment(horizontal="center", vertical="top")

            pos_col, tot_col = 6, 7
        else:
            pos_col, tot_col = 5, 6

        # Формула подсчёта количества ненулевых позиций на листе объекта
        summary.cell(line, pos_col, f"=COUNTIF('{site_sheet_name}'!C9:C{site_total_line - 1}, \">0\")").font = Font(
            name=FONT_FAMILY, size=10, color=COLOR_GRAY_900
        )
        summary.cell(line, pos_col).alignment = Alignment(horizontal="center", vertical="top")

        # Формула ссылки на итог листа объекта
        summary.cell(line, tot_col, f"='{site_sheet_name}'!E{site_total_line}").number_format = '#,##0.00 "€"'
        summary.cell(line, tot_col).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
        summary.cell(line, tot_col).alignment = Alignment(horizontal="right", vertical="top")

        row_fill = PatternFill("solid", fgColor=COLOR_WHITE if line % 2 != 0 else COLOR_GRAY_100)
        for cell in summary[line]:
            if cell.column not in (pos_col, tot_col):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
                top=Side(style="thin", color=COLOR_NAVY_100),
                bottom=Side(style="thin", color=COLOR_NAVY_100),
            )
            cell.fill = row_fill

    summary_total_line = start_data_row + len(snapshot["sites"])
    label_col = 6 if is_preliminary else 5
    tot_col = 7 if is_preliminary else 6

    text_cell(summary.cell(summary_total_line, label_col), total_label)
    summary.cell(summary_total_line, label_col).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
    summary.cell(summary_total_line, label_col).alignment = Alignment(vertical="center", horizontal="right")

    if len(snapshot["sites"]) > 0:
        tot_letter = "G" if is_preliminary else "F"
        summary.cell(
            summary_total_line, tot_col, f"=SUM({tot_letter}{start_data_row}:{tot_letter}{summary_total_line - 1})"
        ).number_format = '#,##0.00 "€"'
    else:
        summary.cell(summary_total_line, tot_col, 0.0).number_format = '#,##0.00 "€"'

    summary.cell(summary_total_line, tot_col).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
    summary.cell(summary_total_line, tot_col).alignment = Alignment(vertical="center", horizontal="right")

    total_fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
    for cell in summary[summary_total_line]:
        cell.fill = total_fill
        cell.border = Border(
            top=Side(style="thin", color=COLOR_GREEN_600),
            bottom=Side(style="double", color=COLOR_NAVY_900),
            left=Side(style="thin", color=COLOR_NAVY_100),
            right=Side(style="thin", color=COLOR_NAVY_100),
        )
    summary.row_dimensions[summary_total_line].height = 28

    if is_preliminary:
        col_widths = {"A": 26, "B": 24, "C": 16, "D": 32, "E": 16, "F": 14, "G": 18}
        last_col_letter = "G"
    else:
        col_widths = {"A": 26, "B": 24, "C": 16, "D": 34, "E": 14, "F": 18}
        last_col_letter = "F"

    for col_letter, width in col_widths.items():
        summary.column_dimensions[col_letter].width = width
    summary.freeze_panes = f"A{start_data_row}"
    summary.print_title_rows = f"1:{header_row}"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    if incomplete:
        warning_row = summary_total_line + 2
        summary.merge_cells(start_row=warning_row, start_column=1, end_row=warning_row, end_column=tot_col)
        text_cell(
            summary.cell(warning_row, 1),
            "Hinweis: Preise oder Mengen nicht vereinbart; Teilbetrag ist kein Endbetrag.",
        )
        summary.cell(warning_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.row_dimensions[warning_row].height = 30
    summary.print_area = f"A1:{last_col_letter}{summary_total_line + 2 if incomplete else summary_total_line}"
    summary.oddFooter.center.text = "Seite &P / &N"

    # ---------------------------------------------------------
    # Листы объектов (Standort)
    # ---------------------------------------------------------
    for site_snapshot in snapshot["sites"]:
        sheet_name = _site_sheet_title(site_snapshot["site_name"], site_snapshot["site_id"])
        sheet = book.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = True

        # Заголовок и метаданные слева
        sheet.merge_cells("A1:C1")
        text_cell(sheet["A1"], doc_title)
        sheet["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)

        is_site_done = site_snapshot.get("status") == "completed"
        status_label = "Abgeschlossen" if is_site_done else "In Bearbeitung (noch nicht abgeschlossen)"

        meta_rows = [
            (2, "Anlage", site_snapshot.get("group_name") or snapshot.get("group_name", "—")),
            (3, "Standort", site_snapshot["site_name"]),
            (4, "P/Nr", site_snapshot["project_number"]),
            (5, "Adresse", site_snapshot["address"]),
            (6, "Status", status_label),
        ]
        for row, label, value in meta_rows:
            sheet.cell(row, 1, label).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            text_cell(sheet.cell(row, 2), value)
            if label == "Status":
                sheet.cell(row, 2).font = Font(
                    name=FONT_FAMILY, size=10, bold=True, color=COLOR_GREEN_600 if is_site_done else "D97706"
                )
            else:
                sheet.cell(row, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

        # Контакты компании справа
        _apply_company_header_block(sheet, start_col=4, end_col=5, max_rows=4)

        headers = ["Leistung", "Einheit", "Menge", "Einzelpreis", "Gesamt"]
        for column, label in enumerate(headers, 1):
            cell = sheet.cell(8, column, label)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
            cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[8].height = 28

        for line, row in enumerate(site_snapshot["rows"], 9):
            text_cell(sheet.cell(line, 1), row["name"])
            sheet.cell(line, 1).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)

            text_cell(sheet.cell(line, 2), row["unit"])
            sheet.cell(line, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            sheet.cell(line, 2).alignment = Alignment(horizontal="center", vertical="top")

            # Col 3: Menge
            if row["quantity"] is None:
                text_cell(sheet.cell(line, 3), "Menge ungeklärt")
            else:
                sheet.cell(line, 3, float(Decimal(row["quantity"]))).number_format = "0.00"
            sheet.cell(line, 3).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            sheet.cell(line, 3).alignment = Alignment(horizontal="right", vertical="top")

            # Col 4: Einzelpreis
            if row["unit_price"] is None:
                sheet.cell(line, 4, 0.0).number_format = '#,##0.00 "€"'
            else:
                sheet.cell(line, 4, float(Decimal(row["unit_price"]))).number_format = '#,##0.00 "€"'
            sheet.cell(line, 4).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            sheet.cell(line, 4).alignment = Alignment(horizontal="right", vertical="top")

            # Col 5: Gesamt = Menge * Einzelpreis (динамическая формула)
            sheet.cell(line, 5, f"=IF(ISNUMBER(C{line}), ROUND(C{line}*D{line}, 2), 0)").number_format = '#,##0.00 "€"'
            sheet.cell(line, 5).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            sheet.cell(line, 5).alignment = Alignment(horizontal="right", vertical="top")

            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if line % 2 != 0 else COLOR_GRAY_100)
            for cell in sheet[line]:
                if cell.column == 1:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(
                    left=Side(style="thin", color=COLOR_NAVY_100),
                    right=Side(style="thin", color=COLOR_NAVY_100),
                    top=Side(style="thin", color=COLOR_NAVY_100),
                    bottom=Side(style="thin", color=COLOR_NAVY_100),
                )
                cell.fill = row_fill

        total_line = 9 + len(site_snapshot["rows"])
        site_incomplete = site_snapshot.get("amount") is None
        text_cell(sheet.cell(total_line, 2), "Teilbetrag" if site_incomplete else "Gesamt")
        sheet.cell(total_line, 2).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
        sheet.cell(total_line, 2).alignment = Alignment(vertical="center", horizontal="center")

        # Итог по колонке Gesamt
        if len(site_snapshot["rows"]) > 0:
            sheet.cell(total_line, 5, f"=SUM(E9:E{total_line - 1})").number_format = '#,##0.00 "€"'
        else:
            sheet.cell(total_line, 5, 0.0).number_format = '#,##0.00 "€"'

        sheet.cell(total_line, 5).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
        sheet.cell(total_line, 5).alignment = Alignment(vertical="center", horizontal="right")

        total_fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
        for cell in sheet[total_line]:
            cell.fill = total_fill
            cell.border = Border(
                top=Side(style="thin", color=COLOR_GREEN_600),
                bottom=Side(style="double", color=COLOR_NAVY_900),
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
            )
        sheet.row_dimensions[total_line].height = 28

        for column, width in {"A": 50, "B": 14, "C": 14, "D": 16, "E": 18}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A9"
        sheet.print_title_rows = "1:8"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
        if site_incomplete:
            warning_row = total_line + 2
            sheet.merge_cells(start_row=warning_row, start_column=1, end_row=warning_row, end_column=5)
            text_cell(
                sheet.cell(warning_row, 1),
                "Hinweis: Mengen oder Preise nicht vereinbart; Teilbetrag ist kein Endbetrag.",
            )
            sheet.cell(warning_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            sheet.row_dimensions[warning_row].height = 30
        sheet.print_area = f"A1:E{total_line + 2 if site_incomplete else total_line}"
        sheet.oddFooter.center.text = "Seite &P / &N"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _completion_text_height(value, width, minimum=20):
    """Оценить высоту переносимого текста, сохранив заданные пользователем переводы строк."""
    lines = sum(max(1, len(wrap(line, width=width))) for line in str(value).splitlines())
    return max(minimum, 15 * lines + 6)


def build_work_completion_workbook(snapshot):
    """Собрать клиентский Leistungsnachweis без цен, норм, формул и подписей."""
    book = Workbook()
    summary = book.active
    summary.title = "Übersicht"
    summary.sheet_view.showGridLines = False
    summary.row_dimensions[1].height = 30
    summary.merge_cells("A1:D1")
    text_cell(summary["A1"], "Leistungsnachweis")
    summary["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)
    group_name = snapshot.get("group_name") or "Ein Standort"
    metadata = [(2, "Anlage", group_name), (3, "Leistungszeitraum", snapshot.get("period_label", "—"))]
    for row, label, value in metadata:
        summary.cell(row, 1, label).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
        summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        text_cell(summary.cell(row, 2), value)
        summary.cell(row, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
        summary.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        summary.row_dimensions[row].height = _completion_text_height(value, 65)
    _apply_company_header_block(summary, start_col=5, end_col=7, max_rows=3)

    headers = ["Anlage", "Standort", "P/Nr", "Adresse", "Status", "Positionen", "Ist-Personenstunden"]
    for column, label in enumerate(headers, 1):
        cell = summary.cell(5, column, label)
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    summary.row_dimensions[5].height = 28

    start_data_row = 6
    group_hours = Decimal("0")
    for line, site in enumerate(snapshot["sites"], start_data_row):
        site_hours = sum((Decimal(row["actual_person_hours"]) for row in site["rows"]), Decimal("0"))
        group_hours += site_hours
        values = [
            site["group_name"],
            site["site_name"],
            site["project_number"],
            site["address"],
            "Abgeschlossen" if site["status"] == "completed" else "In Bearbeitung",
            len(site["rows"]),
            float(site_hours),
        ]
        for column, value in enumerate(values, 1):
            cell = summary.cell(line, column)
            if column <= 5:
                text_cell(cell, value)
            else:
                cell.value = value
                if column == 7:
                    cell.number_format = "0.00"
            cell.font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            cell.alignment = Alignment(
                horizontal="right" if column in (6, 7) else "left", vertical="top", wrap_text=True
            )
            cell.border = Border(
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
                top=Side(style="thin", color=COLOR_NAVY_100),
                bottom=Side(style="thin", color=COLOR_NAVY_100),
            )
            cell.fill = PatternFill("solid", fgColor=COLOR_WHITE if line % 2 else COLOR_GRAY_100)
        summary.row_dimensions[line].height = max(
            _completion_text_height(value, width, minimum=32)
            for value, width in zip(values[:5], (24, 22, 14, 30, 16), strict=True)
        )

    total_row = start_data_row + len(snapshot["sites"])
    text_cell(summary.cell(total_row, 6), "Gesamt")
    summary.cell(total_row, 6).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
    summary.cell(total_row, 7, float(group_hours)).font = Font(
        name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600
    )
    summary.cell(total_row, 7).number_format = "0.00"
    for cell in summary[total_row]:
        cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
        cell.border = Border(
            top=Side(style="thin", color=COLOR_GREEN_600), bottom=Side(style="double", color=COLOR_NAVY_900)
        )

    excluded = snapshot.get("excluded_reports", [])
    unknown_count = sum(site.get("unresolved_quantity_count", 0) for site in snapshot["sites"])
    warnings = []
    if excluded:
        warnings.append(
            f"Hinweis: {len(excluded)} nicht aktuelle oder noch nicht angenommene "
            "MBG-Berichte wurden nicht berücksichtigt."
        )
    if unknown_count:
        warnings.append(f"Hinweis: Bei {unknown_count} Positionen ist die Menge ungeklärt; Zeiten bleiben ausgewiesen.")
    for warning_row, warning in enumerate(warnings, total_row + 2):
        summary.merge_cells(start_row=warning_row, start_column=1, end_row=warning_row, end_column=7)
        text_cell(summary.cell(warning_row, 1), warning)
        summary.cell(warning_row, 1).font = Font(name=FONT_FAMILY, size=9.5, italic=True, color="B45309")
        summary.cell(warning_row, 1).fill = PatternFill("solid", fgColor="FEF3C7")
        summary.cell(warning_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        summary.row_dimensions[warning_row].height = 32

    for column, width in {"A": 26, "B": 24, "C": 16, "D": 32, "E": 18, "F": 14, "G": 22}.items():
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A6"
    summary.print_title_rows = "1:5"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.paperSize = summary.PAPERSIZE_A4
    summary.page_setup.fitToWidth = 1
    summary.page_setup.fitToHeight = 0
    summary.sheet_properties.pageSetUpPr.fitToPage = True
    summary.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
    summary.oddFooter.center.text = "Seite &P / &N"
    summary.print_area = f"A1:G{total_row + 1 + len(warnings) if warnings else total_row}"

    for site in snapshot["sites"]:
        sheet = book.create_sheet(_site_sheet_title(site["site_name"], site["site_id"]))
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        sheet.merge_cells("A1:D1")
        text_cell(sheet["A1"], "Leistungsnachweis")
        sheet["A1"].font = Font(name=FONT_FAMILY, size=18, bold=True, color=COLOR_NAVY_900)
        site_metadata = [
            (2, "Anlage", site["group_name"]),
            (3, "Standort", site["site_name"]),
            (4, "P/Nr", site["project_number"]),
            (5, "Adresse", site["address"]),
            (
                6,
                "Leistungszeitraum",
                f"{site['period_start'][8:10]}.{site['period_start'][5:7]}.{site['period_start'][:4]} - "
                f"{site['period_end'][8:10]}.{site['period_end'][5:7]}.{site['period_end'][:4]}",
            ),
        ]
        for row, label, value in site_metadata:
            sheet.cell(row, 1, label).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
            text_cell(sheet.cell(row, 2), value)
            sheet.cell(row, 2).font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
            sheet.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[row].height = _completion_text_height(value, 48)
        _apply_company_header_block(sheet, start_col=5, end_col=5, max_rows=4)
        headers = ["Leistung", "Einheit", "Menge", "Ist-Personenstunden", "Nachweis"]
        for column, label in enumerate(headers, 1):
            cell = sheet.cell(8, column, label)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
            cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[8].height = 28
        site_hours = Decimal("0")
        for line, row in enumerate(site["rows"], 9):
            text_cell(sheet.cell(line, 1), row["name"])
            text_cell(sheet.cell(line, 2), row["unit"])
            if row["quantity"] is None:
                text_cell(sheet.cell(line, 3), "Menge ungeklärt")
            else:
                sheet.cell(line, 3, float(Decimal(row["quantity"])))
                sheet.cell(line, 3).number_format = "0.00"
            person_hours = Decimal(row["actual_person_hours"])
            site_hours += person_hours
            sheet.cell(line, 4, float(person_hours))
            sheet.cell(line, 4).number_format = "0.00"
            evidence = "; ".join(
                f"{source['work_date'][8:10]}.{source['work_date'][5:7]}.{source['work_date'][:4]}"
                f" · MBG {source['report_id']}"
                + (f" · erfasst: {source['recorded_quantity']}" if row["quantity"] is None else "")
                for source in row.get("sources", [])
            )
            text_cell(sheet.cell(line, 5), evidence)
            for cell in sheet[line][:5]:
                cell.font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                cell.alignment = Alignment(
                    horizontal="right" if cell.column in (3, 4) else "left", vertical="top", wrap_text=True
                )
                cell.border = Border(
                    left=Side(style="thin", color=COLOR_NAVY_100),
                    right=Side(style="thin", color=COLOR_NAVY_100),
                    top=Side(style="thin", color=COLOR_NAVY_100),
                    bottom=Side(style="thin", color=COLOR_NAVY_100),
                )
                cell.fill = PatternFill("solid", fgColor=COLOR_WHITE if line % 2 else COLOR_GRAY_100)
            sheet.row_dimensions[line].height = max(
                _completion_text_height(row["name"], 50, minimum=32),
                _completion_text_height(evidence, 25, minimum=32),
            )
        total_line = 9 + len(site["rows"])
        text_cell(sheet.cell(total_line, 2), "Gesamt")
        sheet.cell(total_line, 2).font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
        sheet.cell(total_line, 4, float(site_hours)).font = Font(
            name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600
        )
        sheet.cell(total_line, 4).number_format = "0.00"
        for cell in sheet[total_line][:4]:
            cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
            cell.border = Border(
                top=Side(style="thin", color=COLOR_GREEN_600), bottom=Side(style="double", color=COLOR_NAVY_900)
            )
        site_unknown = site.get("unresolved_quantity_count", 0)
        if site_unknown:
            warning_line = total_line + 2
            sheet.merge_cells(start_row=warning_line, start_column=1, end_row=warning_line, end_column=5)
            text_cell(sheet.cell(warning_line, 1), f"Hinweis: Menge ungeklärt bei {site_unknown} Positionen.")
            sheet.cell(warning_line, 1).font = Font(name=FONT_FAMILY, size=9.5, italic=True, color="B45309")
            sheet.cell(warning_line, 1).fill = PatternFill("solid", fgColor="FEF3C7")
            sheet.row_dimensions[warning_line].height = 30
        for column, width in {"A": 56, "B": 16, "C": 14, "D": 24, "E": 28}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A9"
        sheet.print_title_rows = "1:8"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.2, footer=0.2)
        sheet.oddFooter.center.text = "Seite &P / &N"
        sheet.print_area = f"A1:E{total_line + 2 if site_unknown else total_line}"
    output = BytesIO()
    book.save(output)
    return output.getvalue()
