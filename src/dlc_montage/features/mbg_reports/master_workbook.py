# ruff: noqa: RUF001, RUF003
"""Feature adapter for building the master settlement workbook.

Strict rules:
- Zero ALL financial prices/rates (transport, hourly catalog, piecework);
- Includes only customer-billable work; excludes breaks and internal-only time;
- Zero in formulas too (no hardcoded 40 or 50);
- Historical master retains paid sites with live status from DB;
- Reusable service function returning openpyxl.Workbook or bytes.
"""

from io import BytesIO
from typing import Sequence

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from features.mbg_reports.reporting_selectors import load_accepted_reports_dataset
from features.mbg_reports.settlement_scope import project_settlement_dataset
from features.mbg_reports.sheet_titles import assign_site_sheet_titles
from features.projects.models import AcObject

# Brand colors & styling
COLOR_NAVY_900 = "0A2540"
COLOR_NAVY_700 = "1A3B5C"
COLOR_NAVY_100 = "E6EAF0"
COLOR_GREEN_600 = "00A86B"
COLOR_GREEN_50 = "E6FBF2"
COLOR_AMBER_700 = "B45309"
COLOR_AMBER_100 = "FEF3C7"
COLOR_GRAY_900 = "111827"
COLOR_GRAY_600 = "4B5563"
COLOR_GRAY_100 = "F8FAFC"
COLOR_WHITE = "FFFFFF"

FONT_FAMILY = "Segoe UI"
GERMAN_WEEKDAYS = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]

COMPANY_NAME = "DLC Montage"
COMPANY_CONTACT = "Dmitri Labodin"
COMPANY_PHONE = "Tel: +49 176 62203619"
COMPANY_EMAIL = "dmitri.labodin@googlemail.com"


def thin_border() -> Border:
    return Border(
        left=Side(style="thin", color=COLOR_NAVY_100),
        right=Side(style="thin", color=COLOR_NAVY_100),
        top=Side(style="thin", color=COLOR_NAVY_100),
        bottom=Side(style="thin", color=COLOR_NAVY_100),
    )


def apply_company_header(sheet, start_col: int = 8, end_col: int = 10) -> None:
    sheet.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
    c1 = sheet.cell(1, start_col, COMPANY_NAME)
    c1.font = Font(name=FONT_FAMILY, size=13, bold=True, color=COLOR_NAVY_900)
    c1.alignment = Alignment(horizontal="right", vertical="center")

    sheet.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
    c2 = sheet.cell(2, start_col, COMPANY_CONTACT)
    c2.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_600)
    c2.alignment = Alignment(horizontal="right", vertical="center")

    sheet.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
    c3 = sheet.cell(3, start_col, f"{COMPANY_PHONE}  ·  {COMPANY_EMAIL}")
    c3.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_600)
    c3.alignment = Alignment(horizontal="right", vertical="center")


def build_master_settlement_workbook(
    *, target_weeks: Sequence[tuple[int, int]] | None = None, dataset=None
) -> openpyxl.Workbook:
    """Сформировать общую сводную книгу Excel со всеми объектами и нулевыми финансовыми ставками."""
    wb = openpyxl.Workbook()

    # Сбор данных по принятым отчетам с сохранением оплаченных объектов
    if dataset is None:
        dataset = load_accepted_reports_dataset(
            target_weeks=list(target_weeks) if target_weeks else None, include_paid=True
        )
    dataset = project_settlement_dataset(dataset)
    all_sites_set = set(dataset.all_sites.keys())
    site_obj_map = {sid: AcObject.objects.get(pk=sid) for sid in all_sites_set}

    site_sheet_names = assign_site_sheet_titles(
        ((site_id, site.name) for site_id, site in site_obj_map.items()), reserved={"Übersicht"}
    )

    # ==========================================================================
    # ЛИСТ 1: Сводная таблица (Übersicht)
    # ==========================================================================
    ws_summary = wb.active
    ws_summary.title = "Übersicht"
    ws_summary.sheet_view.showGridLines = True

    # Data validation для колонки Zahlung
    dv_payment = DataValidation(
        type="list",
        formula1='"Bezahlt,Offen"',
        allow_blank=True,
    )
    ws_summary.add_data_validation(dv_payment)

    # Заголовок страницы
    ws_summary.merge_cells("A1:G1")
    t1 = ws_summary["A1"]
    t1.value = "DLC Montage · AC Gesamtabrechnung & Leistungsnachweis"
    t1.font = Font(name=FONT_FAMILY, size=15, bold=True, color=COLOR_NAVY_900)
    t1.alignment = Alignment(horizontal="left", vertical="center")

    ws_summary.merge_cells("A2:G2")
    t2 = ws_summary["A2"]
    t2.value = "Gesamtübersicht aller AC-Standorte nach Kalenderwochen"
    t2.font = Font(name=FONT_FAMILY, size=10, italic=True, color=COLOR_GRAY_600)
    t2.alignment = Alignment(horizontal="left", vertical="center")

    apply_company_header(ws_summary, start_col=8, end_col=10)

    ws_summary.row_dimensions[1].height = 24
    ws_summary.row_dimensions[2].height = 18
    ws_summary.row_dimensions[3].height = 10
    ws_summary.row_dimensions[4].height = 6

    # Статистика KW38+
    kw38_plus_sites = set()
    kw38_plus_hours = 0.0
    for period, sites_map in dataset.sites_by_period.items():
        if period[0] > 2026 or (period[0] == 2026 and period[1] >= 38):
            for s_id, s_summary in sites_map.items():
                kw38_plus_sites.add(s_id)
                kw38_plus_hours += float(s_summary.total_hours)

    completed_count = sum(1 for s_id in kw38_plus_sites if dataset.all_sites[s_id].is_completed)
    in_progress_count = len(kw38_plus_sites) - completed_count
    in_prog_names = ", ".join(
        sorted(
            dataset.all_sites[s_id].site_name for s_id in kw38_plus_sites if not dataset.all_sites[s_id].is_completed
        )
    )
    in_prog_suffix = f" ({in_prog_names})" if in_prog_names else ""

    ws_summary.merge_cells("A5:J5")
    card_info = ws_summary["A5"]
    card_info.value = (
        f"Im Abrechnungszeitraum bearbeitete Standorte: {len(kw38_plus_sites)} gesamt  |  "
        f"Vollständig abgeschlossen: {completed_count}  |  "
        f"In Bearbeitung: {in_progress_count}{in_prog_suffix}  |  "
        f"Gesamtstunden: {kw38_plus_hours:.2f} h  |  "
        f"Materiallogistik (0,00 €): 0,00 €  |  "
        f"Montagepreise: 0,00 €"
    )
    card_info.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
    card_info.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
    card_info.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_summary.row_dimensions[5].height = 22

    for col in range(1, 11):
        ws_summary.cell(5, col).border = thin_border()

    # --------------------------------------------------------------------------
    # Таблицы по неделям
    # --------------------------------------------------------------------------
    current_row = 7
    week_names = {
        36: "KW 36 (31.08.2026 – 04.09.2026)",
        37: "KW 37 (07.09.2026 – 11.09.2026)",
        38: "KW 38 (14.09.2026 – 18.09.2026) · [Zur Abrechnung]",
        39: "KW 39 (21.09.2026 – 25.09.2026) · [Zur Abrechnung]",
        40: "KW 40 (28.09.2026 – 02.10.2026) · [Zur Abrechnung]",
    }

    summary_headers = [
        "Nr.",
        "Standort (Objekt)",
        "Adresse",
        "Anlage (Projektgruppe)",
        "Status",
        "Arbeitstage",
        "Mitarbeiter",
        "Stunden (h)",
        "Betrag (€)",
        "Zahlung",
    ]

    all_week_total_rows = []
    payment_cell_ranges = []

    for iso_year, iso_week in dataset.periods:
        sites_in_period = dataset.sites_by_period.get((iso_year, iso_week), {})
        if not sites_in_period:
            continue

        period_title = (
            week_names.get(iso_week, f"KW {iso_week} / {iso_year}")
            if iso_year == 2026
            else f"KW {iso_week} / {iso_year}"
        )
        ws_summary.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        sec_cell = ws_summary.cell(current_row, 1, period_title)
        sec_cell.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        sec_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_700 if iso_week < 38 else COLOR_NAVY_900)
        sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_summary.row_dimensions[current_row].height = 24
        current_row += 1

        for col_idx, h_text in enumerate(summary_headers, 1):
            cell = ws_summary.cell(current_row, col_idx, h_text)
            cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_700)
            cell.alignment = Alignment(
                horizontal="center" if col_idx in (1, 5, 6, 8, 9, 10) else "left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border()
        ws_summary.row_dimensions[current_row].height = 24
        current_row += 1

        start_kw_row = current_row
        sorted_sites = sorted(sites_in_period.items(), key=lambda x: x[1].site_name)

        for idx, (s_id, s_summary) in enumerate(sorted_sites, 1):
            _site = site_obj_map[s_id]
            status_text = s_summary.construction_status
            work_days = s_summary.work_days_count
            emp_str = ", ".join(s_summary.unique_employee_names) if s_summary.unique_employee_names else "—"
            site_hours = float(s_summary.total_hours)
            # Вся стоимость строго 0.00 €
            site_amount = 0.00

            sheet_target = site_sheet_names.get(s_id, "Übersicht")

            ws_summary.cell(current_row, 1, idx).alignment = Alignment(horizontal="center", vertical="top")

            obj_cell = ws_summary.cell(current_row, 2, s_summary.site_name)
            obj_cell.hyperlink = f"#'{sheet_target}'!A1"
            obj_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900, underline="single")
            obj_cell.alignment = Alignment(horizontal="left", vertical="top")

            ws_summary.cell(current_row, 3, s_summary.address or "—").alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )
            ws_summary.cell(current_row, 4, s_summary.group_name or "—").alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

            status_cell = ws_summary.cell(current_row, 5, status_text)
            status_cell.alignment = Alignment(horizontal="center", vertical="top")
            if not s_summary.is_completed:
                status_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700)
                status_cell.fill = PatternFill("solid", fgColor=COLOR_AMBER_100)
            else:
                status_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600)
                status_cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)

            ws_summary.cell(current_row, 6, work_days).alignment = Alignment(horizontal="center", vertical="top")
            ws_summary.cell(current_row, 7, emp_str).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )

            h_cell = ws_summary.cell(current_row, 8, site_hours)
            h_cell.number_format = "0.00"
            h_cell.alignment = Alignment(horizontal="right", vertical="top")

            a_cell = ws_summary.cell(current_row, 9, site_amount)
            a_cell.number_format = "#,##0.00 €"
            a_cell.alignment = Alignment(horizontal="right", vertical="top")

            # Колонка 10: Zahlung (Offen / Bezahlt с выпадающим списком из статуса БД)
            is_site_paid = _site.is_paid
            pay_label = "Bezahlt" if is_site_paid else "Offen"
            pay_cell = ws_summary.cell(current_row, 10, pay_label)
            pay_cell.alignment = Alignment(horizontal="center", vertical="top")
            if is_site_paid:
                pay_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600)
                pay_cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
            else:
                pay_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700)
                pay_cell.fill = PatternFill("solid", fgColor=COLOR_AMBER_100)

            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if idx % 2 != 0 else COLOR_GRAY_100)
            for c in range(1, 11):
                cell_item = ws_summary.cell(current_row, c)
                if c not in (5, 10):
                    cell_item.fill = row_fill
                cell_item.border = thin_border()
                if c not in (2, 5, 10):
                    cell_item.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

            ws_summary.row_dimensions[current_row].height = max(24, 15 * max(1, len(emp_str) // 30 + 1))
            current_row += 1

        end_kw_row = current_row - 1
        payment_range_str = f"J{start_kw_row}:J{end_kw_row}"
        dv_payment.add(payment_range_str)
        payment_cell_ranges.append(payment_range_str)

        # Подитог недели
        ws_summary.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        sub_label = ws_summary.cell(current_row, 1, f"Zwischensumme {period_title.split(' · ')[0]}")
        sub_label.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)
        sub_label.alignment = Alignment(horizontal="right", vertical="center", indent=1)

        kw_h_sum = ws_summary.cell(current_row, 8, f"=SUM(H{start_kw_row}:H{end_kw_row})")
        kw_h_sum.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)
        kw_h_sum.number_format = "0.00"
        kw_h_sum.alignment = Alignment(horizontal="right", vertical="center")

        kw_a_sum = ws_summary.cell(current_row, 9, f"=SUM(I{start_kw_row}:I{end_kw_row})")
        kw_a_sum.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)
        kw_a_sum.number_format = "#,##0.00 €"
        kw_a_sum.alignment = Alignment(horizontal="right", vertical="center")

        ws_summary.cell(current_row, 10, "")
        all_week_total_rows.append(current_row)

        for c in range(1, 11):
            cc = ws_summary.cell(current_row, c)
            cc.fill = PatternFill("solid", fgColor=COLOR_NAVY_100)
            cc.border = Border(
                top=Side(style="thin", color=COLOR_NAVY_700),
                bottom=Side(style="thin", color=COLOR_NAVY_700),
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
            )
        ws_summary.row_dimensions[current_row].height = 24
        current_row += 2

    # Итоговый общий итог по всем неделям
    ws_summary.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
    tot_label = ws_summary.cell(current_row, 1, "Gesamtsumme aller Wochen")
    tot_label.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY_900)
    tot_label.alignment = Alignment(horizontal="right", vertical="center", indent=1)

    formula_h = "+".join(f"H{r}" for r in all_week_total_rows) if all_week_total_rows else "0"
    tot_h = ws_summary.cell(current_row, 8, f"={formula_h}")
    tot_h.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
    tot_h.number_format = "0.00"
    tot_h.alignment = Alignment(horizontal="right", vertical="center")

    formula_a = "+".join(f"I{r}" for r in all_week_total_rows) if all_week_total_rows else "0"
    tot_a = ws_summary.cell(current_row, 9, f"={formula_a}")
    tot_a.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_GREEN_600)
    tot_a.number_format = "#,##0.00 €"
    tot_a.alignment = Alignment(horizontal="right", vertical="center")

    ws_summary.cell(current_row, 10, "")
    for c in range(1, 11):
        cc = ws_summary.cell(current_row, c)
        cc.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
        cc.border = Border(
            top=Side(style="thin", color=COLOR_GREEN_600),
            bottom=Side(style="double", color=COLOR_NAVY_900),
            left=Side(style="thin", color=COLOR_NAVY_100),
            right=Side(style="thin", color=COLOR_NAVY_100),
        )
    ws_summary.row_dimensions[current_row].height = 28

    # Условное форматирование для Bezahlt / Offen
    rule_bezahlt = CellIsRule(
        operator="equal",
        formula=['"Bezahlt"'],
        fill=PatternFill("solid", fgColor=COLOR_GREEN_50),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600),
    )
    rule_offen = CellIsRule(
        operator="equal",
        formula=['"Offen"'],
        fill=PatternFill("solid", fgColor=COLOR_AMBER_100),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700),
    )

    for prange in payment_cell_ranges:
        ws_summary.conditional_formatting.add(prange, rule_bezahlt)
        ws_summary.conditional_formatting.add(prange, rule_offen)

    col_widths = {
        "A": 6,
        "B": 26,
        "C": 30,
        "D": 32,
        "E": 17,
        "F": 12,
        "G": 38,
        "H": 14,
        "I": 16,
        "J": 14,
    }
    for col_letter, width in col_widths.items():
        ws_summary.column_dimensions[col_letter].width = width

    ws_summary.freeze_panes = "A7"
    ws_summary.page_setup.orientation = "landscape"
    ws_summary.page_setup.paperSize = ws_summary.PAPERSIZE_A4
    ws_summary.page_setup.fitToWidth = 1
    ws_summary.page_setup.fitToHeight = 0
    ws_summary.sheet_properties.pageSetUpPr.fitToPage = True

    # ==========================================================================
    # ЛИСТЫ ПО КАЖДОМУ ОБЪЕКТУ (Standorte)
    # ==========================================================================
    sorted_all_sites = sorted(dataset.all_sites.items(), key=lambda x: x[1].site_name)

    for s_id, s_summary in sorted_all_sites:
        _site = site_obj_map[s_id]
        sheet_title = site_sheet_names[s_id]
        ws = wb.create_sheet(title=sheet_title)
        ws.sheet_view.showGridLines = True

        site_periods = sorted(list(set((d.iso_year, d.iso_week) for d in s_summary.day_facts)))
        kw_display_str = ", ".join(f"KW {k} ({y})" if y != 2026 else f"KW {k}" for (y, k) in site_periods)
        first_date = s_summary.day_facts[0].work_date.strftime("%d.%m.%Y") if s_summary.day_facts else "—"
        last_date = s_summary.day_facts[-1].work_date.strftime("%d.%m.%Y") if s_summary.day_facts else "—"
        period_str = f"{first_date} – {last_date} ({kw_display_str})"

        site_status_label = s_summary.construction_status
        is_in_prog = not s_summary.is_completed

        ws.merge_cells("A1:E1")
        ws["A1"] = f"DLC Montage · Standort-Abrechnung: {s_summary.site_name}"
        ws["A1"].font = Font(name=FONT_FAMILY, size=15, bold=True, color=COLOR_NAVY_900)

        ws.cell(1, 6, "← Zurück zur Übersicht").hyperlink = "#'Übersicht'!A1"
        ws.cell(1, 6).font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_700, underline="single")
        ws.cell(1, 6).alignment = Alignment(horizontal="right", vertical="center")

        apply_company_header(ws, start_col=7, end_col=8)

        meta_rows = [
            (2, "Standort:", s_summary.site_name, "Anlage (Gruppe):", s_summary.group_name or "—"),
            (3, "Adresse:", s_summary.address or "—", "Projekt-Nr:", s_summary.project_number or "—"),
            (4, "Zeitraum:", period_str, "Bau-Status:", site_status_label),
        ]
        for r_idx, l1, v1, l2, v2 in meta_rows:
            ws.cell(r_idx, 1, l1).font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_700)
            ws.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=4)
            c_v1 = ws.cell(r_idx, 2, v1)
            c_v1.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_900)

            ws.cell(r_idx, 5, l2).font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_700)
            ws.cell(r_idx, 6, v2).font = Font(
                name=FONT_FAMILY,
                size=9.5,
                bold=True if l2 == "Bau-Status:" else False,
                color=COLOR_AMBER_700
                if (l2 == "Bau-Status:" and is_in_prog)
                else (COLOR_GREEN_600 if l2 == "Bau-Status:" else COLOR_GRAY_900),
            )
            ws.row_dimensions[r_idx].height = 18

        # ----------------------------------------------------------------------
        # БЛОК 1 (СВЕРХУ): Почасовой отчёт (Stundennachweis)
        # ----------------------------------------------------------------------
        ws.merge_cells("A6:H6")
        b1_title = ws["A6"]
        b1_title.value = "1. STUNDENNACHWEIS (Abrechenbare Arbeitszeit nach Kalenderwochen & Tagen)"
        b1_title.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        b1_title.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        b1_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[6].height = 24

        h_cols = ["KW", "Datum", "Wochentag", "Mitarbeiter", "Beginn", "Ende", "Ist (h)", "Beschreibung der Tätigkeit"]
        for c_idx, c_name in enumerate(h_cols, 1):
            c_cell = ws.cell(7, c_idx, c_name)
            c_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
            c_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_700)
            c_cell.alignment = Alignment(
                horizontal="center" if c_idx in (1, 2, 3, 5, 6, 7) else "left", vertical="center"
            )
            c_cell.border = thin_border()
        ws.row_dimensions[7].height = 22

        row_h = 8
        start_h_row = row_h

        for day in s_summary.day_facts:
            r_kw = f"KW {day.iso_week}"
            r_date_str = day.work_date.strftime("%d.%m.%Y")
            r_wday = GERMAN_WEEKDAYS[day.work_date.weekday()]

            for entry in day.entries:
                p_str = (
                    ", ".join(entry.participant_names) if entry.participant_names else ", ".join(day.workforce_names)
                )
                dur = float(entry.actual_person_hours)

                ws.cell(row_h, 1, r_kw).alignment = Alignment(horizontal="center", vertical="top")
                ws.cell(row_h, 2, r_date_str).alignment = Alignment(horizontal="center", vertical="top")
                ws.cell(row_h, 3, r_wday).alignment = Alignment(horizontal="center", vertical="top")
                ws.cell(row_h, 4, p_str).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

                b_time = entry.starts_at or "—"
                e_time = entry.ends_at or "—"
                ws.cell(row_h, 5, b_time).alignment = Alignment(horizontal="center", vertical="top")
                ws.cell(row_h, 6, e_time).alignment = Alignment(horizontal="center", vertical="top")

                dur_cell = ws.cell(row_h, 7, dur)
                dur_cell.number_format = "0.00"
                dur_cell.alignment = Alignment(horizontal="right", vertical="top")

                desc_cell = ws.cell(row_h, 8, entry.description)
                desc_cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

                row_fill = PatternFill("solid", fgColor=COLOR_WHITE if row_h % 2 == 0 else COLOR_GRAY_100)
                for c in range(1, 9):
                    cc = ws.cell(row_h, c)
                    cc.fill = row_fill
                    cc.border = thin_border()
                    cc.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

                ws.row_dimensions[row_h].height = max(22, 14 * max(1, len(entry.description) // 55 + 1))
                row_h += 1

        end_h_row = row_h - 1

        # Итог по часам
        ws.merge_cells(start_row=row_h, start_column=1, end_row=row_h, end_column=6)
        ws.cell(row_h, 1, "Summe Stunden (Stundennachweis)").alignment = Alignment(
            horizontal="right", vertical="center", indent=1
        )
        ws.cell(row_h, 1).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)

        sum_formula = f"=SUM(G{start_h_row}:G{end_h_row})" if end_h_row >= start_h_row else "0.00"
        sum_h_cell = ws.cell(row_h, 7, sum_formula)
        sum_h_cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_GREEN_600)
        sum_h_cell.number_format = "0.00"
        sum_h_cell.alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row_h, 8, "Stunden Gesamt").font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GRAY_600)

        for c in range(1, 9):
            cc = ws.cell(row_h, c)
            cc.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
            cc.border = Border(
                top=Side(style="thin", color=COLOR_GREEN_600),
                bottom=Side(style="double", color=COLOR_NAVY_900),
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
            )
        ws.row_dimensions[row_h].height = 24
        row_h += 2

        # ----------------------------------------------------------------------
        # БЛОК 2 (СНИЗУ): Сдельный отчёт / Выполненные объёмы (Leistungsnachweis)
        # ----------------------------------------------------------------------
        ws.merge_cells(start_row=row_h, start_column=1, end_row=row_h, end_column=8)
        b2_title = ws.cell(row_h, 1, "2. LEISTUNGSNACHWEIS (Erbrachte Montageleistungen / Werkvertrag)")
        b2_title.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        b2_title.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        b2_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row_h].height = 24
        row_h += 1

        l_cols = [
            "Pos.",
            "Leistung / Montageposition",
            "Einheit",
            "Menge",
            "Ist (h)",
            "Einzelpreis (€)",
            "Gesamtbetrag (€)",
            "Status",
        ]
        for c_idx, c_name in enumerate(l_cols, 1):
            c_cell = ws.cell(row_h, c_idx, c_name)
            c_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
            c_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_700)
            c_cell.alignment = Alignment(
                horizontal="center" if c_idx in (1, 3, 4, 5, 6, 7, 8) else "left", vertical="center"
            )
            c_cell.border = thin_border()
        ws.row_dimensions[row_h].height = 22
        row_h += 1

        start_l_row = row_h
        perf_rows = s_summary.position_rows

        for p_idx, prow in enumerate(perf_rows, 1):
            ws.cell(row_h, 1, p_idx).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row_h, 2, prow.name).alignment = Alignment(horizontal="left", vertical="top")
            ws.cell(row_h, 3, prow.unit or "Stk.").alignment = Alignment(horizontal="center", vertical="top")

            if prow.quantity is not None:
                qty_cell = ws.cell(row_h, 4, float(prow.quantity))
                qty_cell.alignment = Alignment(horizontal="right", vertical="top")
                qty_cell.number_format = "0.##"
            else:
                qty_cell = ws.cell(row_h, 4, "—")
                qty_cell.alignment = Alignment(horizontal="center", vertical="top")

            hours_val = float(prow.actual_person_hours) if prow.actual_person_hours is not None else 0.0
            ws.cell(row_h, 5, hours_val).alignment = Alignment(horizontal="right", vertical="top")
            ws.cell(row_h, 5).number_format = "0.00"

            # Все расценки строго 0.00 €
            ws.cell(row_h, 6, 0.0).alignment = Alignment(horizontal="right", vertical="top")
            ws.cell(row_h, 6).number_format = "#,##0.00 €"

            ws.cell(row_h, 7, 0.0).alignment = Alignment(horizontal="right", vertical="top")
            ws.cell(row_h, 7).number_format = "#,##0.00 €"

            ws.cell(row_h, 8, site_status_label).alignment = Alignment(horizontal="center", vertical="top")

            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if row_h % 2 == 0 else COLOR_GRAY_100)
            for c in range(1, 9):
                cc = ws.cell(row_h, c)
                cc.fill = row_fill
                cc.border = thin_border()
                cc.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

            ws.row_dimensions[row_h].height = 20
            row_h += 1

        end_l_row = row_h - 1

        # Итог по сдельщине
        ws.merge_cells(start_row=row_h, start_column=1, end_row=row_h, end_column=4)
        ws.cell(row_h, 1, "Summe erbrachte Leistungen (Werkvertrag)").alignment = Alignment(
            horizontal="right", vertical="center", indent=1
        )
        ws.cell(row_h, 1).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)

        sum_hours_formula = f"=SUM(E{start_l_row}:E{end_l_row})" if end_l_row >= start_l_row else "0.00"
        sum_lh_cell = ws.cell(row_h, 5, sum_hours_formula)
        sum_lh_cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_NAVY_900)
        sum_lh_cell.number_format = "0.00"
        sum_lh_cell.alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row_h, 6, "").alignment = Alignment(horizontal="right", vertical="center")

        sum_amount_formula = f"=SUM(G{start_l_row}:G{end_l_row})" if end_l_row >= start_l_row else "0.00"
        sum_la_cell = ws.cell(row_h, 7, sum_amount_formula)
        sum_la_cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_GREEN_600)
        sum_la_cell.number_format = "#,##0.00 €"
        sum_la_cell.alignment = Alignment(horizontal="right", vertical="center")

        ws.cell(row_h, 8, "").alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 9):
            cc = ws.cell(row_h, c)
            cc.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
            cc.border = Border(
                top=Side(style="thin", color=COLOR_GREEN_600),
                bottom=Side(style="double", color=COLOR_NAVY_900),
                left=Side(style="thin", color=COLOR_NAVY_100),
                right=Side(style="thin", color=COLOR_NAVY_100),
            )
        ws.row_dimensions[row_h].height = 24

        site_col_widths = {
            "A": 9,
            "B": 13,
            "C": 12,
            "D": 30,
            "E": 10,
            "F": 10,
            "G": 14,
            "H": 48,
        }
        for col_l, width in site_col_widths.items():
            ws.column_dimensions[col_l].width = width

        ws.freeze_panes = "A8"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    return wb


def build_master_settlement_bytes(*, target_weeks: Sequence[tuple[int, int]] | None = None) -> bytes:
    """Сформировать байты общей сводной книги Excel."""
    wb = build_master_settlement_workbook(target_weeks=target_weeks)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
