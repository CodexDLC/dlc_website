"""Безценовые клиентские XLSX-документы: акт выполненных работ и табель MBG."""

import re
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .selectors import accepted_completion_rows, accepted_daily_work, customer_site_detail

FONT_FAMILY = "Segoe UI"
COLOR_NAVY_900 = "0A2540"
COLOR_NAVY_700 = "1E3A5F"
COLOR_NAVY_100 = "E6EAF0"
COLOR_GREEN_600 = "00A86B"
COLOR_GREEN_50 = "E6FBF2"
COLOR_GRAY_900 = "111827"
COLOR_GRAY_600 = "4B5563"
COLOR_GRAY_100 = "F8FAFC"
COLOR_WHITE = "FFFFFF"
COLOR_AMBER_50 = "FFFBEB"
COLOR_AMBER_700 = "B45309"

THIN_BORDER_SIDE = Side(style="thin", color=COLOR_NAVY_100)
BORDER_ROW = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
BORDER_TOTAL = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=Side(style="thin", color=COLOR_NAVY_700),
    bottom=Side(style="double", color=COLOR_NAVY_700),
)


def _safe_sheet_title(name: str, used_titles: set[str] | None = None) -> str:
    """Подготовить безопасное уникальное название листа Excel (до 31 символа)."""
    cleaned = re.sub(r"[\\/?*:[\]]", "_", str(name or "")).strip()
    if not cleaned:
        cleaned = "Standort"
    base = cleaned[:28].rstrip() if len(cleaned) > 28 else cleaned
    title = base
    if used_titles is not None:
        counter = 1
        while title in used_titles:
            counter += 1
            suffix = f" {counter}"
            title = f"{cleaned[: 31 - len(suffix)].rstrip()}{suffix}"
        used_titles.add(title)
    return title


def _text_cell(
    sheet,
    row,
    column,
    value,
    *,
    bold=False,
    color=None,
    align="left",
    bg_color=None,
    border=None,
    size=10,
):
    """Записать внешний текст строкой, предотвращая выполнение формул."""
    cell = sheet.cell(row=row, column=column, value=str(value) if value is not None else "")
    cell.data_type = "s"
    cell.font = Font(name=FONT_FAMILY, size=size, bold=bold, color=color or COLOR_GRAY_900)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    if bg_color:
        cell.fill = PatternFill("solid", fgColor=bg_color)
    if border:
        cell.border = border
    return cell


def _number_cell(
    sheet,
    row,
    column,
    value,
    *,
    decimals=2,
    bold=False,
    color=None,
    bg_color=None,
    border=None,
    size=10,
):
    """Записать числовое значение с сохранением числового типа в Excel."""
    if value is None or str(value).lower() in ("ungeklärt", "nicht bestätigt", "не подтверждено", "—", "-"):
        val_str = str(value) if value is not None else "ungeklärt"
        return _text_cell(
            sheet,
            row,
            column,
            val_str,
            bold=bold,
            color=color or COLOR_GRAY_600,
            align="right",
            bg_color=bg_color,
            border=border,
            size=size,
        )
    try:
        num = float(value) if isinstance(value, (Decimal, str)) else value
    except (ValueError, TypeError):
        return _text_cell(
            sheet,
            row,
            column,
            value,
            bold=bold,
            color=color,
            align="right",
            bg_color=bg_color,
            border=border,
            size=size,
        )
    cell = sheet.cell(row=row, column=column, value=num)
    cell.data_type = "n"
    cell.number_format = "#,##0.00" if decimals == 2 else ("#,##0" if decimals == 0 else "0")
    cell.font = Font(name=FONT_FAMILY, size=size, bold=bold, color=color or COLOR_GRAY_900)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    if bg_color:
        cell.fill = PatternFill("solid", fgColor=bg_color)
    if border:
        cell.border = border
    return cell


def _date_cell(
    sheet,
    row,
    column,
    value,
    *,
    bold=False,
    color=None,
    bg_color=None,
    border=None,
    size=10,
):
    """Записать дату с сохранением типа и формата отображения."""
    date_val = value
    if isinstance(value, str):
        try:
            date_val = date.fromisoformat(value)
        except (ValueError, TypeError):
            return _text_cell(
                sheet,
                row,
                column,
                value,
                bold=bold,
                color=color,
                align="center",
                bg_color=bg_color,
                border=border,
                size=size,
            )
    cell = sheet.cell(row=row, column=column, value=date_val)
    if isinstance(date_val, (date, datetime)):
        cell.number_format = "YYYY-MM-DD"
    cell.font = Font(name=FONT_FAMILY, size=size, bold=bold, color=color or COLOR_GRAY_900)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if bg_color:
        cell.fill = PatternFill("solid", fgColor=bg_color)
    if border:
        cell.border = border
    return cell


def _apply_table_header(sheet, row, headers):
    """Оформить заголовки таблицы в едином стиле DLC."""
    sheet.row_dimensions[row].height = 24
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER_ROW


def _format_sheet_view(sheet, *, landscape=True):
    """Включить сетку и компактную настройку печати."""
    sheet.sheet_view.showGridLines = True
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _site_sheet(workbook, site, title, *, iso_year=None, iso_week=None):
    """Сформировать лист акта выполненных работ без цен для одного Standort."""
    detail = customer_site_detail(site)
    daily_works = accepted_daily_work(site, iso_year=iso_year, iso_week=iso_week)
    completed_rows = (
        accepted_completion_rows(site, iso_year=iso_year, iso_week=iso_week)
        if (iso_year and iso_week)
        else detail.completed_rows
    )

    sheet = workbook.create_sheet(title)
    _format_sheet_view(sheet, landscape=True)

    # Заголовок объекта
    sheet.merge_cells("A1:F1")
    _text_cell(
        sheet, 1, 1, f"Leistungsnachweis · {site.name}", bold=True, color=COLOR_NAVY_900, size=14, align="left"
    )
    sheet.row_dimensions[1].height = 28

    period_str = f"KW {iso_week:02d} / {iso_year}" if (iso_year and iso_week) else "Gesamter Zeitraum"
    info_line = f"Anlage: {site.group.name}  |  Projekt-Nr.: {site.project_number or '—'}  |  Zeitraum: {period_str}"
    sheet.merge_cells("A2:F2")
    _text_cell(sheet, 2, 1, info_line, color=COLOR_GRAY_600, size=10, align="left")

    current_row = 3
    if detail.plan.excluded_report_count:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        _text_cell(
            sheet,
            current_row,
            1,
            f"Unvollständig: {detail.plan.excluded_report_count} Berichte ohne aktuellen Nachweis",
            bold=True,
            color=COLOR_AMBER_700,
            bg_color=COLOR_AMBER_50,
        )
        current_row += 1

    if detail.review_stale:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        _text_cell(
            sheet,
            current_row,
            1,
            "Hinweis: Veralteter Freigabestand / Abweichung zum aktuellen Arbeitsstand",
            bold=True,
            color=COLOR_AMBER_700,
            bg_color=COLOR_AMBER_50,
        )
        current_row += 1

    current_row += 1

    # Раздел 1: План работ
    sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
    _text_cell(sheet, current_row, 1, "1. Planpositionen und Freigabestand", bold=True, color=COLOR_NAVY_700, size=11)
    current_row += 1

    plan_headers = ["Position", "Plan", "Bestätigt", "Rest", "Einheit", "Status"]
    _apply_table_header(sheet, current_row, plan_headers)
    current_row += 1

    for row_item in detail.plan.rows:
        sheet.row_dimensions[current_row].height = 20
        _text_cell(sheet, current_row, 1, row_item.item.label, border=BORDER_ROW)
        _number_cell(sheet, current_row, 2, row_item.item.planned_quantity, border=BORDER_ROW)
        _number_cell(sheet, current_row, 3, row_item.completed_quantity, border=BORDER_ROW)
        _number_cell(sheet, current_row, 4, row_item.remaining_quantity, border=BORDER_ROW)
        _text_cell(sheet, current_row, 5, row_item.item.unit_snapshot, align="center", border=BORDER_ROW)
        _text_cell(sheet, current_row, 6, row_item.status, align="center", border=BORDER_ROW)
        current_row += 1

    current_row += 1

    # Раздел 2: Подтверждённые физические объёмы
    sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    _text_cell(sheet, current_row, 1, "2. Erbrachte Montageleistungen", bold=True, color=COLOR_NAVY_700, size=11)
    current_row += 1

    completed_headers = ["Erbrachte Arbeiten", "Menge", "Einheit", "Personenstunden", "Nachweis"]
    _apply_table_header(sheet, current_row, completed_headers)
    current_row += 1

    total_completed_hours = Decimal("0")
    for completed in completed_rows:
        sheet.row_dimensions[current_row].height = 20
        _text_cell(sheet, current_row, 1, completed.name, border=BORDER_ROW)
        _number_cell(
            sheet,
            current_row,
            2,
            completed.quantity if completed.quantity is not None else "ungeklärt",
            border=BORDER_ROW,
        )
        _text_cell(sheet, current_row, 3, completed.unit, align="center", border=BORDER_ROW)
        _number_cell(sheet, current_row, 4, completed.actual_person_hours, border=BORDER_ROW)
        _text_cell(sheet, current_row, 5, completed.kind, align="center", border=BORDER_ROW)
        total_completed_hours += completed.actual_person_hours
        current_row += 1

    # Итог раздела 2
    sheet.row_dimensions[current_row].height = 22
    _text_cell(sheet, current_row, 1, "Gesamt Montageleistungen", bold=True, border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 2, "", border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 3, "", border=BORDER_TOTAL)
    _number_cell(
        sheet,
        current_row,
        4,
        total_completed_hours,
        bold=True,
        bg_color=COLOR_GREEN_50,
        color=COLOR_NAVY_900,
        border=BORDER_TOTAL,
    )
    _text_cell(sheet, current_row, 5, "", border=BORDER_TOTAL)
    current_row += 2

    # Раздел 3: Дневные записи
    sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    _text_cell(sheet, current_row, 1, "3. Angenommene Tagesnachweise", bold=True, color=COLOR_NAVY_700, size=11)
    current_row += 1

    daily_headers = ["Datum", "Arbeit", "Menge", "Einheit", "Personenstunden"]
    _apply_table_header(sheet, current_row, daily_headers)
    current_row += 1

    total_daily_hours = Decimal("0")
    for work in daily_works:
        sheet.row_dimensions[current_row].height = 20
        _date_cell(sheet, current_row, 1, work.work_date, border=BORDER_ROW)
        _text_cell(sheet, current_row, 2, work.description, border=BORDER_ROW)
        _number_cell(sheet, current_row, 3, work.quantity, border=BORDER_ROW)
        _text_cell(sheet, current_row, 4, work.unit, align="center", border=BORDER_ROW)
        _number_cell(sheet, current_row, 5, work.actual_person_hours, border=BORDER_ROW)
        total_daily_hours += work.actual_person_hours
        current_row += 1

    # Итог дневных записей
    sheet.row_dimensions[current_row].height = 22
    _text_cell(sheet, current_row, 1, "Summe Personenstunden Tagesberichte", bold=True, border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 2, "", border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 3, "", border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 4, "", border=BORDER_TOTAL)
    _number_cell(
        sheet,
        current_row,
        5,
        total_daily_hours,
        bold=True,
        bg_color=COLOR_GREEN_50,
        color=COLOR_NAVY_900,
        border=BORDER_TOTAL,
    )

    sheet.freeze_panes = "A5"
    widths = {"A": 36, "B": 24, "C": 20, "D": 18, "E": 18, "F": 16}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    return sheet


def site_workbook(site, *, iso_year=None, iso_week=None):
    """Сформировать отдельный XLSX Standort без коммерческих полей."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _site_sheet(workbook, site, "Arbeiten", iso_year=iso_year, iso_week=iso_week)
    data = BytesIO()
    workbook.save(data)
    return data.getvalue()


def group_workbook(group):
    """Сформировать обзор Anlage и листы её Standorte без цен."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Übersicht"
    _format_sheet_view(summary, landscape=True)

    summary.merge_cells("A1:F1")
    _text_cell(
        summary, 1, 1, f"Leistungsübersicht · {group.name}", bold=True, color=COLOR_NAVY_900, size=14, align="left"
    )
    summary.row_dimensions[1].height = 28

    client_name = group.client.name if group.client else "—"
    summary.merge_cells("A2:F2")
    _text_cell(
        summary,
        2,
        1,
        f"Auftraggeber: {client_name}  |  Standorte: {group.ac_objects.count()}  |  Gesamter Zeitraum",
        color=COLOR_GRAY_600,
        size=10,
        align="left",
    )

    summary_headers = [
        "Standort",
        "Projektnummer",
        "Status",
        "Planpositionen",
        "Angenommene Arbeiten",
        "Personenstunden gesamt",
    ]
    _apply_table_header(summary, 4, summary_headers)

    used_sheet_titles = {"Übersicht"}
    total_plan_items = 0
    total_accepted_works = 0
    total_group_hours = Decimal("0")

    sites = group.ac_objects.all().order_by("name", "pk")
    for number, site in enumerate(sites, start=1):
        row = number + 4
        summary.row_dimensions[row].height = 20
        detail = customer_site_detail(site)
        daily_works = accepted_daily_work(site)
        site_hours = sum((w.actual_person_hours for w in daily_works), Decimal("0"))

        plan_count = len(detail.plan.rows)
        works_count = len(daily_works)

        total_plan_items += plan_count
        total_accepted_works += works_count
        total_group_hours += site_hours

        _text_cell(summary, row, 1, site.name, border=BORDER_ROW)
        _text_cell(summary, row, 2, site.project_number or "—", align="center", border=BORDER_ROW)
        _text_cell(summary, row, 3, site.get_status_display(), align="center", border=BORDER_ROW)
        _number_cell(summary, row, 4, plan_count, decimals=0, border=BORDER_ROW)
        _number_cell(summary, row, 5, works_count, decimals=0, border=BORDER_ROW)
        _number_cell(summary, row, 6, site_hours, decimals=2, border=BORDER_ROW)

        site_title = _safe_sheet_title(site.name, used_sheet_titles)
        _site_sheet(workbook, site, site_title)

    # Итог по группе
    total_row = len(sites) + 5
    summary.row_dimensions[total_row].height = 24
    _text_cell(summary, total_row, 1, "Gesamt Anlage", bold=True, border=BORDER_TOTAL)
    _text_cell(summary, total_row, 2, "", border=BORDER_TOTAL)
    _text_cell(summary, total_row, 3, "", border=BORDER_TOTAL)
    _number_cell(summary, total_row, 4, total_plan_items, decimals=0, bold=True, border=BORDER_TOTAL)
    _number_cell(summary, total_row, 5, total_accepted_works, decimals=0, bold=True, border=BORDER_TOTAL)
    _number_cell(
        summary,
        total_row,
        6,
        total_group_hours,
        decimals=2,
        bold=True,
        bg_color=COLOR_GREEN_50,
        color=COLOR_NAVY_900,
        border=BORDER_TOTAL,
    )

    summary.freeze_panes = "A5"
    for col, width in {"A": 36, "B": 20, "C": 20, "D": 18, "E": 22, "F": 24}.items():
        summary.column_dimensions[col].width = width

    data = BytesIO()
    workbook.save(data)
    return data.getvalue()


def _site_timesheet_sheet(workbook, site, title, *, iso_year=None, iso_week=None):
    """Сформировать отдельный лист табеля MBG с фактическими человеко-часами."""
    detail = customer_site_detail(site)
    daily_works = accepted_daily_work(site, iso_year=iso_year, iso_week=iso_week)

    sheet = workbook.create_sheet(title)
    _format_sheet_view(sheet, landscape=True)

    sheet.merge_cells("A1:E1")
    _text_cell(
        sheet, 1, 1, f"MBG-Stundennachweis · {site.name}", bold=True, color=COLOR_NAVY_900, size=14, align="left"
    )
    sheet.row_dimensions[1].height = 28

    period_str = f"KW {iso_week:02d} / {iso_year}" if (iso_year and iso_week) else "Gesamter Zeitraum"
    info_line = (
        f"Anlage: {site.group.name}  |  Projekt-Nr.: {site.project_number or '—'}  |  "
        f"Status: {site.get_status_display()}  |  Zeitraum: {period_str}"
    )
    sheet.merge_cells("A2:E2")
    _text_cell(sheet, 2, 1, info_line, color=COLOR_GRAY_600, size=10, align="left")

    current_row = 3
    if detail.plan.excluded_report_count:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        _text_cell(
            sheet,
            current_row,
            1,
            f"Unvollständig: {detail.plan.excluded_report_count} Berichte ohne aktuellen Nachweis",
            bold=True,
            color=COLOR_AMBER_700,
            bg_color=COLOR_AMBER_50,
        )
        current_row += 1

    if detail.review_stale:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        _text_cell(
            sheet,
            current_row,
            1,
            "Hinweis: Veralteter Freigabestand / Abweichung zum aktuellen Arbeitsstand",
            bold=True,
            color=COLOR_AMBER_700,
            bg_color=COLOR_AMBER_50,
        )
        current_row += 1

    current_row += 1
    headers = ["Datum", "Arbeit / Tätigkeit", "Menge", "Einheit", "Personenstunden"]
    _apply_table_header(sheet, current_row, headers)
    current_row += 1

    total_hours = Decimal("0")
    for work in daily_works:
        sheet.row_dimensions[current_row].height = 20
        _date_cell(sheet, current_row, 1, work.work_date, border=BORDER_ROW)
        _text_cell(sheet, current_row, 2, work.description, border=BORDER_ROW)
        _number_cell(sheet, current_row, 3, work.quantity, border=BORDER_ROW)
        _text_cell(sheet, current_row, 4, work.unit, align="center", border=BORDER_ROW)
        _number_cell(sheet, current_row, 5, work.actual_person_hours, border=BORDER_ROW)
        total_hours += work.actual_person_hours
        current_row += 1

    # Итог по табелю
    sheet.row_dimensions[current_row].height = 24
    _text_cell(sheet, current_row, 1, "Gesamt Personenstunden", bold=True, border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 2, "", border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 3, "", border=BORDER_TOTAL)
    _text_cell(sheet, current_row, 4, "", border=BORDER_TOTAL)
    _number_cell(
        sheet,
        current_row,
        5,
        total_hours,
        bold=True,
        bg_color=COLOR_GREEN_50,
        color=COLOR_NAVY_900,
        border=BORDER_TOTAL,
    )

    sheet.freeze_panes = "A5"
    for col, width in {"A": 16, "B": 48, "C": 18, "D": 16, "E": 20}.items():
        sheet.column_dimensions[col].width = width
    return sheet


def site_timesheet_workbook(site, *, iso_year=None, iso_week=None):
    """Сформировать отдельный табель MBG одного Standort."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _site_timesheet_sheet(workbook, site, "Stundennachweis", iso_year=iso_year, iso_week=iso_week)
    data = BytesIO()
    workbook.save(data)
    return data.getvalue()


def group_timesheet_workbook(group, *, iso_year=None, iso_week=None):
    """Сформировать табель MBG всей Anlage: сводный лист и листы Standorte."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Übersicht"
    _format_sheet_view(summary, landscape=True)

    summary.merge_cells("A1:F1")
    _text_cell(
        summary,
        1,
        1,
        f"MBG-Stundennachweis Übersicht · {group.name}",
        bold=True,
        color=COLOR_NAVY_900,
        size=14,
        align="left",
    )
    summary.row_dimensions[1].height = 28

    period_str = f"KW {iso_week:02d} / {iso_year}" if (iso_year and iso_week) else "Gesamter Zeitraum"
    client_name = group.client.name if group.client else "—"
    summary.merge_cells("A2:F2")
    _text_cell(
        summary,
        2,
        1,
        f"Auftraggeber: {client_name}  |  Standorte: {group.ac_objects.count()}  |  Zeitraum: {period_str}",
        color=COLOR_GRAY_600,
        size=10,
        align="left",
    )

    summary_headers = [
        "Standort",
        "Projektnummer",
        "Status",
        "Arbeitstage",
        "Einträge",
        "Personenstunden gesamt",
    ]
    _apply_table_header(summary, 4, summary_headers)

    used_sheet_titles = {"Übersicht"}
    total_days = 0
    total_entries = 0
    total_hours = Decimal("0")

    sites = group.ac_objects.all().order_by("name", "pk")
    for number, site in enumerate(sites, start=1):
        row = number + 4
        summary.row_dimensions[row].height = 20
        daily_works = accepted_daily_work(site, iso_year=iso_year, iso_week=iso_week)

        distinct_days = len({w.work_date for w in daily_works})
        entries_count = len(daily_works)
        site_hours = sum((w.actual_person_hours for w in daily_works), Decimal("0"))

        total_days += distinct_days
        total_entries += entries_count
        total_hours += site_hours

        _text_cell(summary, row, 1, site.name, border=BORDER_ROW)
        _text_cell(summary, row, 2, site.project_number or "—", align="center", border=BORDER_ROW)
        _text_cell(summary, row, 3, site.get_status_display(), align="center", border=BORDER_ROW)
        _number_cell(summary, row, 4, distinct_days, decimals=0, border=BORDER_ROW)
        _number_cell(summary, row, 5, entries_count, decimals=0, border=BORDER_ROW)
        _number_cell(summary, row, 6, site_hours, decimals=2, border=BORDER_ROW)

        site_title = _safe_sheet_title(site.name, used_sheet_titles)
        _site_timesheet_sheet(workbook, site, site_title, iso_year=iso_year, iso_week=iso_week)

    # Итог по группе
    total_row = len(sites) + 5
    summary.row_dimensions[total_row].height = 24
    _text_cell(summary, total_row, 1, "Gesamt Anlage", bold=True, border=BORDER_TOTAL)
    _text_cell(summary, total_row, 2, "", border=BORDER_TOTAL)
    _text_cell(summary, total_row, 3, "", border=BORDER_TOTAL)
    _number_cell(summary, total_row, 4, total_days, decimals=0, bold=True, border=BORDER_TOTAL)
    _number_cell(summary, total_row, 5, total_entries, decimals=0, bold=True, border=BORDER_TOTAL)
    _number_cell(
        summary,
        total_row,
        6,
        total_hours,
        decimals=2,
        bold=True,
        bg_color=COLOR_GREEN_50,
        color=COLOR_NAVY_900,
        border=BORDER_TOTAL,
    )

    summary.freeze_panes = "A5"
    for col, width in {"A": 36, "B": 20, "C": 20, "D": 18, "E": 18, "F": 24}.items():
        summary.column_dimensions[col].width = width

    data = BytesIO()
    workbook.save(data)
    return data.getvalue()
