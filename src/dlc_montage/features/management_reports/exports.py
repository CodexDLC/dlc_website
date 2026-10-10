"""Немецкая выгрузка внутренних часов; не использует MBG."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

COLOR_NAVY_900 = "0A2540"
COLOR_NAVY_100 = "E6EAF0"
COLOR_GRAY_900 = "111827"
COLOR_GRAY_100 = "F8FAFC"
COLOR_WHITE = "FFFFFF"
FONT_FAMILY = "Segoe UI"


def build_month_workbook(summary):
    book = Workbook()
    sheet = book.active
    sheet.title = "Stunden"
    sheet.append(
        [
            "Mitarbeiter",
            "Bestätigt (h)",
            "Vorläufig (h)",
            "Stundenlohn brutto (€)",
            "Bestätigtes Brutto (€)",
            "Vorläufiges Brutto (€)",
        ]
    )
    for row in summary["people"]:
        sheet.append([row["name"], row["accepted"] / 60, row["submitted"] / 60, float(row["hourly_rate"])])
        sheet.cell(sheet.max_row, 1).data_type = "s"
        sheet.cell(sheet.max_row, 5, f"=B{sheet.max_row}*D{sheet.max_row}")
        sheet.cell(sheet.max_row, 6, f"=C{sheet.max_row}*D{sheet.max_row}")
    if summary["people"]:
        total_row = sheet.max_row + 1
        sheet.cell(total_row, 1, "Summe")
        for column in (2, 3, 5, 6):
            letter = sheet.cell(1, column).column_letter
            sheet.cell(total_row, column, f"=SUM({letter}2:{letter}{total_row - 1})")
        for cell in sheet[total_row]:
            cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_NAVY_900)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_100)
    details = book.create_sheet("Einzelheiten")
    details.append(["Datum", "Mitarbeiter", "Anlage", "Bereich", "Stunden", "Status"])
    for row in summary["details"]:
        details.append(
            [
                row["date"],
                row["name"],
                row["group"],
                row["direction"].upper(),
                row["minutes"] / 60,
                "Bestätigt" if row["status"] == "accepted" else "Vorläufig",
            ]
        )
        details.cell(details.max_row, 1).number_format = "dd.mm.yyyy"
        for c in (2, 3, 4, 6):
            details.cell(details.max_row, c).data_type = "s"
    for ws in book:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name=FONT_FAMILY, size=10.5, bold=True, color=COLOR_WHITE)
            cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 26
        for row_idx in range(2, ws.max_row + 1):
            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if row_idx % 2 == 0 else COLOR_GRAY_100)
            for cell in ws[row_idx]:
                cell.font = Font(name=FONT_FAMILY, size=10, color=COLOR_GRAY_900)
                cell.fill = row_fill
                cell.border = Border(
                    left=Side(style="thin", color=COLOR_NAVY_100),
                    right=Side(style="thin", color=COLOR_NAVY_100),
                    top=Side(style="thin", color=COLOR_NAVY_100),
                    bottom=Side(style="thin", color=COLOR_NAVY_100),
                )
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 24
            for cell in col:
                if cell.row > 1:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(cell.value, float):
                        cell.number_format = "0.00"
                        cell.alignment = Alignment(horizontal="right", vertical="top")
        if ws.title == "Stunden":
            for row_idx in range(2, ws.max_row + 1):
                for column in (4, 5, 6):
                    ws.cell(row_idx, column).number_format = '#,##0.00 "€"'
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
