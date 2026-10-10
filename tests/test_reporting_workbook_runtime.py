"""Приёмка реальных Excel-генераторов на синтетических снимках без файлового IO."""

import sys
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from features.mbg_reports.models import MbgDailyReport, MbgReportRevision
from features.projects.models import AcObject, ProjectGroup
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))


@pytest.fixture
def accepted_runtime_site(db):
    owner = get_user_model().objects.create_user(username="runtime-owner", is_superuser=True)
    group = ProjectGroup.objects.create(name="Live group", has_ac=True)
    site = AcObject.objects.create(group=group, name="Live site", address="Live address", status="completed")
    for day in (date(2025, 10, 1), date(2026, 9, 30)):
        report = MbgDailyReport.objects.create(
            site=site,
            author=owner,
            work_date=day,
            status="accepted",
            version=1,
        )
        MbgReportRevision.objects.create(
            report=report,
            actor=owner,
            version=1,
            action="accept",
            snapshot={
                "status": "accepted",
                "date": str(day),
                "site_id": site.pk,
                "site_name": "Saved site",
                "address": "Saved address",
                "group_name": "Saved group",
                "project_number": "P-SAVED",
                "workforce": [{"employee_id": 123, "name": "Synthetic worker"}],
                "entries": [
                    {
                        "description": "Runtime cable",
                        "work_type_code": "cable-laying",
                        "calculation_rule": "cable",
                        "unit": "m",
                        "quantity": "1",
                        "completed_quantity": None,
                        "work_item_id": None,
                        "actual_person_hours": "2.00",
                        "start": "08:00:00",
                        "end": "10:00:00",
                        "participants": [{"employee_id": 123, "name": "Synthetic worker"}],
                    },
                    {
                        "description": "Runtime pause",
                        "work_type_code": "mittagspause",
                        "is_break": True,
                        "actual_person_hours": "0.00",
                        "start": "12:00:00",
                        "end": "13:00:00",
                    },
                ],
            },
        )
    return site


@pytest.mark.parametrize(
    "module_name,function_name",
    [
        ("generate_ac_master_report", "main"),
        ("generate_emil_abnahme_report", "generate_abnahme_report"),
    ],
)
def test_real_workbook_paths_preserve_periods_metadata_and_unaccepted_review(
    accepted_runtime_site,
    monkeypatch,
    module_name,
    function_name,
):
    """Генераторы вызываются полностью, но сохранение перехватывается в памяти."""
    import importlib

    module = importlib.import_module(module_name)
    captures = []
    original_save = module.openpyxl.Workbook.save

    def capture(book, _filename):
        buffer = BytesIO()
        original_save(book, buffer)
        captures.append(buffer.getvalue())

    monkeypatch.setattr(module.openpyxl.Workbook, "save", capture)
    monkeypatch.setattr(module.os, "makedirs", lambda *args, **kwargs: None)
    getattr(module, function_name)()
    assert captures
    workbook = load_workbook(BytesIO(captures[0]))
    values = [cell.value for sheet in workbook for row in sheet for cell in row if cell.value is not None]
    text = "\n".join(str(value) for value in values)
    assert "Saved address" in text
    assert "Live address" not in text
    assert "2025" in text and "2026" in text
    assert "Runtime cable" in text
    overview_rows = [
        row for row in workbook.worksheets[0].iter_rows(values_only=True) if len(row) > 1 and row[1] == "Saved site"
    ]
    assert len(overview_rows) == 2
    if module_name == "generate_emil_abnahme_report":
        assert "Freigegeben" not in values
        assert "Restarbeiten Trassenverbindung vor finaler Abnahme fertigstellen" not in values
    else:
        overview_numbers = [
            cell.value
            for row in workbook.worksheets[0]
            for cell in row
            if isinstance(cell.value, (int, float, Decimal))
        ]
        assert 2 in overview_numbers


@pytest.mark.parametrize("duplicate_existing", [True, False])
def test_ambiguous_duplicate_work_never_receives_manual_annotation(duplicate_existing):
    """Одно замечание нельзя переносить между одинаково названными позициями."""
    from sync_emil_google_sheet import merge_or_preserve_site_protocol

    annotated = [1, "Same cable", "m", "1", "Erledigt", "OK", "Manual note"]
    blank = [2, "Same cable", "m", "1", "Erledigt", "—", ""]
    existing = [annotated, blank] if duplicate_existing else [annotated]
    new = [blank] if duplicate_existing else [blank, list(blank)]
    with pytest.raises(RuntimeError, match=r"[Dd]uplicat|[Aa]mbiguous"):
        merge_or_preserve_site_protocol(existing, new)


@pytest.mark.parametrize("quantity", [None, "invalid"])
def test_unknown_flat_quantity_is_not_invented(quantity):
    """Паушальная единица тоже не даёт права придумывать количество в старом снимке."""
    from features.mbg_reports.completion_projection import project_accepted_entries

    rows = project_accepted_entries(
        [
            (
                1,
                1,
                date(2026, 9, 30),
                {
                    "entries": [
                        {
                            "work_type_code": "material-delivery",
                            "unit": "psch.",
                            "quantity": quantity,
                            "actual_person_hours": "2.00",
                            "description": "Delivery",
                        }
                    ]
                },
            )
        ]
    )
    assert rows[0].quantity is None
    assert rows[0].actual_person_hours == Decimal("2.00")


@pytest.mark.parametrize(
    "module_name,function_name",
    [
        ("generate_ac_master_report", "main"),
        ("generate_emil_abnahme_report", "generate_abnahme_report"),
        ("sync_google_sheets", "sync_all"),
        ("sync_emil_google_sheet", "sync_emil_sheet"),
    ],
)
def test_default_reporting_period_does_not_stop_at_sample_weeks(module_name, function_name):
    """Новая неделя не должна исчезать из-за диапазона демонстрационных недель."""
    import ast
    import importlib
    import inspect

    function = getattr(importlib.import_module(module_name), function_name)
    parameter = inspect.signature(function).parameters.get("target_weeks")
    if parameter:
        assert parameter.default is None
    for node in ast.walk(ast.parse(inspect.getsource(function))):
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
            if node.func.id == "load_accepted_reports_dataset":
                assert not any(
                    isinstance(arg.value, (ast.List, ast.Tuple)) for arg in node.keywords if arg.arg == "target_weeks"
                )
