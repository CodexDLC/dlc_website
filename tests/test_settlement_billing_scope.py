"""Billing exports must never reuse the full employee-time scope."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock

import pytest
from django.contrib.auth import get_user_model
from features.mbg_reports.models import MbgDailyReport, MbgReportRevision
from features.mbg_reports.reporting_selectors import load_accepted_reports_dataset
from features.projects.models import AcObject, ProjectGroup
from openpyxl import load_workbook


@pytest.fixture
def billing_scope(db):
    owner = get_user_model().objects.create_user(username="billing-scope-owner")
    group = ProjectGroup.objects.create(name="Billing group", has_ac=True)
    site = AcObject.objects.create(group=group, name="Billing site", status="completed")
    internal_site = AcObject.objects.create(group=group, name="Internal site")
    entries = [
        {"description": "External cable ohne Berechnung text", "actual_person_hours": "4.00"},
        {"description": "Free inverter replacement", "actual_person_hours": "2.50", "internal_only": True},
        {"description": "Legacy lunch", "actual_person_hours": "0.50", "is_break": True},
        {"description": "Legacy pause code", "actual_person_hours": "0.25", "work_type_code": "pause"},
        {"description": "External zero-hour work", "actual_person_hours": "0.00"},
    ]
    for target, day, rows in (
        (site, date(2026, 9, 22), entries),
        (site, date(2026, 9, 23), [entries[1], entries[2]]),
        (internal_site, date(2026, 9, 29), [entries[1], entries[2]]),
    ):
        report = MbgDailyReport.objects.create(site=target, author=owner, work_date=day, status="accepted", version=1)
        MbgReportRevision.objects.create(
            report=report,
            actor=owner,
            version=1,
            action="accept",
            snapshot={
                "status": "accepted",
                "address": "Saved address",
                "project_number": "P-SAVED",
                "site_name": target.name,
                "site_id": target.pk,
                "date": str(day),
                "workforce": [{"name": "Worker"}],
                "entries": [
                    dict(
                        {
                            "work_type_code": "hourly",
                            "unit": "Std.",
                            "quantity": "1",
                            "participants": [{"name": "Worker"}],
                        },
                        **row,
                    )
                    for row in rows
                ],
            },
        )
    return site, internal_site


@pytest.mark.parametrize("inject_dataset", [True, False])
def test_actual_master_workbook_scopes_even_injected_employee_dataset(billing_scope, inject_dataset):
    from features.mbg_reports.master_workbook import build_master_settlement_workbook

    site, internal_site = billing_scope
    full = load_accepted_reports_dataset()
    internal = full.all_sites[site.pk]
    assert internal.total_hours == Decimal("9.00")
    assert len(internal.day_facts) == 2
    assert any(entry.is_break for entry in internal.day_facts[0].entries)
    book = build_master_settlement_workbook(dataset=full if inject_dataset else None)
    buffer = BytesIO()
    book.save(buffer)
    exported = load_workbook(BytesIO(buffer.getvalue()))
    text = "\n".join(str(c.value) for s in exported for row in s for c in row if c.value is not None)
    for forbidden in ("Free inverter replacement", "Legacy lunch", "Legacy pause code", internal_site.name):
        assert forbidden not in text
    assert "External cable ohne Berechnung text" in text
    assert "External zero-hour work" in text
    summary_row = next(row for row in exported["Übersicht"].values if row[1] == site.name)
    assert summary_row[5] == 1
    assert summary_row[7] == 4
    # Projection must not mutate the shared internal dataset or accepted history.
    assert full.all_sites[site.pk] == internal
    assert (
        MbgReportRevision.objects.get(report_id=internal.day_facts[0].report_id).snapshot["entries"][2][
            "actual_person_hours"
        ]
        == "0.50"
    )


@pytest.mark.parametrize("adapter", ["feature", "script"])
def test_google_billing_payload_uses_same_scope_and_preserves_payment(billing_scope, monkeypatch, adapter):
    site, internal_site = billing_scope
    sheet = MagicMock()
    sheet.title = "Fake billing spreadsheet"
    summary = MagicMock()
    summary.title = "Übersicht"
    summary.id = 1
    summary.get_all_values.return_value = [
        ["KW 39 (2026)"],
        ["", site.name, "", "", "", "", "", "", "", "Bezahlt"],
    ]
    state = summary.get_all_values.return_value + [[] for _ in range(30)]
    state[-1] = ["", "Stale private billed row", "", "", "", "", "", "25.00", "", "Offen"]
    summary.get_all_values.side_effect = lambda: [list(row) for row in state]

    def apply_update(*args, **kwargs):
        written = args[0] if args else kwargs["values"]
        state[: len(written)] = [list(row) for row in written]

    summary.update.side_effect = apply_update
    sheet.worksheet.return_value = summary
    sheet.worksheets.return_value = [summary]
    if adapter == "feature":
        from features.mbg_reports.google_sync import sync_master_to_google_sheets

        sync_master_to_google_sheets(spreadsheet=sheet)
    else:
        from pathlib import Path

        monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / "scripts"))
        import sync_google_sheets

        monkeypatch.setattr(sync_google_sheets, "get_spreadsheet", lambda: sheet)
        sync_google_sheets.sync_all(target_weeks=[39])
    call = summary.update.call_args
    rows = call.args[0] if call.args else call.kwargs["values"]
    site_row = next(row for row in rows if len(row) > 1 and row[1] == site.name)
    assert float(str(site_row[7]).replace(",", ".")) == 4
    assert site_row[5] == 1
    assert site_row[9] == "Bezahlt"
    assert not any("Stale private billed row" in row for row in state)
    assert not any(len(row) > 1 and row[1] == internal_site.name for row in rows)


def test_internal_weekly_workbook_keeps_full_rows_and_original_accounting(billing_scope):
    from features.mbg_reports.export_workbooks import build_weekly_workbook

    site, _ = billing_scope
    snapshot = MbgReportRevision.objects.filter(report__site=site).order_by("report__work_date").first().snapshot
    book = load_workbook(BytesIO(build_weekly_workbook({"grouped": {site.pk: [snapshot]}, "year": 2026, "week": 39})))
    sheet = book.worksheets[1]
    records = {row[6]: row[5] for row in sheet.values if len(row) > 6}
    assert records["Free inverter replacement"] == 2.5
    assert records["Legacy lunch"] == 0.5
    assert records["Legacy pause code"] == 0.25


def test_settlement_projection_removes_internal_only_periods_without_mutating_facts(billing_scope):
    from features.mbg_reports.settlement_scope import project_settlement_dataset

    site, internal_site = billing_scope
    full = load_accepted_reports_dataset()
    projected = project_settlement_dataset(full)
    assert full.periods == ((2026, 39), (2026, 40))
    assert projected.periods == ((2026, 39),)
    assert set(projected.all_sites) == {site.pk}
    assert internal_site.pk in full.all_sites
    assert projected.all_sites[site.pk].total_hours == Decimal("4.00")
    assert projected == project_settlement_dataset(projected)


def test_final_hourly_renderer_excludes_private_rows_without_changing_weekly_source(billing_scope):
    from features.mbg_reports.export_workbooks import build_final_hourly_workbook

    site, _ = billing_scope
    snapshot = MbgReportRevision.objects.filter(report__site=site).order_by("report__work_date").first().snapshot
    book = load_workbook(BytesIO(build_final_hourly_workbook([snapshot], date(2026, 9, 22), date(2026, 9, 22))))
    values = [str(c.value) for s in book for row in s for c in row if c.value is not None]
    text = "\n".join(values)
    assert "Free inverter replacement" not in text
    assert "Legacy lunch" not in text
    assert "Legacy pause code" not in text
    assert "External zero-hour work" in text
    assert len(snapshot["entries"]) == 5


def test_final_hourly_artifact_snapshot_is_billing_scoped(billing_scope, monkeypatch):
    from features.mbg_reports import export_services

    site, internal_site = billing_scope
    internal_site.status = "completed"
    internal_site.save(update_fields=["status"])
    owner = get_user_model().objects.get(username="billing-scope-owner")
    owner.is_superuser = True
    owner.save(update_fields=["is_superuser"])
    for revision in MbgReportRevision.objects.all():
        revision.snapshot["address"] = "Saved address"
        revision.snapshot["project_number"] = "P-SAVED"
        revision.save(update_fields=["snapshot"])
    monkeypatch.setattr(export_services, "store_export_artifact", lambda artifact, payload: artifact)
    artifact = export_services.generate_final_hours(actor=owner, group_id=site.group_id)
    names = [entry["description"] for snapshot in artifact.snapshot["reports"] for entry in snapshot["entries"]]
    assert names == ["External cable ohne Berechnung text", "External zero-hour work"]
    assert len(artifact.snapshot["reports"]) == 1

    assert artifact.period_start == artifact.period_end == date(2026, 9, 22)
