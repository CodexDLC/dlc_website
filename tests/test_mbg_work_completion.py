"""Контракт клиентского нефинансового Leistungsnachweis MBG."""

from datetime import date, time
from decimal import Decimal
from uuid import uuid4

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.storage import FileSystemStorage
from django.urls import reverse
from features.mbg_reports.models import MbgReportRevision, PerformanceExport
from features.mbg_reports.services import review_report, save_report
from features.projects.models import AcObject, ProjectGroup
from features.projects.services import save_site_work_item
from features.staff.models import Employee
from features.work_catalog.models import CableRoute, WorkType
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db


@pytest.fixture
def completion_data():
    admin = get_user_model().objects.create_user(username="completion-owner", is_superuser=True)
    crew = [
        Employee.objects.create(first_name="Ada", last_name="Montage", has_ac_access=True),
        Employee.objects.create(first_name="Ben", last_name="Montage", has_ac_access=True),
    ]
    group = ProjectGroup.objects.create(name="Anlage Nord")
    site = AcObject.objects.create(
        group=group,
        name="Haus 10",
        address="Nordallee 10",
        project_number="PN-10",
        status="started",
    )
    return admin, crew, group, site


def _accepted_report(*, actor, site, workforce, work_date, entries):
    report = save_report(
        actor=actor,
        site=site,
        work_date=work_date,
        workforce=workforce,
        entries=entries,
        submit=True,
    )
    return review_report(actor=actor, report_id=report.pk, version=report.version, decision="accept")


def _entry(*, work, people, start=time(8), end=time(9), quantity=1, route=None, cable_type=None, description=""):
    return {
        "participants": people,
        "starts_at": start,
        "ends_at": end,
        "work_type": work,
        "quantity": Decimal(str(quantity)),
        "route": route,
        "cable_type": cable_type,
        "standard_connections": 0,
        "heavy_connections": 0,
        "with_underconstruction": False,
        "description": description,
    }


def _all_workbook_text(book):
    return "\n".join(
        str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value
    )


def _performance_export(*, actor, group, document_type, filename, snapshot):
    return PerformanceExport.objects.create(
        group=group,
        document_type=document_type,
        fingerprint=uuid4().hex,
        snapshot=snapshot,
        filename=filename,
        author=actor,
    )


def test_work_completion_export_is_snapshot_only_unpriced_and_idempotent(completion_data, monkeypatch, tmp_path):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="wr-montage")
    _accepted_report(
        actor=admin,
        site=site,
        workforce=crew,
        work_date=date(2026, 8, 28),
        entries=[_entry(work=work, people=crew, quantity=2, description="=Montage Wechselrichter")],
    )

    artifact = generate_work_completion_report(actor=admin, group_id=group.pk)
    assert generate_work_completion_report(actor=admin, group_id=group.pk).pk == artifact.pk
    assert artifact.document_type == "work_completion"
    assert artifact.snapshot["sites"][0]["rows"][0]["actual_person_hours"] == "2.00"
    assert "unit_price" not in str(artifact.snapshot)
    assert "normative" not in str(artifact.snapshot)
    with artifact.file.open("rb") as source:
        book = load_workbook(source, data_only=False)
    assert book.sheetnames[0] == "Übersicht"
    assert book.worksheets[1]["A1"].value == "Leistungsnachweis"
    assert book.worksheets[1]["B3"].value == site.name
    assert book.worksheets[1]["A8"].value == "Leistung"
    assert book.worksheets[1]["D8"].value == "Ist-Personenstunden"
    assert book.worksheets[1]["A9"].data_type == "s"  # Excel formula injection remains text.
    assert "€" not in _all_workbook_text(book)
    assert "Einzelpreis" not in _all_workbook_text(book)
    assert "Norm" not in _all_workbook_text(book)


def test_work_completion_uses_item_ids_and_deltas_for_distinct_and_continued_work(
    completion_data, monkeypatch, tmp_path
):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="cable-tray")
    route = CableRoute.objects.create(code="completion-same-route", name_de="NH → HV")
    wr1 = save_site_work_item(
        actor=admin, site=site, work_type=work, label="WR1", planned_quantity=1, route=route
    )
    wr2 = save_site_work_item(
        actor=admin, site=site, work_type=work, label="WR2", planned_quantity=1, route=route
    )
    _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 26),
        entries=[
            {**_entry(work=work, people=crew[:1], route=route, description="WR1 Trasse"),
             "work_item": wr1, "completed_quantity": Decimal("0")},
            {**_entry(work=work, people=crew[:1], start=time(9), end=time(10), route=route,
                      description="WR2 Trasse"), "work_item": wr2, "completed_quantity": Decimal("1")},
        ],
    )
    _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 27),
        entries=[{**_entry(work=work, people=crew[:1], route=route, description="WR1 Trasse"),
                  "work_item": wr1, "completed_quantity": Decimal("1")}],
    )
    work.name_de = "Каталог теперь изменён"
    work.save(update_fields=["name_de"])
    route.name_de = "Новый путь"
    route.save(update_fields=["name_de"])

    artifact = generate_work_completion_report(actor=admin, group_id=group.pk)
    rows = artifact.snapshot["sites"][0]["rows"]
    assert {row["work_item_id"]: row["quantity"] for row in rows} == {wr1.pk: "1.00", wr2.pk: "1.00"}
    assert {row["work_item_id"]: row["actual_person_hours"] for row in rows} == {
        wr1.pk: "2.00", wr2.pk: "1.00"
    }
    assert {row["name"] for row in rows} == {"WR1 Trasse · NH → HV", "WR2 Trasse · NH → HV"}
    with artifact.file.open("rb") as source:
        book = load_workbook(source)
    assert [book.worksheets[1].cell(line, 3).value for line in (9, 10)] == [1, 1]


def test_work_completion_unit_mismatch_stays_unresolved_for_linked_item(completion_data):
    from features.mbg_reports.selectors import work_completion_projection
    from features.mbg_reports.work_items import accepted_work_item_progress

    admin, crew, _group, site = completion_data
    work = WorkType.objects.get(code="cable-tray")
    route = CableRoute.objects.create(code="completion-mixed-unit", name_de="NH → HV")
    item = save_site_work_item(
        actor=admin, site=site, work_type=work, label="WR1", planned_quantity=2, route=route
    )
    first = _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 26),
        entries=[{**_entry(work=work, people=crew[:1], route=route, description="WR1"),
                  "work_item": item, "completed_quantity": Decimal("1")}],
    )
    second = _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 27),
        entries=[{**_entry(work=work, people=crew[:1], route=route, description="WR1"),
                  "work_item": item, "completed_quantity": Decimal("1")}],
    )
    revision = second.revisions.get(version=second.version)
    snapshot = revision.snapshot
    snapshot["entries"][0]["unit"] = "Stk."
    revision.snapshot = snapshot
    revision.save(update_fields=["snapshot"])

    projection = work_completion_projection(site_id=site.pk)
    rows = projection["sites"][0]["rows"]
    assert projection["accepted_report_count"] == 2
    assert len(rows) == 2
    assert {row["kind"] for row in rows} == {"unresolved_unit_mismatch"}
    assert all(row["work_item_id"] == item.pk and row["quantity"] is None for row in rows)
    assert {row["sources"][0]["report_id"] for row in rows} == {first.pk, second.pk}
    progress = accepted_work_item_progress(site_id=site.pk)
    assert progress.items == ()
    assert len(progress.unresolved_entries) == 2
    assert {row.reason for row in progress.unresolved_entries} == {"unit_mismatch"}


def test_work_completion_linked_correction_replaces_current_delta_and_preserves_export_history(
    completion_data, monkeypatch, tmp_path
):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="cable-tray")
    route = CableRoute.objects.create(code="completion-revision-route", name_de="NH → HV")
    item = save_site_work_item(
        actor=admin, site=site, work_type=work, label="WR1", planned_quantity=2, route=route
    )
    entry = {**_entry(work=work, people=crew[:1], route=route, description="WR1"),
             "work_item": item, "completed_quantity": Decimal("1")}
    report = _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 26), entries=[entry]
    )
    first = generate_work_completion_report(actor=admin, group_id=group.pk)
    assert generate_work_completion_report(actor=admin, group_id=group.pk).pk == first.pk

    corrected = save_report(
        actor=admin, report_id=report.pk, version=report.version, site=site,
        work_date=report.work_date, workforce=crew[:1],
        entries=[{**entry, "completed_quantity": Decimal("0")}], reason="Korrektur der Menge",
    )
    assert corrected.status == "accepted"
    second = generate_work_completion_report(actor=admin, group_id=group.pk)
    assert second.pk != first.pk
    assert first.snapshot["sites"][0]["rows"][0]["quantity"] == "1.00"
    assert second.snapshot["sites"][0]["rows"][0]["quantity"] == "0.00"
    assert second.snapshot["sites"][0]["rows"][0]["actual_person_hours"] == "1.00"
    assert generate_work_completion_report(actor=admin, group_id=group.pk).pk == second.pk
    assert PerformanceExport.objects.filter(group=group, document_type="work_completion").count() == 2


def test_work_completion_hourly_hours_and_unknown_legacy_quantity_are_explicit(
    completion_data, monkeypatch, tmp_path
):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    hourly = WorkType.objects.create(
        code="completion-delivery", name_de="Lieferung", direction="ac", unit="hour",
        calculation_rule="actual", price_mode="actual_hourly",
    )
    tray = WorkType.objects.get(code="cable-tray")
    route = CableRoute.objects.create(code="completion-legacy-route", name_de="Alte Trasse")
    first = _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 28),
        entries=[
            _entry(work=hourly, people=crew[:1], end=time(10, 30), quantity=2, description="Lieferung"),
            _entry(work=tray, people=crew[:1], start=time(11), end=time(12), quantity=3,
                   route=route, description="Alte Trasse"),
        ],
    )
    second = _accepted_report(
        actor=admin, site=site, workforce=crew[:1], work_date=date(2026, 8, 29),
        entries=[_entry(work=hourly, people=crew[:1], end=time(12), quantity=9, description="Lieferung")],
    )
    legacy_revision = second.revisions.get(version=second.version)
    legacy_snapshot = legacy_revision.snapshot
    legacy_snapshot["entries"][0]["unit"] = "Stk."
    legacy_revision.snapshot = legacy_snapshot
    legacy_revision.save(update_fields=["snapshot"])

    artifact = generate_work_completion_report(actor=admin, group_id=group.pk)
    rows = artifact.snapshot["sites"][0]["rows"]
    delivery = next(row for row in rows if row["name"] == "Lieferung")
    legacy = next(row for row in rows if "Alte Trasse" in row["name"])
    assert delivery["quantity"] == "6.50"
    assert delivery["unit"] == "Std."
    assert delivery["actual_person_hours"] == "6.50"
    assert {source["report_version"] for source in delivery["sources"]} == {first.version, second.version}
    assert [source["actual_person_hours"] for source in delivery["sources"]] == ["2.50", "4.00"]
    assert delivery["sources"][1]["unit"] == "Stk."
    assert legacy["quantity"] is None
    assert legacy["sources"][0]["recorded_quantity"] == "3.00"
    assert legacy["sources"][0]["work_date"] == "2026-08-28"
    with artifact.file.open("rb") as source:
        book = load_workbook(source)
    sheet = book.worksheets[1]
    assert "Menge ungeklärt" in _all_workbook_text(book)
    assert any(sheet.cell(line, 3).value == "Menge ungeklärt" for line in (9, 10))
    assert any(sheet.cell(line, 3).value == 6.5 for line in (9, 10))
    assert "€" not in _all_workbook_text(book)


def test_work_completion_uses_only_current_accepted_snapshot_and_excludes_breaks(
    completion_data, monkeypatch, tmp_path
):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="wr-montage")
    lunch = WorkType.objects.get(code="mittagspause")
    report = _accepted_report(
        actor=admin,
        site=site,
        workforce=crew,
        work_date=date(2026, 8, 29),
        entries=[
            _entry(work=work, people=crew, start=time(8), end=time(9), description="Montage"),
            _entry(work=lunch, people=crew, start=time(12), end=time(13), description="Mittagspause"),
        ],
    )
    accepted_revision = report.revisions.get(version=report.version)
    accepted_snapshot = accepted_revision.snapshot
    accepted_snapshot["entries"][0]["actual_person_hours"] = "99.00"
    accepted_revision.snapshot = accepted_snapshot
    accepted_revision.save(update_fields=["snapshot"])
    report.version += 1
    report.save(update_fields=["version"])

    with pytest.raises(ValidationError, match="текущих принятых снимков"):
        generate_work_completion_report(actor=admin, group_id=group.pk)

    current = MbgReportRevision.objects.create(
        report=report,
        actor=admin,
        version=report.version,
        action="save",
        snapshot={**accepted_snapshot, "status": "accepted"},
    )
    artifact = generate_work_completion_report(actor=admin, group_id=group.pk)
    row = artifact.snapshot["sites"][0]["rows"][0]
    assert row["actual_person_hours"] == "99.00"
    assert "Mittagspause" not in str(artifact.snapshot["sites"])
    assert current.version == report.version


def test_work_completion_preserves_legacy_snapshot_rows_and_distinct_descriptions(completion_data):
    from features.mbg_reports.selectors import work_completion_projection

    admin, crew, _group, site = completion_data
    work = WorkType.objects.get(code="wr-montage")
    old_report = _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 8, 30),
        entries=[_entry(work=work, people=crew[:1], quantity=1, description="Montage")],
    )
    old_snapshot = old_report.revisions.get(version=old_report.version).snapshot
    for field in ("with_underconstruction", "standard_connections", "heavy_connections"):
        old_snapshot["entries"][0].pop(field)
    old_report.revisions.filter(version=old_report.version).update(snapshot=old_snapshot)

    current_report = _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 8, 31),
        entries=[
            _entry(work=work, people=crew[:1], quantity=1, description="Montage"),
            {
                **_entry(
                    work=work,
                    people=crew[:1],
                    start=time(9),
                    end=time(10),
                    quantity=1,
                    description="Montage",
                ),
                "with_underconstruction": True,
            },
        ],
    )

    rows = work_completion_projection(site_id=site.pk)["sites"][0]["rows"]
    assert len(rows) == 3
    assert all(row["quantity"] is None for row in rows)
    assert len([row for row in rows if row["name"] == "Montage"]) == 2
    assert "mit Unterkonstruktion" in next(row for row in rows if "Unterkonstruktion" in row["name"])["name"]
    assert current_report.status == "accepted"


def test_work_completion_keeps_legacy_routes_as_unknown_rows_and_group_isolation(
    completion_data, monkeypatch, tmp_path
):
    from features.mbg_reports.exports import generate_work_completion_report

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="cable-tray")
    route = CableRoute.objects.create(code="nord-route", name_de="Nordtrasse")
    other_route = CableRoute.objects.create(code="sued-route", name_de="Südtrasse")
    _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 8, 1),
        entries=[_entry(work=work, people=crew[:1], quantity=1, route=route, description="Trasse Nord")],
    )
    _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 9, 1),
        entries=[_entry(work=work, people=crew[:1], quantity=3, route=route, description="Trasse Nord")],
    )
    _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 9, 2),
        entries=[_entry(work=work, people=crew[:1], quantity=2, route=other_route, description="Trasse Süd")],
    )
    other = AcObject.objects.create(
        group=group, name="Haus 11", address="Nordallee 11", project_number="PN-11", status="started"
    )
    _accepted_report(
        actor=admin,
        site=other,
        workforce=crew[:1],
        work_date=date(2026, 9, 3),
        entries=[_entry(work=work, people=crew[:1], quantity=8, route=route, description="Andere Trasse")],
    )
    other.status = "completed"
    other.save(update_fields=["status"])

    group_artifact = generate_work_completion_report(actor=admin, group_id=group.pk)
    site_artifact = generate_work_completion_report(actor=admin, site_id=site.pk)
    group_rows = group_artifact.snapshot["sites"][0]["rows"]
    north = [row for row in group_rows if row["route"] == "Nordtrasse"]
    assert len(north) == 2
    assert all(row["quantity"] is None for row in north)
    assert sum(Decimal(row["actual_person_hours"]) for row in north) == Decimal("2")
    assert {item["site_id"] for item in site_artifact.snapshot["sites"]} == {site.pk}


def test_work_completion_validates_scope_and_export_endpoint_is_post_only(completion_data, client):
    from features.mbg_reports.exports import generate_work_completion_report

    admin, _crew, group, site = completion_data
    with pytest.raises(ValidationError):
        generate_work_completion_report(actor=admin, site_id=site.pk)
    stranger = get_user_model().objects.create_user(username="completion-stranger")
    with pytest.raises(PermissionDenied):
        generate_work_completion_report(actor=stranger, group_id=group.pk)
    client.force_login(admin)
    assert client.get(reverse("mbg_reports:work_completion_export")).status_code == 405
    assert client.post(reverse("mbg_reports:work_completion_export"), {"scope": "invalid"}).status_code == 302


def test_work_completion_history_is_bounded_and_keeps_priced_exports_separate(completion_data, client):
    from features.mbg_reports.selectors import work_completion_export_history

    admin, _crew, group, site = completion_data
    priced = _performance_export(
        actor=admin,
        group=group,
        document_type="priced_performance",
        filename="priced-performance.xlsx",
        snapshot={"document_type": "leistungsnachweis"},
    )
    for index in range(11):
        _performance_export(
            actor=admin,
            group=group,
            document_type="work_completion",
            filename=f"work-completion-{index}.xlsx",
            snapshot={
                "document_type": "work_completion",
                "scope": "site" if index == 10 else "group",
                "group_name": group.name,
                "sites": [{"site_name": site.name}],
            },
        )

    client.force_login(admin)
    history = work_completion_export_history()
    assert len(history) == 10
    assert history[0]["artifact"].filename == "work-completion-10.xlsx"
    assert history[-1]["artifact"].filename == "work-completion-1.xlsx"
    assert history[0]["scope_label"] == site.name

    hub = client.get(reverse("mbg_reports:export_hub"))
    content = hub.content.decode()
    assert hub.status_code == 200
    assert "Последние отчёты о выполненных работах" in content
    assert all(
        reverse("mbg_reports:performance_download", args=[item["artifact"].pk]) in content for item in history
    )
    assert reverse("mbg_reports:performance_download", args=[priced.pk]) not in content

    performance_page = client.get(reverse("mbg_reports:performance_exports", args=[group.pk]))
    performance_content = performance_page.content.decode()
    assert priced.filename in performance_content
    assert "work-completion-0.xlsx" not in performance_content


def test_work_completion_hub_hides_inactive_scopes_and_requires_export_permission(completion_data, client):
    from features.mbg_reports.selectors import work_completion_groups

    admin, crew, group, site = completion_data
    work = WorkType.objects.get(code="wr-montage")
    _accepted_report(
        actor=admin,
        site=site,
        workforce=crew[:1],
        work_date=date(2026, 9, 5),
        entries=[_entry(work=work, people=crew[:1], description="Montage")],
    )
    assert work_completion_groups() == [{"group": group, "sites": [site]}]

    site.is_active = False
    site.save(update_fields=["is_active"])
    assert work_completion_groups() == []

    site.is_active = True
    site.save(update_fields=["is_active"])
    group.status = "archived"
    group.save(update_fields=["status"])
    assert work_completion_groups() == []

    stranger = get_user_model().objects.create_user(username="completion-history-stranger")
    artifact = _performance_export(
        actor=admin,
        group=group,
        document_type="work_completion",
        filename="private-work-completion.xlsx",
        snapshot={"document_type": "work_completion", "scope": "group", "group_name": group.name},
    )
    client.force_login(stranger)
    response = client.post(reverse("mbg_reports:work_completion_export"), {"scope": "group", "scope_id": group.pk})
    assert response.status_code == 403
    assert client.get(reverse("mbg_reports:performance_download", args=[artifact.pk])).status_code == 403
