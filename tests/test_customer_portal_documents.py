"""Тесты раздела «Документы» и клиентских экспортов XLSX."""

from datetime import date
from io import BytesIO
from uuid import uuid4

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from features.customer_portal.models import SiteReview
from features.customer_portal.services import grant_membership
from features.mbg_reports.models import MbgDailyReport, MbgReportRevision
from features.projects.models import AcObject, Client, ProjectGroup
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db


@pytest.fixture
def portal_setup():
    users = get_user_model()
    admin = users.objects.create_superuser(username="portal-doc-admin", email="admin@doc.test", password=None)
    alice = users.objects.create_user(username="portal-alice")
    bob = users.objects.create_user(username="portal-bob")
    stranger = users.objects.create_user(username="portal-stranger")

    client_a = Client.objects.create(name="Auftraggeber Alpha")
    client_b = Client.objects.create(name="Auftraggeber Beta")

    grant_membership(actor=admin, user=alice, client=client_a, can_manage_projects=True)
    grant_membership(actor=admin, user=bob, client=client_b, can_manage_projects=False)
    # stranger has no ClientMembership

    group_a = ProjectGroup.objects.create(name="Anlage Alpha", client=client_a)
    group_b = ProjectGroup.objects.create(name="Anlage Beta", client=client_b)

    site_a1 = AcObject.objects.create(group=group_a, name="Standort Alpha 1", project_number="P-A1")
    long_name = "Standort Alpha 2 mit [sehr langem] / verbotenem Namen: 12345678901234567890"
    site_a2 = AcObject.objects.create(group=group_a, name=long_name, project_number="P-A2")
    site_b = AcObject.objects.create(group=group_b, name="Standort Beta 1", project_number="P-B1")

    return {
        "admin": admin,
        "alice": alice,
        "bob": bob,
        "stranger": stranger,
        "client_a": client_a,
        "client_b": client_b,
        "group_a": group_a,
        "group_b": group_b,
        "site_a1": site_a1,
        "site_a2": site_a2,
        "site_b": site_b,
    }


def _create_report_with_snapshot(site, author, work_date, entries, status="accepted", version=1):
    report = MbgDailyReport.objects.create(
        author=author,
        site=site,
        work_date=work_date,
        status=status,
        version=version,
        submission_key=uuid4(),
    )
    revision = MbgReportRevision.objects.create(
        report=report,
        actor=author,
        version=version,
        action="accept",
        snapshot={
            "status": status,
            "entries": entries,
        },
    )
    return report, revision


def test_documents_view_and_access_scope(portal_setup, client):
    ctx = portal_setup
    client.force_login(ctx["alice"])
    res = client.get(reverse("customer_portal:documents"))
    assert res.status_code == 200
    content = res.content.decode()
    assert "Anlage Alpha" in content
    assert "Standort Alpha 1" in content
    assert "Anlage Beta" not in content

    # Stranger without membership is denied
    client.force_login(ctx["stranger"])
    assert client.get(reverse("customer_portal:documents")).status_code == 403
    assert client.get(reverse("customer_portal:group_work_xlsx", args=[ctx["group_a"].pk])).status_code == 403
    assert client.get(reverse("customer_portal:accepted_work_xlsx", args=[ctx["site_a1"].pk])).status_code == 403
    assert client.get(reverse("customer_portal:group_timesheet_xlsx", args=[ctx["group_a"].pk])).status_code == 403
    assert client.get(reverse("customer_portal:site_timesheet_xlsx", args=[ctx["site_a1"].pk])).status_code == 403

    # Bob (Client B) cannot access Alice's (Client A) exports
    client.force_login(ctx["bob"])
    assert client.get(reverse("customer_portal:group_work_xlsx", args=[ctx["group_a"].pk])).status_code == 404
    assert client.get(reverse("customer_portal:accepted_work_xlsx", args=[ctx["site_a1"].pk])).status_code == 404
    assert client.get(reverse("customer_portal:group_timesheet_xlsx", args=[ctx["group_a"].pk])).status_code == 404
    assert client.get(reverse("customer_portal:site_timesheet_xlsx", args=[ctx["site_a1"].pk])).status_code == 404

    # Admin can access any
    client.force_login(ctx["admin"])
    res_admin = client.get(reverse("customer_portal:group_work_xlsx", args=[ctx["group_a"].pk]))
    assert res_admin.status_code == 200
    assert res_admin["Cache-Control"] == "private, no-store"


def test_multi_person_row_and_actual_hours_sum(portal_setup, client):
    ctx = portal_setup
    # Multi-person row: 3 workers, 2.5 hours each => 7.5 actual person hours
    entries = [
        {
            "description": "Kabeltrasse montieren",
            "unit": "m",
            "completed_quantity": "25",
            "actual_person_hours": "7.50",
            "normative_person_hours": "4.00",
            "participants": [{"name": "Worker 1"}, {"name": "Worker 2"}, {"name": "Worker 3"}],
        },
        {
            "description": "Unterkonstruktion",
            "unit": "Stk.",
            "completed_quantity": "5",
            "actual_person_hours": "3.50",
            "normative_person_hours": "2.00",
            "participants": [{"name": "Worker 1"}],
        },
    ]
    _create_report_with_snapshot(ctx["site_a1"], ctx["admin"], date(2026, 10, 1), entries)

    client.force_login(ctx["alice"])
    # Download timesheet
    res = client.get(reverse("customer_portal:site_timesheet_xlsx", args=[ctx["site_a1"].pk]))
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    sheet = wb.active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None]

    # Must contain 7.5 and 3.5, and total sum 11.0
    assert 7.5 in values or 7.50 in values
    assert 3.5 in values or 3.50 in values
    assert 11.0 in values or 11 in values

    # Must NOT expose worker names
    str_values = " ".join(str(v) for v in values)
    assert "Worker 1" not in str_values
    assert "Worker 2" not in str_values


def test_hourly_and_unresolved_quantities_not_guessed(portal_setup, client):
    ctx = portal_setup
    # Hourly row has unit Std., actual hours 4.0, but no physical volume
    # Unresolved piece row has completed_quantity=None
    entries = [
        {
            "description": "Fehlersuche Wechselrichter",
            "unit": "Std.",
            "calculation_rule": "actual",
            "quantity": "1",  # dummy in form, should not be mistaken for physical volume
            "completed_quantity": None,
            "actual_person_hours": "4.00",
        },
        {
            "description": "Kabelverlegung ohne Mengenerfassung",
            "unit": "m",
            "calculation_rule": "cable",
            "completed_quantity": None,
            "actual_person_hours": "2.00",
        },
    ]
    _create_report_with_snapshot(ctx["site_a1"], ctx["admin"], date(2026, 10, 2), entries)

    client.force_login(ctx["alice"])
    res = client.get(reverse("customer_portal:accepted_work_xlsx", args=[ctx["site_a1"].pk]))
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    sheet = wb["Arbeiten"] if "Arbeiten" in wb.sheetnames else wb.active
    rows = list(sheet.iter_rows(values_only=True))

    # Find the row for 'Kabelverlegung ohne Mengenerfassung'
    unresolved_row = next(r for r in rows if r and "Kabelverlegung ohne Mengenerfassung" in str(r[0]))
    # Quantity must be 'ungeklärt' / 'не подтверждено', NOT 1 or 0!
    qty_val = unresolved_row[1]
    assert str(qty_val).lower() in ("ungeklärt", "nicht bestätigt", "не подтверждено")
    assert qty_val not in (0, 1, "0", "1")


def test_stale_accepted_snapshot_not_masked(portal_setup, client):
    ctx = portal_setup
    report, _rev = _create_report_with_snapshot(
        ctx["site_a1"], ctx["admin"], date(2026, 10, 2),
        [{"description": "DC Montage", "unit": "Stk.", "completed_quantity": "10", "actual_person_hours": "5.00"}],
    )
    # Create a review that accepted earlier digest, then report updated to version 2 without revision
    SiteReview.objects.create(
        site=ctx["site_a1"],
        status="accepted",
        snapshot_digest="old_stale_digest_12345",
    )
    report.version = 2
    report.save()

    client.force_login(ctx["alice"])
    res = client.get(reverse("customer_portal:accepted_work_xlsx", args=[ctx["site_a1"].pk]))
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    sheet = wb.active
    all_texts = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    # Warning about stale or incomplete reports must appear
    assert "Unvollständig" in all_texts or "Veraltet" in all_texts or "Freigabestand" in all_texts


def test_iso_year_boundary_and_calendar_week_filter(portal_setup, client):
    ctx = portal_setup
    # 2020-12-31 is Thursday in ISO year 2020, week 53
    _create_report_with_snapshot(
        ctx["site_a1"],
        ctx["admin"],
        date(2020, 12, 31),
        [
            {
                "description": "Silvester Arbeit",
                "unit": "Std.",
                "calculation_rule": "actual",
                "actual_person_hours": "3.00",
            }
        ],
    )
    # 2021-01-04 is Monday in ISO year 2021, week 1
    _create_report_with_snapshot(
        ctx["site_a1"],
        ctx["admin"],
        date(2021, 1, 4),
        [
            {
                "description": "Neujahr Arbeit",
                "unit": "Std.",
                "calculation_rule": "actual",
                "actual_person_hours": "4.00",
            }
        ],
    )

    client.force_login(ctx["alice"])
    # Filter for KW 53 / 2020
    ts_url = reverse("customer_portal:site_timesheet_xlsx", args=[ctx["site_a1"].pk])
    res_w53 = client.get(ts_url + "?year=2020&week=53")
    assert res_w53.status_code == 200
    wb53 = load_workbook(BytesIO(res_w53.content))
    txt53 = " ".join(str(cell.value) for row in wb53.active.iter_rows() for cell in row if cell.value)
    assert "Silvester Arbeit" in txt53
    assert "Neujahr Arbeit" not in txt53

    # Filter for KW 1 / 2021
    res_w1 = client.get(ts_url + "?year=2021&week=1")
    assert res_w1.status_code == 200
    wb1 = load_workbook(BytesIO(res_w1.content))
    txt1 = " ".join(str(cell.value) for row in wb1.active.iter_rows() for cell in row if cell.value)
    assert "Neujahr Arbeit" in txt1
    assert "Silvester Arbeit" not in txt1

    # Invalid week query returns 400
    res_invalid = client.get(ts_url + "?year=2020&week=99")
    assert res_invalid.status_code == 400


def test_typed_cells_formula_defense_and_safe_sheet_names(portal_setup, client):
    ctx = portal_setup
    _create_report_with_snapshot(
        ctx["site_a2"], ctx["admin"], date(2026, 10, 3),
        [
            {
                "description": "=SUM(A1:A10)",
                "unit": "Stk.",
                "completed_quantity": "15",
                "actual_person_hours": "6.50",
            },
            {
                "description": "+cmd|' /C calc'!A0",
                "unit": "m",
                "completed_quantity": "50",
                "actual_person_hours": "8.00",
            },
        ],
    )

    client.force_login(ctx["alice"])
    # Test group timesheet and group work exports
    res_group_ts = client.get(reverse("customer_portal:group_timesheet_xlsx", args=[ctx["group_a"].pk]))
    assert res_group_ts.status_code == 200
    wb_group = load_workbook(BytesIO(res_group_ts.content))

    # All sheet names must be <= 31 chars and contain no forbidden characters
    for name in wb_group.sheetnames:
        assert len(name) <= 31
        assert not any(c in name for c in r"\/?*[]:")

    # Summary sheet must exist
    assert "Übersicht" in wb_group.sheetnames

    # Check cell data types
    sheet2 = wb_group[wb_group.sheetnames[1]]
    for row in sheet2.iter_rows():
        for cell in row:
            if cell.value in ("=SUM(A1:A10)", "+cmd|' /C calc'!A0"):
                assert cell.data_type == "s"
            if cell.value in (6.5, 8.0, 14.5, 15, 50):
                assert cell.data_type == "n"
