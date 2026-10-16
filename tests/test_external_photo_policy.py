"""External work policy keeps cleanup internal and photos as a free site package."""

from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO

from features.customer_portal.completion_projection import project_customer_entries
from features.mbg_reports.completion_projection import project_accepted_entries
from features.mbg_reports.export_workbooks import build_final_hourly_workbook, build_weekly_workbook
from features.mbg_reports.master_workbook import build_master_settlement_workbook
from features.mbg_reports.models import MbgReportRevision
from features.mbg_reports.pricing_projection import price_completion_row
from features.mbg_reports.reporting_selectors import load_accepted_reports_dataset
from features.mbg_reports.settlement_scope import project_settlement_dataset
from openpyxl import load_workbook
from test_settlement_billing_scope import billing_scope  # noqa: F401


def snapshot(site=1):
    return {
        "site_id": site,
        "entries": [
            {
                "work_type_code": code,
                "description": name,
                "unit": "Std.",
                "quantity": "1",
                "calculation_rule": "zero" if code == "photo-documentation" else "actual",
                "actual_person_hours": hours,
                "starts_at": "17:00",
                "ends_at": "17:30",
                "internal_only": internal,
                "participants": [{"name": "Worker"}],
            }
            for code, name, hours, internal in [
                ("freiraeumen", "Cleanup", "1.00", False),
                ("photo-documentation", "Photo first", "0.50", False),
                ("photo-documentation", "Photo second", "0.75", False),
                ("photo-documentation", "Private photo", "0.25", True),
            ]
        ],
    }


def test_customer_photos_one_package_per_site_preserve_indexes_and_history():
    original = snapshot()
    before = deepcopy(original)
    rows = project_customer_entries(
        [
            (10, 1, date(2026, 10, 16), original),
            (11, 1, date(2026, 10, 17), snapshot()),
            (12, 1, date(2026, 10, 17), snapshot(2)),
        ]
    )
    assert len(rows) == 2
    for row in rows:
        assert (row.quantity, row.unit, row.actual_person_hours) == (Decimal("1"), "psch.", Decimal("0"))
        assert {s["entry_index"] for s in row.sources} == {1, 2}
        assert all("actual_person_hours" not in s and "normative_person_hours" not in s for s in row.sources)
        assert price_completion_row(row).state == "free"
    assert original == before
    internal = project_accepted_entries([(10, 1, date(2026, 10, 16), original)], include_internal=True)
    assert sum(r.actual_person_hours for r in internal) == Decimal("2.50")


def test_real_billing_renderers_hide_photo_time_and_keep_internal_weekly(billing_scope):  # noqa: F811
    site, _ = billing_scope
    revision = MbgReportRevision.objects.filter(report__site=site).order_by("report__work_date").first()
    for revision in MbgReportRevision.objects.filter(report__site=site):
        revision.snapshot["entries"].extend(snapshot()["entries"])
        revision.save(update_fields=["snapshot"])
    revision = MbgReportRevision.objects.filter(report__site=site).order_by("report__work_date").first()
    full = load_accepted_reports_dataset()
    projected = project_settlement_dataset(full)
    summary = projected.all_sites[site.pk]
    assert summary.total_hours == Decimal("4")
    photos = [e for d in summary.day_facts for e in d.entries if e.work_type_code == "photo-documentation"]
    assert len(photos) == 1
    assert (photos[0].quantity, photos[0].unit, photos[0].starts_at, photos[0].ends_at) == (
        Decimal("1"),
        "psch.",
        "",
        "",
    )
    assert projected == project_settlement_dataset(projected)
    book = build_master_settlement_workbook(dataset=full)
    text = "\n".join(str(c.value) for s in book for row in s for c in row)
    assert "Cleanup" not in text and "Private photo" not in text
    assert "Photo first" in text
    from features.customer_portal.services import review_work_snapshot

    actual_review = review_work_snapshot(site)
    assert not any(row["name"] == "Cleanup" for row in actual_review)
    photo = next(row for row in actual_review if row["kind"] == "photo_package")
    assert (photo["quantity"], photo["unit"]) == ("1", "psch.")
    assert len(photo["sources"]) == 4
    assert all("actual_person_hours" not in source for source in photo["sources"])
    final = load_workbook(
        BytesIO(build_final_hourly_workbook([revision.snapshot], date(2026, 9, 22), date(2026, 9, 22)))
    )
    text = "\n".join(str(c.value) for s in final for row in s for c in row)
    assert "Photo first" not in text and "Cleanup" not in text
    weekly = load_workbook(
        BytesIO(build_weekly_workbook({"grouped": {site.pk: [revision.snapshot]}, "year": 2026, "week": 39}))
    )
    text = "\n".join(str(c.value) for s in weekly for row in s for c in row)
    assert "Cleanup" in text and "Photo first" in text and "Private photo" in text


def test_new_performance_artifact_uses_photo_package_and_excludes_cleanup(billing_scope, monkeypatch):  # noqa: F811
    from django.contrib.auth import get_user_model
    from features.mbg_reports import export_services
    from features.mbg_reports.export_workbooks import build_performance_workbook

    site, _ = billing_scope
    site.address = "Test street"
    site.project_number = "P-TEST"
    site.save(update_fields=["address", "project_number"])
    owner = get_user_model().objects.get(username="billing-scope-owner")
    owner.is_superuser = True
    owner.save(update_fields=["is_superuser"])
    for revision in MbgReportRevision.objects.filter(report__site=site):
        revision.snapshot["entries"].extend(snapshot()["entries"])
        revision.save(update_fields=["snapshot"])
    monkeypatch.setattr(export_services, "store_export_artifact", lambda artifact, payload: artifact)
    artifact = export_services.generate_performance_report(actor=owner, group_id=site.group_id)
    rows = artifact.snapshot["sites"][0]["rows"]
    assert not any(row["code"] == "freiraeumen" for row in rows)
    photo = next(row for row in rows if row["code"] == "photo-documentation")
    assert (photo["quantity"], photo["unit"], photo["price_state"], photo["amount"]) == ("1", "psch.", "free", "0.00")
    assert photo["actual_person_hours"] == "0"
    book = load_workbook(BytesIO(build_performance_workbook(artifact.snapshot)))
    row = next(row for row in book.worksheets[1].values if row[0] == "Photo first")
    assert row[1:4] == ("psch.", 1, 0)
