"""Characterization tests for MBG export rendering, generation services, and storage."""

from datetime import date, time
from io import BytesIO
from unittest.mock import patch

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.storage import FileSystemStorage
from features.mbg_reports.exports import (
    build_final_hourly_workbook,
    build_performance_workbook,
    build_weekly_workbook,
    build_work_completion_workbook,
    build_workbook,
    generate_final_hours,
    generate_performance_report,
    generate_week,
    generate_work_completion_report,
)
from features.mbg_reports.models import WeeklyExport
from features.mbg_reports.services import review_report, save_report
from features.projects.models import AcObject, ProjectGroup
from features.staff.models import Employee
from features.work_catalog.models import WorkType
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db


@pytest.fixture
def export_context():
    admin = get_user_model().objects.create_user(username="export-admin", is_superuser=True)
    regular_user = get_user_model().objects.create_user(username="regular-user", is_superuser=False)
    employee = Employee.objects.create(first_name="Max", last_name="Mustermann", has_ac_access=True)
    group = ProjectGroup.objects.create(name="Solarpark Ost", has_ac=True, status="active")
    site = AcObject.objects.create(
        group=group,
        name="Trafostation 1",
        address="Industriestr. 10",
        project_number="P-1001",
        status="in_progress",
        is_active=True,
    )
    work = WorkType.objects.get(code="wr-montage")
    row = {
        "participants": [employee],
        "starts_at": time(8),
        "ends_at": time(12),
        "work_type": work,
        "quantity": 1,
        "cable_type": None,
        "route": None,
        "standard_connections": 0,
        "heavy_connections": 0,
        "with_underconstruction": False,
        "description": "WR Montage characterization",
    }
    return admin, regular_user, employee, group, site, row


def test_export_authorization_enforced(export_context):
    _admin, regular_user, _employee, group, _site, _row = export_context

    with pytest.raises(PermissionDenied):
        generate_week(actor=regular_user, year=2026, week=38)

    with pytest.raises(PermissionDenied):
        generate_final_hours(actor=regular_user, group_id=group.pk)

    with pytest.raises(PermissionDenied):
        generate_performance_report(actor=regular_user, group_id=group.pk)

    with pytest.raises(PermissionDenied):
        generate_work_completion_report(actor=regular_user, group_id=group.pk)


def test_weekly_export_strict_rejection_rules(export_context):
    admin, _regular_user, employee, _group, site, row = export_context

    # 1. Invalid calendar week
    with pytest.raises(ValidationError, match="Некорректная календарная неделя"):
        generate_week(actor=admin, year=2026, week=99)

    # 2. No reports for week
    with pytest.raises(ValidationError, match="Для выгрузки нужны принятые отчёты"):
        generate_week(actor=admin, year=2026, week=38)

    # 3. Draft/unaccepted report
    target_date = date.fromisocalendar(2026, 38, 1)
    save_report(
        actor=admin,
        site=site,
        work_date=target_date,
        workforce=[employee],
        entries=[row],
        submit=False,
    )
    with pytest.raises(ValidationError, match="Для выгрузки нужны принятые отчёты"):
        generate_week(actor=admin, year=2026, week=38)


def test_weekly_export_snapshot_immutability_and_filename(export_context, monkeypatch, tmp_path):
    monkeypatch.setattr(WeeklyExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, _regular_user, employee, _group, site, row = export_context
    target_date = date.fromisocalendar(2026, 38, 2)
    report = save_report(
        actor=admin,
        site=site,
        work_date=target_date,
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")

    artifact = generate_week(actor=admin, year=2026, week=38)
    assert artifact.filename == "AC - KW 38 - 2026.xlsx"
    assert artifact.snapshot[0]["address"] == "Industriestr. 10"

    # Mutate live site
    site.address = "Changed Address"
    site.save()

    # Artifact snapshot remains unchanged
    assert artifact.snapshot[0]["address"] == "Industriestr. 10"

    repeated = generate_week(actor=admin, year=2026, week=38)
    assert repeated.pk == artifact.pk
    assert repeated.fingerprint == artifact.fingerprint
    assert repeated.file.name == artifact.file.name
    assert WeeklyExport.objects.count() == 1


def test_export_file_cleanup_on_save_failure(export_context, monkeypatch, tmp_path):
    monkeypatch.setattr(WeeklyExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, _regular_user, employee, _group, site, row = export_context
    target_date = date.fromisocalendar(2026, 39, 1)
    report = save_report(
        actor=admin,
        site=site,
        work_date=target_date,
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")

    with patch.object(WeeklyExport, "save", side_effect=RuntimeError("Simulated DB failure")):
        with pytest.raises(RuntimeError, match="Simulated DB failure"):
            generate_week(actor=admin, year=2026, week=39)

    # Ensure no orphaned files left in storage
    files = [f for f in tmp_path.rglob("*") if f.is_file()]
    assert files == []


def test_build_work_completion_workbook_in_memory_rendering():
    snapshot = {
        "group_name": "Test Group",
        "period_label": "01.09.2026 - 15.09.2026",
        "sites": [
            {
                "site_id": 1,
                "site_name": "Site Alpha",
                "group_name": "Test Group",
                "project_number": "P-101",
                "address": "Testweg 1",
                "status": "completed",
                "period_start": "2026-09-01",
                "period_end": "2026-09-15",
                "unresolved_quantity_count": 1,
                "rows": [
                    {
                        "name": "AC Trasse",
                        "unit": "m",
                        "quantity": None,
                        "actual_person_hours": "3.50",
                        "sources": [{"work_date": "2026-09-01", "report_id": 10, "recorded_quantity": "ungeklärt"}],
                    }
                ],
            }
        ],
    }
    content = build_work_completion_workbook(snapshot)
    wb = load_workbook(BytesIO(content))
    assert "Übersicht" in wb.sheetnames
    detail_sheet = wb["Site Alpha"]
    # Check text protection for unknown quantity
    assert detail_sheet["C9"].value == "Menge ungeklärt"
    assert detail_sheet["C9"].data_type == "s"
    assert detail_sheet["D9"].value == 3.50
    # Warning present
    assert any("Menge ungeklärt" in str(cell.value) for row in detail_sheet.iter_rows() for cell in row)


def test_build_performance_workbook_in_memory_rendering():
    snapshot = {
        "group_name": "Anlage Test",
        "is_preliminary": False,
        "export_date": "04.10.2026",
        "sites": [
            {
                "site_id": 10,
                "site_name": "Dach 1",
                "project_number": "P-99",
                "address": "Solarallee 5",
                "status": "completed",
                "amount": "150.00",
                "rows": [
                    {
                        "name": "WR Montage",
                        "unit": "Stk.",
                        "quantity": "2",
                        "unit_price": "75.00",
                    }
                ],
            }
        ],
    }
    content = build_performance_workbook(snapshot)
    wb = load_workbook(BytesIO(content))
    summary = wb["Übersicht"]
    assert summary["A1"].value == "Leistungsnachweis"
    detail = wb["Dach 1"]
    assert detail["E9"].value == "=IF(ISNUMBER(C9), ROUND(C9*D9, 2), 0)"
    assert detail["E10"].value == "=SUM(E9:E9)"


def test_build_workbook_and_final_hourly_in_memory_rendering():
    snapshots = [
        {
            "site_id": 1,
            "site_name": "Haus A",
            "group_name": "Gruppe 1",
            "project_number": "P-01",
            "address": "Weg 1",
            "date": "2026-09-14",
            "entries": [
                {
                    "description": "Kabelverlegung",
                    "participants": [{"name": "Worker 1"}],
                    "start": "08:00",
                    "end": "12:00",
                    "actual_person_hours": "4.00",
                }
            ],
        }
    ]
    # 1. Weekly workbook
    weekly_bytes = build_workbook(snapshots, 2026, 38)
    wb_weekly = load_workbook(BytesIO(weekly_bytes))
    assert wb_weekly.sheetnames == ["Übersicht", "Haus A"]
    assert wb_weekly["Übersicht"]["A1"].value == "Stundennachweis"
    assert wb_weekly["Haus A"]["F8"].value == "=SUM(F7:F7)"

    # Direct build_weekly_workbook invocation
    weekly_direct_bytes = build_weekly_workbook({"grouped": {1: snapshots}, "week": 38, "year": 2026})
    wb_direct = load_workbook(BytesIO(weekly_direct_bytes))
    assert wb_direct.sheetnames == ["Übersicht", "Haus A"]

    # 2. Final hourly workbook
    final_bytes = build_final_hourly_workbook(snapshots, date(2026, 9, 1), date(2026, 9, 30))
    wb_final = load_workbook(BytesIO(final_bytes))
    assert wb_final.sheetnames == ["Übersicht", "Haus A"]
    assert wb_final["Übersicht"]["A1"].value == "Gesamtstundennachweis"
    assert wb_final["Haus A"]["B8"].value == "Summe Zeitraum"
