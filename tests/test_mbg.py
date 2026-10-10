"""Клиентское время независимо от внутренних смен и от других объектов."""

from datetime import date, time, timedelta
from decimal import Decimal
from uuid import uuid4

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from features.internal_time.services import save_report as save_internal
from features.mbg_reports.forms import ReportForm
from features.mbg_reports.services import review_report, save_report
from features.projects.models import AcObject, ProjectGroup
from features.staff.models import Employee
from features.work_catalog.models import WorkType

pytestmark = pytest.mark.django_db


@pytest.fixture
def mbg_data():
    admin = get_user_model().objects.create_user(username="owner", is_superuser=True)
    employee = Employee.objects.create(first_name="Test", last_name="Employee", has_ac_access=True)
    group = ProjectGroup.objects.create(name="Berlin")
    site = AcObject.objects.create(group=group, name="Haus 42", address="Berlin 42", project_number="MBG-42")
    row = dict(
        participants=[employee],
        starts_at=time(9),
        ends_at=time(10),
        work_type=WorkType.objects.get(code="wr-montage"),
        quantity=1,
        cable_type=None,
        route=None,
        standard_connections=0,
        heavy_connections=0,
        with_underconstruction=False,
        description="Montage WR",
    )
    return admin, employee, group, site, row


def test_mbg_time_is_independent(mbg_data):
    admin, employee, group, site, row = mbg_data
    save_internal(
        actor=admin,
        data=dict(
            project_group=group,
            direction="ac",
            work_date=timezone.localdate(),
            starts_at=time(7),
            ends_at=time(18),
            lunch_minutes=30,
            note="",
            work_description="",
            submission_key=uuid4(),
        ),
        participants=[dict(employee=employee, deduction_minutes=0)],
    )
    report = save_report(
        actor=admin, site=site, work_date=timezone.localdate(), workforce=[employee], entries=[row], submit=True
    )
    assert report.revisions.first().snapshot["entries"][0]["minutes"] == 60
    reviewed = review_report(actor=admin, report_id=report.pk, version=1, decision="accept")
    assert reviewed.status == "accepted"
    assert reviewed.revisions.count() == 2
    with pytest.raises(ValidationError):
        save_report(
            actor=admin,
            site=site,
            work_date=timezone.localdate(),
            workforce=[employee],
            entries=[row],
            report_id=report.pk,
            version=2,
        )


def test_first_saved_mbg_report_starts_site_and_ac_direction(mbg_data):
    admin, employee, group, site, row = mbg_data

    assert site.status == "not_started"
    assert group.ac_status == "not_started"

    save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )

    site.refresh_from_db()
    group.refresh_from_db()
    assert site.status == "started"
    assert group.ac_status == "started"


def test_completed_site_rejects_new_report_until_reopened(mbg_data):
    admin, employee, _group, site, row = mbg_data
    site.status = "completed"
    site.save(update_fields=["status"])

    with pytest.raises(ValidationError, match="завершён"):
        save_report(
            actor=admin,
            site=site,
            work_date=timezone.localdate(),
            workforce=[employee],
            entries=[row],
        )


def test_site_overview_is_read_first_and_shows_operational_totals(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    colleague = Employee.objects.create(first_name="Second", last_name="Worker", has_ac_access=True)
    save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee, colleague],
        entries=[{**row, "participants": [employee, colleague]}],
    )
    client.force_login(admin)

    response = client.get(reverse("projects:site_detail", args=[group.pk, site.pk]))

    assert response.status_code == 200
    assert response.context["activity"].recorded_person_hours == Decimal("2.00")
    assert response.context["activity"].employee_count == 2
    assert response.context["activity"].report_count == 1
    assert response.context["positions"] == []  # Черновик не становится выполненной работой.
    content = response.content.decode()
    assert "Обзор объекта" in content
    assert "Подтверждённая сумма" in content


def test_ac_object_list_is_grouped_by_anlage_and_links_to_site_overview(mbg_data, client):
    admin, _employee, group, site, _row = mbg_data
    client.force_login(admin)

    response = client.get("/ru/cabinet/ac/objects/")

    assert response.status_code == 200
    assert len(response.context["group_cards"]) == 1
    assert response.context["group_cards"][0].group == group
    assert reverse("projects:site_detail", args=[group.pk, site.pk]) in response.content.decode()


def test_anlage_detail_shows_aggregated_mbg_metrics(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("projects:detail", args=[group.pk]))

    assert response.status_code == 200
    assert response.context["group_activity"].recorded_person_hours == Decimal("1.00")
    assert "человеко-часов MBG" in response.content.decode()


def test_overlap_across_objects_and_permissions(mbg_data):
    admin, _employee, group, site, row = mbg_data
    save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
    )
    other = AcObject.objects.create(group=group, name="Haus 40")
    with pytest.raises(ValidationError):
        save_report(
            actor=admin,
            site=other,
            work_date=timezone.localdate(),
            workforce=[row["participants"][0]],
            entries=[row],
        )
    user = get_user_model().objects.create_user(username="worker")
    with pytest.raises(PermissionDenied):
        save_report(
            actor=user,
            site=other,
            work_date=timezone.localdate(),
            workforce=[row["participants"][0]],
            entries=[row],
        )


def test_weekly_export_snapshot_and_idempotency(mbg_data, monkeypatch, tmp_path):
    from django.core.files.storage import FileSystemStorage
    from features.mbg_reports.exports import generate_week
    from features.mbg_reports.models import WeeklyExport
    from openpyxl import load_workbook

    monkeypatch.setattr(WeeklyExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")
    iso = timezone.localdate().isocalendar()
    other_group = ProjectGroup.objects.create(name="Leipzig")
    other_site = AcObject.objects.create(
        group=other_group,
        name="Haus 8",
        address="Leipzig 8",
        project_number="MBG-8",
    )
    other_report = save_report(
        actor=admin,
        site=other_site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[{**row, "starts_at": time(10), "ends_at": time(11)}],
        submit=True,
    )
    review_report(actor=admin, report_id=other_report.pk, version=1, decision="accept")
    artifact = generate_week(actor=admin, year=iso.year, week=iso.week)
    assert generate_week(actor=admin, year=iso.year, week=iso.week).pk == artifact.pk
    with artifact.file.open("rb") as source:
        book = load_workbook(source)
        assert len(book.worksheets) == 3
        assert book.active.title == "Übersicht"
        assert {book.active.cell(row, 1).value for row in range(5, 7)} == {"Berlin", "Leipzig"}
        sheet = next(sheet for sheet in book.worksheets if sheet["B3"].value == "MBG-42")
        assert sheet["B3"].value == "MBG-42"
        assert sheet["E7"].value is None
        assert sheet["F7"].value == 1
        assert sheet["F8"].value in (1, "=SUM(F7:F7)")
    site.project_number = "CHANGED"
    site.save()
    assert artifact.snapshot[0]["project_number"] == "MBG-42"


def test_final_hourly_export_uses_full_completed_anlage_period(mbg_data, monkeypatch, tmp_path):
    from django.core.files.storage import FileSystemStorage
    from features.mbg_reports.exports import generate_final_hours
    from features.mbg_reports.models import FinalHourlyExport
    from openpyxl import load_workbook

    monkeypatch.setattr(FinalHourlyExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, employee, group, site, row = mbg_data
    first_date = timezone.localdate() - timedelta(days=8)
    last_date = timezone.localdate()
    for work_date in (first_date, last_date):
        report = save_report(
            actor=admin,
            site=site,
            work_date=work_date,
            workforce=[employee],
            entries=[row],
            submit=True,
        )
        review_report(actor=admin, report_id=report.pk, version=report.version, decision="accept")

    with pytest.raises(ValidationError, match="все активные объекты"):
        generate_final_hours(actor=admin, group_id=group.pk)

    site.status = "completed"
    site.save(update_fields=["status"])
    artifact = generate_final_hours(actor=admin, group_id=group.pk)

    assert generate_final_hours(actor=admin, group_id=group.pk).pk == artifact.pk
    assert artifact.period_start == first_date
    assert artifact.period_end == last_date
    assert artifact.snapshot["reports"][0]["date"] == str(first_date)
    with artifact.file.open("rb") as source:
        book = load_workbook(source, data_only=False)
        assert book.active.title == "Übersicht"
        assert book.active["A1"].value == "Gesamtstundennachweis"
        assert book.active["A2"].value == "Zeitraum:"
        assert book.active["B2"].value == f"{first_date:%d.%m.%Y}-{last_date:%d.%m.%Y}"
        site_sheet = book.worksheets[1]
        assert site_sheet["A4"].value == "Zeitraum"
        assert site_sheet["B4"].value == f"{first_date:%d.%m.%Y}-{last_date:%d.%m.%Y}"
        assert site_sheet["B9"].value == "Summe Zeitraum"
        assert site_sheet["F9"].value == "=SUM(F7:F8)"


def test_export_hub_offers_final_hours_only_for_fully_completed_anlage(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=report.version, decision="accept")
    site.status = "completed"
    site.save(update_fields=["status"])
    client.force_login(admin)

    content = client.get(reverse("mbg_reports:export_hub")).content.decode()

    assert "Итоговый почасовой отчёт" in content
    assert reverse("mbg_reports:final_hours_export", args=[group.pk]) in content


def test_leistungsnachweis_requires_completed_site_and_is_idempotent(mbg_data, monkeypatch, tmp_path):
    from django.core.files.storage import FileSystemStorage
    from features.mbg_reports.exports import generate_performance_report
    from features.mbg_reports.models import PerformanceExport
    from openpyxl import load_workbook

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")

    with pytest.raises(ValidationError, match="завершён"):
        generate_performance_report(actor=admin, group_id=group.pk)

    site.status = "completed"
    site.save(update_fields=["status"])
    other_site = AcObject.objects.create(
        group=group,
        name="Haus 44",
        address="Berlin 44",
        project_number="MBG-44",
    )
    other_report = save_report(
        actor=admin,
        site=other_site,
        work_date=timezone.localdate() - timedelta(days=1),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=other_report.pk, version=1, decision="accept")
    other_site.status = "completed"
    other_site.save(update_fields=["status"])
    artifact = generate_performance_report(actor=admin, group_id=group.pk)
    assert generate_performance_report(actor=admin, group_id=group.pk).pk == artifact.pk
    with artifact.file.open("rb") as source:
        book = load_workbook(source, data_only=True)
        assert len(book.worksheets) == 3
        assert book.active.title == "Übersicht"
        assert {book.active.cell(row, 1).value for row in range(5, 7)} == {group.name}
        assert {book.active.cell(row, 2).value for row in range(5, 7)} == {site.name, other_site.name}
        sheet = next(sheet for sheet in book.worksheets[1:] if sheet["B3"].value == site.name)
        assert sheet["B2"].value == site.group.name
        assert sheet["B3"].value == site.name
        assert sheet["B4"].value == site.project_number
        assert any("Montage WR" in str(row[0].value) for row in sheet.iter_rows(min_row=8))


def test_mbg_pages_and_partial_rows(mbg_data, client):
    admin, _employee, _group, site, _row = mbg_data
    client.force_login(admin)
    for name in ("mbg_reports:list", "mbg_reports:new", "work_catalog:list", "projects:list"):
        assert client.get(reverse(name)).status_code == 200
    response = client.post(
        reverse("mbg_reports:new"),
        {"action": "add", "entries-TOTAL_FORMS": "1", "site": str(site.pk), "submission_key": str(uuid4())},
        HTTP_HX_REQUEST="true",
    )
    assert response.status_code == 200
    assert "entries-1-participants" in response.content.decode()


def test_editor_can_move_a_later_work_card_up_without_losing_entered_values(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    client.force_login(admin)

    response = client.post(
        reverse("mbg_reports:new"),
        {
            "action": "move:1:up",
            "group": str(group.pk),
            "site": str(site.pk),
            "work_date": str(timezone.localdate()),
            "workforce": [str(employee.pk)],
            "submission_key": str(uuid4()),
            "entries-TOTAL_FORMS": "2",
            "entries-INITIAL_FORMS": "0",
            "entries-MIN_NUM_FORMS": "1",
            "entries-MAX_NUM_FORMS": "100",
            "entries-0-participants": [str(employee.pk)],
            "entries-0-starts_at": "07:00",
            "entries-0-ends_at": "08:00",
            "entries-0-work_type": str(row["work_type"].pk),
            "entries-0-quantity": "1",
            "entries-0-standard_connections": "0",
            "entries-0-heavy_connections": "0",
            "entries-0-description": "First work",
            "entries-1-participants": [str(employee.pk)],
            "entries-1-starts_at": "08:00",
            "entries-1-ends_at": "09:00",
            "entries-1-work_type": str(row["work_type"].pk),
            "entries-1-quantity": "1",
            "entries-1-standard_connections": "0",
            "entries-1-heavy_connections": "0",
            "entries-1-description": "Second work",
        },
        HTTP_HX_REQUEST="true",
    )
    content = response.content.decode()

    assert response.status_code == 200
    assert content.index('value="Second work"') < content.index('value="First work"')
    assert 'value="move:0:up"' in content
    assert 'value="move:1:down"' in content


def test_invalid_mbg_form_returns_visible_error_toast(mbg_data, client):
    admin, _employee, _group, _site, _row = mbg_data
    client.force_login(admin)

    response = client.post(
        reverse("mbg_reports:new"),
        {
            "action": "save",
            "entries-TOTAL_FORMS": "1",
            "entries-INITIAL_FORMS": "0",
            "entries-MIN_NUM_FORMS": "1",
            "entries-MAX_NUM_FORMS": "100",
        },
        HTTP_HX_REQUEST="true",
    )
    content = response.content.decode()

    assert response.status_code == 200
    assert "data-error-toast" in content
    assert "Отчёт не сохранён" in content
    assert 'role="alert"' in content


def test_report_list_has_filters_without_export_projects(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    other_group = ProjectGroup.objects.create(name="Leipzig")
    other_site = AcObject.objects.create(group=other_group, name="Haus 8", project_number="MBG-8")
    other_report = save_report(
        actor=admin,
        site=other_site,
        work_date=timezone.localdate() - timedelta(days=1),
        workforce=[employee],
        entries=[{**row, "starts_at": time(11), "ends_at": time(12)}],
        submit=True,
    )
    client.force_login(admin)

    response = client.get(
        reverse("mbg_reports:list"),
        {"group": group.pk, "site": site.pk, "status": "draft"},
    )
    content = response.content.decode()

    assert response.status_code == 200
    assert 'name="group"' in content
    assert 'name="site"' in content
    assert 'name="status"' in content
    assert 'name="date_from"' in content
    assert 'name="date_to"' in content
    assert "Проекты и выгрузки" not in content
    assert reverse("mbg_reports:detail", args=[report.pk]) in content
    assert reverse("mbg_reports:detail", args=[other_report.pk]) not in content
    assert 'data-searchable-choice="group"' in content
    assert 'data-searchable-choice="site"' in content


def test_export_workspace_is_superuser_only(mbg_data, client):
    admin, _employee, group, _site, _row = mbg_data
    manager = get_user_model().objects.create_user(username="mbg-manager")
    manager.user_permissions.add(Permission.objects.get(codename="manage_reports"))

    client.force_login(admin)
    response = client.get(reverse("mbg_reports:export_hub"))
    assert response.status_code == 200
    assert client.get(reverse("mbg_reports:exports")).status_code == 200

    client.force_login(manager)
    assert client.get(reverse("mbg_reports:list")).status_code == 200
    assert client.get(reverse("mbg_reports:export_hub")).status_code == 200
    assert client.get(reverse("mbg_reports:exports")).status_code == 403
    assert client.get(reverse("mbg_reports:group_overview", args=[group.pk])).status_code == 403


def test_export_hub_lists_week_cards_with_anlage_metrics_newest_first(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    colleague = Employee.objects.create(first_name="Second", last_name="Worker", has_ac_access=True)
    other_group = ProjectGroup.objects.create(name="Leipzig")
    other_site = AcObject.objects.create(
        group=other_group,
        name="Haus 8",
        address="Leipzig 8",
        project_number="MBG-8",
    )
    save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 10),
        workforce=[employee],
        entries=[row],
    )
    save_report(
        actor=admin,
        site=other_site,
        work_date=date(2026, 9, 10),
        workforce=[colleague],
        entries=[{**row, "participants": [colleague]}],
    )
    save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 3),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:export_hub"))

    cards = response.context["weekly_cards"]
    assert [(card.year, card.week) for card in cards] == [(2026, 37), (2026, 36)]
    assert [(metric.name, metric.employee_count, metric.person_hours) for metric in cards[0].groups] == [
        ("Berlin", 1, Decimal("1.00")),
        ("Leipzig", 1, Decimal("1.00")),
    ]
    content = response.content.decode()
    assert content.index("Календарная неделя 37") < content.index("Календарная неделя 36")
    assert content.count('data-week-card="true"') == 2


def test_week_card_uses_only_current_global_export(mbg_data):
    from features.mbg_reports.models import WeeklyExport
    from features.mbg_reports.selectors import weekly_export_cards

    admin, employee, _group, site, row = mbg_data
    save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 10),
        workforce=[employee],
        entries=[row],
    )
    current = WeeklyExport.objects.create(
        year=2026,
        week=37,
        fingerprint="current-weekly-export",
        template_version="ac-weekly-4",
        snapshot=[],
        file=SimpleUploadedFile("current.xlsx", b"current"),
        filename="current.xlsx",
        author=admin,
    )
    WeeklyExport.objects.create(
        year=2026,
        week=37,
        fingerprint="legacy-weekly-export",
        template_version="prototype-1",
        snapshot=[],
        file=SimpleUploadedFile("legacy.xlsx", b"legacy"),
        filename="legacy.xlsx",
        author=admin,
    )

    assert weekly_export_cards()[0].latest_export == current


def test_export_hub_counts_only_accepted_preliminary_reports(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    accepted = save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 10),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=accepted.pk, version=1, decision="accept")
    save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 11),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:export_hub"))

    assert response.context["total_preliminary_reports"] == 1
    assert response.context["preliminary_groups_with_sites"][0]["sites"][0].accepted_reports_count == 1


def test_export_hub_uses_external_interaction_script(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 10),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")
    client.force_login(admin)

    content = client.get(reverse("mbg_reports:export_hub")).content.decode()

    assert "mbg_reports/js/export_hub.js" in content
    assert "onclick=" not in content
    assert "onsubmit=" not in content
    assert "onchange=" not in content


def test_report_detail_supports_legacy_snapshot_without_calculated_hours(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=date(2026, 9, 10),
        workforce=[employee],
        entries=[row],
    )
    revision = report.revisions.get(version=report.version)
    snapshot = revision.snapshot
    snapshot["entries"][0].pop("minutes", None)
    snapshot["entries"][0].pop("actual_person_hours", None)
    revision.snapshot = snapshot
    revision.save(update_fields=["snapshot"])
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:detail", args=[report.pk]))

    assert response.status_code == 200
    assert response.context["total_minutes"] == 60


def test_completed_site_is_listed_in_performance_export_workspace(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")
    site.status = "completed"
    site.save(update_fields=["status"])
    client.force_login(admin)

    hub = client.get(reverse("mbg_reports:export_hub"))
    assert "Недельный табель часов" in hub.content.decode()
    assert "Промежуточный акт" in hub.content.decode()
    assert "итоговый акт выполненных работ" in hub.content.decode().lower()
    assert reverse("mbg_reports:performance_exports", args=[site.group_id]) in hub.content.decode()

    page = client.get(reverse("mbg_reports:performance_exports", args=[site.group_id]))
    assert page.status_code == 200
    assert site.name in page.content.decode()
    assert "Montage WR" in page.content.decode()


def test_export_hub_shows_preliminary_for_active_projects(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report.pk, version=1, decision="accept")
    site.status = "started"
    site.save(update_fields=["status"])
    client.force_login(admin)

    hub = client.get(reverse("mbg_reports:export_hub"))
    assert "Промежуточный акт" in hub.content.decode()
    assert site.group.name in hub.content.decode()
    assert reverse("mbg_reports:performance_exports", args=[site.group_id]) in hub.content.decode()



def test_report_form_filters_sites_by_selected_group(mbg_data):
    _admin, employee, group, site, _row = mbg_data
    other_group = ProjectGroup.objects.create(name="Leipzig")
    other_site = AcObject.objects.create(group=other_group, name="Haus 8", project_number="MBG-8")

    empty_form = ReportForm()
    assert list(empty_form.fields["site"].queryset) == []

    form = ReportForm(
        data={
            "group": str(group.pk),
            "site": str(other_site.pk),
            "work_date": str(timezone.localdate()),
            "workforce": [str(employee.pk)],
            "submission_key": str(uuid4()),
        }
    )

    assert list(form.fields["site"].queryset) == [site]
    assert not form.is_valid()
    assert "site" in form.errors


def test_site_options_endpoint_returns_only_selected_group_sites(mbg_data, client):
    admin, _employee, group, site, _row = mbg_data
    other_group = ProjectGroup.objects.create(name="Leipzig")
    other_site = AcObject.objects.create(group=other_group, name="Haus 8", project_number="MBG-8")
    inactive_site = AcObject.objects.create(group=group, name="Haus 44", is_active=False)
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:site_options"), {"group": group.pk})
    content = response.content.decode()

    assert response.status_code == 200
    assert site.name in content
    assert other_site.name not in content
    assert inactive_site.name not in content


def test_site_history_endpoint_shows_only_selected_site_work(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    saved = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    other_site = AcObject.objects.create(group=group, name="Haus 40", project_number="MBG-40")
    save_report(
        actor=admin,
        site=other_site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[{**row, "starts_at": time(10), "ends_at": time(11), "description": "Other site work"}],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:site_history"), {"site": site.pk})
    content = response.content.decode()

    assert response.status_code == 200
    assert 'data-site-history' in content
    assert f'data-loaded-site="{site.pk}"' in content
    assert "Уже внесено по объекту" in content
    assert saved.entries.get().description in content
    assert "Other site work" not in content
    assert str(employee) in content
    assert "1,00 чел.-ч" in content


def test_editor_site_history_excludes_report_being_edited(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:edit", args=[report.pk]))
    content = response.content.decode()
    history_panel = content[content.index('<aside class="mbg-site-history"') : content.index("</aside>")]

    assert response.status_code == 200
    assert "По этому объекту ещё нет других сохранённых работ." in content
    assert report.entries.get().description not in history_panel


def test_edit_form_restores_group_and_site(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:edit", args=[report.pk]))
    content = response.content.decode()

    assert response.status_code == 200
    assert f'<option value="{group.pk}" selected>' in content
    assert f'<option value="{site.pk}" selected>' in content


def test_editing_accepted_report_marks_missing_reason_field_invalid(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    report = review_report(actor=admin, report_id=report.pk, version=report.version, decision="accept")
    client.force_login(admin)

    response = client.post(
        reverse("mbg_reports:edit", args=[report.pk]),
        {
            "action": "save",
            "group": str(group.pk),
            "site": str(site.pk),
            "work_date": str(report.work_date),
            "workforce": [str(employee.pk)],
            "version": str(report.version),
            "submission_key": str(report.submission_key),
            "reason": "",
            "entries-TOTAL_FORMS": "1",
            "entries-INITIAL_FORMS": "1",
            "entries-MIN_NUM_FORMS": "1",
            "entries-MAX_NUM_FORMS": "100",
            "entries-0-participants": [str(employee.pk)],
            "entries-0-starts_at": "09:00",
            "entries-0-ends_at": "10:00",
            "entries-0-work_type": str(row["work_type"].pk),
            "entries-0-quantity": "1",
            "entries-0-cable_type": "",
            "entries-0-route": "",
            "entries-0-standard_connections": "0",
            "entries-0-heavy_connections": "0",
            "entries-0-description": "Montage WR",
        },
        HTTP_HX_REQUEST="true",
    )

    content = response.content.decode()
    assert response.status_code == 200
    assert "reason" in response.context["form"].errors
    assert "Укажите причину исправления принятого отчёта." in content
    assert 'id="id_reason"' in content
    assert 'aria-invalid="true"' in content


def test_employee_availability_returns_other_saved_mbg_intervals(mbg_data, client):
    admin, employee, group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    AcObject.objects.create(group=group, name="Haus 40", project_number="MBG-40")
    client.force_login(admin)

    response = client.get(
        reverse("mbg_reports:employee_availability"),
        {"work_date": str(timezone.localdate()), "employee": [employee.pk]},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload == {
        "employees": [
            {
                "id": employee.pk,
                "name": str(employee),
                "entries": [
                    {
                        "start": "09:00",
                        "end": "10:00",
                        "group": "Berlin",
                        "site": "Haus 42",
                        "report_id": report.pk,
                    }
                ],
            }
        ]
    }

    editing_response = client.get(
        reverse("mbg_reports:employee_availability"),
        {
            "work_date": str(timezone.localdate()),
            "employee": [employee.pk],
            "exclude_report": report.pk,
        },
    )
    assert editing_response.json()["employees"][0]["entries"] == []


def test_project_site_can_start_prefilled_mbg_report(mbg_data, client):
    admin, employee, group, site, _row = mbg_data
    client.force_login(admin)

    project_response = client.get(reverse("projects:detail", args=[group.pk]))
    expected_url = f'{reverse("mbg_reports:new")}?group={group.pk}&site={site.pk}'
    assert expected_url in project_response.content.decode().replace("&amp;", "&")

    form_response = client.get(
        reverse("mbg_reports:new"),
        {
            "group": group.pk,
            "site": site.pk,
            "work_date": str(timezone.localdate()),
            "workforce": [employee.pk],
        },
    )
    content = form_response.content.decode()
    assert f'<option value="{group.pk}" selected>' in content
    assert f'<option value="{site.pk}" selected>' in content
    assert f'value="{timezone.localdate():%Y-%m-%d}"' in content
    assert f'<option value="{employee.pk}" selected>' in content


def test_saved_report_offers_next_day_and_other_site_starts(mbg_data, client):
    admin, employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
    )
    client.force_login(admin)

    response = client.get(reverse("mbg_reports:detail", args=[report.pk]))
    next_day_url = response.context["next_day_url"]
    other_site_url = response.context["other_site_url"]

    assert f"group={site.group_id}" in next_day_url
    assert f"site={site.pk}" in next_day_url
    assert f"work_date={timezone.localdate() + timedelta(days=1):%Y-%m-%d}" in next_day_url
    assert f"workforce={employee.pk}" in next_day_url
    assert f"group={site.group_id}" in other_site_url
    assert "site=" not in other_site_url
    assert f"work_date={timezone.localdate():%Y-%m-%d}" in other_site_url


def test_stale_revision_and_return_reason(mbg_data):
    admin, _employee, _group, site, row = mbg_data
    report = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
        submit=True,
    )
    with pytest.raises(ValidationError):
        review_report(actor=admin, report_id=report.pk, version=1, decision="return")
    review_report(actor=admin, report_id=report.pk, version=1, decision="return", reason="Check time")
    with pytest.raises(ValidationError):
        save_report(
            actor=admin,
            site=site,
            work_date=timezone.localdate(),
            workforce=[row["participants"][0]],
            entries=[row],
            report_id=report.pk,
            version=1,
        )
    saved = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
        report_id=report.pk,
        version=2,
        submit=True,
    )
    assert saved.status == "submitted"
    assert saved.revisions.count() == 3


def test_export_multiple_sites_and_formula_text():
    from io import BytesIO

    from features.mbg_reports.exports import build_workbook
    from openpyxl import load_workbook

    data = [
        {
            "site_id": i,
            "site_name": f"House {i}",
            "address": "=2+2",
            "project_number": f"P-{i}",
            "date": "2026-09-09",
            "entries": [
                {
                    "participants": [
                        {"employee_id": 1, "name": "=2+2"},
                        {"employee_id": 2, "name": "Second Person"},
                    ],
                    "start": "09:00:00",
                    "end": "10:30:00",
                    "minutes": 90,
                    "description": "=HYPERLINK(test)",
                    "quantity": "1",
                    "route": "",
                    "cable_type": "",
                }
            ],
        }
        for i in (1, 2)
    ]
    book = load_workbook(BytesIO(build_workbook(data, 2026, 37)))
    assert len(book.worksheets) == 3
    assert book.active.title == "Übersicht"
    for sheet in book.worksheets[1:]:
        assert sheet["B7"].data_type == "s"
        assert sheet["G7"].data_type == "s"
        assert sheet["B2"].data_type == "s"
        assert sheet["B7"].value == "=2+2\nSecond Person"
        assert sheet["F7"].value == 3
        assert sheet["F8"].value in (3, "=SUM(F7:F7)")
        assert sheet["E7"].value is None


def test_submission_key_returns_existing_report(mbg_data):
    admin, _employee, _group, site, row = mbg_data
    key = uuid4()
    first = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
        submission_key=key,
    )
    repeated = save_report(
        actor=admin,
        site=site,
        work_date=timezone.localdate(),
        workforce=[row["participants"][0]],
        entries=[row],
        submission_key=key,
    )
    assert first.pk == repeated.pk
    assert first.revisions.count() == 1


def test_multi_group_zwischennachweis_and_view(mbg_data, monkeypatch, tmp_path, client):
    from django.core.files.storage import FileSystemStorage
    from features.mbg_reports.exports import generate_performance_report
    from features.mbg_reports.models import PerformanceExport
    from openpyxl import load_workbook

    monkeypatch.setattr(PerformanceExport._meta.get_field("file"), "storage", FileSystemStorage(location=tmp_path))
    admin, employee, group1, site1, row = mbg_data
    report1 = save_report(
        actor=admin,
        site=site1,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[row],
        submit=True,
    )
    review_report(actor=admin, report_id=report1.pk, version=1, decision="accept")

    group2 = ProjectGroup.objects.create(name="Potsdam", has_ac=True)
    site2 = AcObject.objects.create(
        group=group2,
        name="Haus 99",
        address="Potsdam 99",
        project_number="P-99",
    )
    report2 = save_report(
        actor=admin,
        site=site2,
        work_date=timezone.localdate(),
        workforce=[employee],
        entries=[{**row, "starts_at": time(10, 0), "ends_at": time(11, 0)}],
        submit=True,
    )
    review_report(actor=admin, report_id=report2.pk, version=1, decision="accept")

    # Generate multi-group preliminary export
    artifact = generate_performance_report(actor=admin, is_preliminary=True)
    assert artifact.group is None
    assert artifact.filename == "AC - Gesamt - Zwischennachweis.xlsx"

    with artifact.file.open("rb") as source:
        book = load_workbook(source)
        assert len(book.worksheets) == 3
        summary = book["Übersicht"]
        assert summary.cell(6, 1).value in (group1.name, group2.name)
        assert summary.cell(7, 1).value in (group1.name, group2.name)
        assert {summary.cell(6, 2).value, summary.cell(7, 2).value} == {site1.name, site2.name}

    # Test multi_performance_export POST view
    client.force_login(admin)
    resp = client.post(
        reverse("mbg_reports:multi_performance_export"),
        {"mode": "preliminary", "scope": "all"},
    )
    assert resp.status_code == 302
    assert "performance" in resp.url
