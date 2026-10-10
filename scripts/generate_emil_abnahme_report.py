# ruff: noqa: RUF001
"""Генератор ведомости приёмки и согласования выполненных работ для куратора (Эмиль).

Двусторонний формат:
- DLC (Михаил): заявляет выполненные работы («✓ Erledigt» или «In Ausführung»).
- Куратор (Эмиль): выбирает «✓ i.O.» / «✗ Nicht i.O.» и в соседнюю пустую ячейку пишет замечания.
- Блок приёмки внизу листа: общий статус («Freigegeben»), чек-листы и свободные строки для замечаний.
"""

import os
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "src"))
sys.path.insert(0, str(REPO_ROOT / "src" / "dlc_montage"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "dlc_montage.core.settings.dev")

import django

django.setup()

import openpyxl
from features.mbg_reports.reporting_selectors import load_accepted_reports_dataset
from features.projects.models import AcObject
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from sheet_titles import assign_site_sheet_titles

# Цветовая палитра DLC Montage
COLOR_NAVY_900 = "0A2540"
COLOR_NAVY_700 = "1E3A5F"
COLOR_NAVY_100 = "E6F0FA"

COLOR_GREEN_600 = "00875A"
COLOR_GREEN_50 = "E3FCEF"

COLOR_AMBER_700 = "B76E00"
COLOR_AMBER_100 = "FFF0B3"

COLOR_RED_700 = "DE350B"
COLOR_RED_100 = "FFEBE6"

COLOR_GRAY_900 = "172B4D"
COLOR_GRAY_600 = "5E6C84"
COLOR_GRAY_200 = "EBECF0"
COLOR_GRAY_100 = "F4F5F7"
COLOR_WHITE = "FFFFFF"

FONT_FAMILY = "Segoe UI"


def thin_border():
    b = Side(style="thin", color=COLOR_GRAY_200)
    return Border(left=b, right=b, top=b, bottom=b)


def generate_abnahme_report(target_weeks=None):
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Übersicht Abnahme"
    ws_summary.sheet_view.showGridLines = True

    # Считываем отчёты через единый контракт датасета
    dataset = load_accepted_reports_dataset(target_weeks=target_weeks)
    selected_site_ids = set(dataset.all_sites.keys())
    site_obj_map = {sid: AcObject.objects.get(pk=sid) for sid in selected_site_ids}

    # Имена листов
    site_sheet_names = assign_site_sheet_titles(
        ((s_id, site_obj_map[s_id].name) for s_id in selected_site_ids), reserved={"Übersicht Abnahme"}
    )

    # Валидация общего статуса приёмки объекта
    dv_status = DataValidation(
        type="list",
        formula1='"Freigegeben,Nacharbeit erforderlich,In Prüfung,Nicht abgenommen"',
        allow_blank=True,
    )
    ws_summary.add_data_validation(dv_status)

    dv_yes_no = DataValidation(
        type="list",
        formula1='"Ja,Nein"',
        allow_blank=True,
    )
    ws_summary.add_data_validation(dv_yes_no)

    # Шапка главного листа
    ws_summary.merge_cells("A1:E1")
    t1 = ws_summary["A1"]
    t1.value = "DLC Montage · Abnahme & Leistungsfreigabe (Bauleitung)"
    t1.font = Font(name=FONT_FAMILY, size=15, bold=True, color=COLOR_NAVY_900)

    ws_summary.merge_cells("A2:E2")
    t2 = ws_summary["A2"]
    t2.value = "Prüfprotokoll und Freigabeübersicht der fertiggestellten Montageleistungen"
    t2.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_700)

    # Справа инфо
    ws_summary["G1"] = "Auftragnehmer: DLC Montage (Mikhail)"
    ws_summary["G1"].font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
    ws_summary["G2"] = "Bauleitung / Abnahme: Emil (SPIE)"
    ws_summary["G2"].font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_GREEN_600)
    ws_summary["G3"] = "Projekt: AC-Montage Berlin"
    ws_summary["G3"].font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_600)

    # Карточка состояния
    ws_summary.merge_cells("A4:I4")
    card_h = ws_summary["A4"]
    card_h.value = "★ Prüfübersicht zur Freigabe"
    card_h.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
    card_h.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
    card_h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_summary.row_dimensions[4].height = 24

    completed_cnt = sum(1 for sid in selected_site_ids if dataset.all_sites[sid].is_completed)
    in_prog_cnt = len(selected_site_ids) - completed_cnt
    in_prog_names = ", ".join(
        sorted(dataset.all_sites[sid].site_name for sid in selected_site_ids if not dataset.all_sites[sid].is_completed)
    )
    in_prog_suffix = f" ({in_prog_names})" if in_prog_names else ""

    ws_summary.merge_cells("A5:I5")
    card_sub = ws_summary["A5"]
    card_sub.value = (
        f"Geprüfte Standorte im Zeitraum: {len(selected_site_ids)} gesamt  |  "
        f"Montage abgeschlossen: {completed_cnt}  |  "
        f"In Bearbeitung: {in_prog_cnt}{in_prog_suffix}  |  "
        f"Freigabe zur Abrechnung durch Bauleitung (Emil)"
    )
    card_sub.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
    card_sub.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)
    card_sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_summary.row_dimensions[5].height = 22

    for c in range(1, 10):
        ws_summary.cell(5, c).border = thin_border()

    summary_headers = [
        "Nr.",
        "Standort (Objekt)",
        "Adresse",
        "Anlage (Projektgruppe)",
        "Bau-Status (DLC)",
        "Abnahme-Status (Emil)",
        "Prüfnotiz / Bemerkung des Bauleiters",
        "Freigabe",
        "Prüfdatum",
    ]

    week_titles = {
        39: "KW 39 (21.09.2026 – 25.09.2026)",
        40: "KW 40 (28.09.2026 – 02.10.2026)",
    }

    current_row = 7
    status_cell_ranges = []
    yes_no_ranges = []

    for iso_year, iso_week in dataset.periods:
        sites_in_period = dataset.sites_by_period.get((iso_year, iso_week), {})
        if not sites_in_period:
            continue

        ws_summary.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        period_title = (
            week_titles.get(iso_week, f"KW {iso_week} / {iso_year}")
            if iso_year == 2026
            else f"KW {iso_week} / {iso_year}"
        )
        wk_cell = ws_summary.cell(current_row, 1, period_title)
        wk_cell.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        wk_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        wk_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_summary.row_dimensions[current_row].height = 24
        current_row += 1

        for c_idx, h_text in enumerate(summary_headers, 1):
            h_c = ws_summary.cell(current_row, c_idx, h_text)
            h_c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
            h_c.fill = PatternFill("solid", fgColor=COLOR_NAVY_700)
            h_c.alignment = Alignment(horizontal="center" if c_idx in (1, 5, 6, 8, 9) else "left", vertical="center")
            h_c.border = thin_border()
        ws_summary.row_dimensions[current_row].height = 24
        current_row += 1

        start_kw_row = current_row
        sorted_sites = sorted(sites_in_period.items(), key=lambda x: x[1].site_name)

        for idx, (s_id, s_summary) in enumerate(sorted_sites, 1):
            _site = site_obj_map[s_id]
            is_in_prog = not s_summary.is_completed
            status_text = s_summary.construction_status
            sheet_target = site_sheet_names[s_id]

            ws_summary.cell(current_row, 1, idx).alignment = Alignment(horizontal="center", vertical="top")

            # Гиперссылка на лист объекта
            obj_cell = ws_summary.cell(current_row, 2, s_summary.site_name)
            obj_cell.hyperlink = f"#'{sheet_target}'!A1"
            obj_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900, underline="single")
            obj_cell.alignment = Alignment(horizontal="left", vertical="top")

            ws_summary.cell(current_row, 3, s_summary.address or "—").alignment = Alignment(
                horizontal="left", vertical="top"
            )
            ws_summary.cell(current_row, 4, s_summary.group_name or "—").alignment = Alignment(
                horizontal="left", vertical="top"
            )

            # Строительный статус DLC
            st_cell = ws_summary.cell(current_row, 5, status_text)
            st_cell.alignment = Alignment(horizontal="center", vertical="top")
            if is_in_prog:
                st_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700)
                st_cell.fill = PatternFill("solid", fgColor=COLOR_AMBER_100)
            else:
                st_cell.font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600)
                st_cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50)

            # Колонка 6: Abnahme-Status (Emil)
            default_status = s_summary.curator_review.status
            ab_cell = ws_summary.cell(current_row, 6, default_status)
            ab_cell.alignment = Alignment(horizontal="center", vertical="top")
            ab_cell.font = Font(
                name=FONT_FAMILY,
                size=9,
                bold=True,
                color=COLOR_GREEN_600 if default_status == "Freigegeben" else COLOR_AMBER_700,
            )

            # Колонка 7: Prüfnotiz / Bemerkung des Bauleiters
            default_note = s_summary.curator_review.note
            note_cell = ws_summary.cell(current_row, 7, default_note)
            note_cell.alignment = Alignment(horizontal="left", vertical="top")
            note_cell.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

            # Колонка 8: Freigabe
            freigabe_val = s_summary.curator_review.freigabe
            fr_cell = ws_summary.cell(current_row, 8, freigabe_val)
            fr_cell.alignment = Alignment(horizontal="center", vertical="top")
            fr_cell.font = Font(
                name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600 if freigabe_val == "Ja" else COLOR_AMBER_700
            )

            # Колонка 9: Prüfdatum
            dt_cell = ws_summary.cell(current_row, 9, s_summary.curator_review.date)
            dt_cell.alignment = Alignment(horizontal="center", vertical="top")
            dt_cell.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_600)

            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if idx % 2 != 0 else COLOR_GRAY_100)
            for c in range(1, 10):
                cell_item = ws_summary.cell(current_row, c)
                if c not in (5, 6):
                    cell_item.fill = row_fill
                cell_item.border = thin_border()
                if c not in (2, 5, 6, 8):
                    cell_item.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

            ws_summary.row_dimensions[current_row].height = 24
            current_row += 1

        end_kw_row = current_row - 1
        st_range = f"F{start_kw_row}:F{end_kw_row}"
        dv_status.add(st_range)
        status_cell_ranges.append(st_range)

        yn_range = f"H{start_kw_row}:H{end_kw_row}"
        dv_yes_no.add(yn_range)
        yes_no_ranges.append(yn_range)

        current_row += 1  # отступ

    # Условное форматирование статуса приёмки
    rule_frei = CellIsRule(
        operator="equal",
        formula=['"Freigegeben"'],
        fill=PatternFill("solid", fgColor=COLOR_GREEN_50),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600),
    )
    rule_nach = CellIsRule(
        operator="equal",
        formula=['"Nacharbeit erforderlich"'],
        fill=PatternFill("solid", fgColor=COLOR_AMBER_100),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700),
    )
    rule_pruf = CellIsRule(
        operator="equal",
        formula=['"In Prüfung"'],
        fill=PatternFill("solid", fgColor=COLOR_NAVY_100),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_NAVY_700),
    )
    rule_abg = CellIsRule(
        operator="equal",
        formula=['"Nicht abgenommen"'],
        fill=PatternFill("solid", fgColor=COLOR_RED_100),
        font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_RED_700),
    )

    for rng in status_cell_ranges:
        ws_summary.conditional_formatting.add(rng, rule_frei)
        ws_summary.conditional_formatting.add(rng, rule_nach)
        ws_summary.conditional_formatting.add(rng, rule_pruf)
        ws_summary.conditional_formatting.add(rng, rule_abg)

    # Ширины колонок Übersicht
    widths = {
        "A": 6,
        "B": 24,
        "C": 30,
        "D": 32,
        "E": 18,
        "F": 24,
        "G": 44,
        "H": 12,
        "I": 14,
    }
    for col_l, w in widths.items():
        ws_summary.column_dimensions[col_l].width = w

    ws_summary.freeze_panes = "A7"

    # ==========================================================================
    # ЛИСТЫ ОБЪЕКТОВ (Standorte) ДЛЯ ЭМИЛЯ (7 КОЛОНОК)
    # ==========================================================
    for s_id in sorted(selected_site_ids, key=lambda sid: site_obj_map[sid].name):
        _site = site_obj_map[s_id]
        sheet_title = site_sheet_names[s_id]
        ws = wb.create_sheet(title=sheet_title)
        ws.sheet_view.showGridLines = True

        s_summary = dataset.all_sites[s_id]
        site_periods = sorted(list(set((d.iso_year, d.iso_week) for d in s_summary.day_facts)))
        kw_display = ", ".join(f"KW {k} ({y})" if y != 2026 else f"KW {k}" for (y, k) in site_periods)
        is_in_prog = not s_summary.is_completed

        # Валидация для строчной приёмки (✓ i.O. / ✗ Nicht i.O. / —)
        dv_item_status = DataValidation(
            type="list",
            formula1='"✓ i.O.,✗ Nicht i.O.,—"',
            allow_blank=True,
        )
        ws.add_data_validation(dv_item_status)

        # Шапка объекта
        ws.merge_cells("A1:E1")
        h1 = ws["A1"]
        h1.value = f"DLC Montage · Abnahmeprotokoll Standort: {s_summary.site_name}"
        h1.font = Font(name=FONT_FAMILY, size=14, bold=True, color=COLOR_NAVY_900)
        h1.alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("F1:G1")
        back_cell = ws["F1"]
        back_cell.value = "← Zurück zur Übersicht"
        back_cell.hyperlink = "#'Übersicht Abnahme'!A1"
        back_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_GREEN_600, underline="single")
        back_cell.alignment = Alignment(horizontal="right", vertical="center")

        ws.merge_cells("A2:D2")
        ws["A2"] = f"Standort: {s_summary.site_name}"
        ws["A2"].font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900)
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("E2:G2")
        ws["E2"] = f"Anlage (Gruppe): {s_summary.group_name or '—'}"
        ws["E2"].font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_NAVY_900)
        ws["E2"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A3:D3")
        ws["A3"] = f"Adresse: {s_summary.address or '—'}"
        ws["A3"].font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_900)
        ws["A3"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("E3:G3")
        ws["E3"] = f"Bau-Status (DLC): {'In Bearbeitung' if is_in_prog else 'Vollständig abgeschlossen'}"
        ws["E3"].font = Font(
            name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_AMBER_700 if is_in_prog else COLOR_GREEN_600
        )
        ws["E3"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A4:D4")
        ws["A4"] = f"Abrechnungszeitraum: {kw_display}"
        ws["A4"].font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_GRAY_900)
        ws["A4"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("E4:G4")
        ws["E4"] = "Prüfender Bauleiter: Emil (SPIE)"
        ws["E4"].font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
        ws["E4"].alignment = Alignment(horizontal="left", vertical="center")

        card_fill = PatternFill("solid", fgColor=COLOR_GRAY_100)
        for r_h in range(2, 5):
            for c_h in range(1, 8):
                ws.cell(r_h, c_h).fill = card_fill
            ws.row_dimensions[r_h].height = 22
        ws.row_dimensions[1].height = 28

        # Таблица выполненных работ (7 колонок)
        row_cur = 6
        ws.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=7)
        tbl_t = ws.cell(row_cur, 1, "1. ERBRACHTE MONTAGELEISTUNGEN (Zur Abnahme durch den Bauleiter)")
        tbl_t.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        tbl_t.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        tbl_t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row_cur].height = 24
        row_cur += 1

        l_cols = [
            "Pos.",
            "Leistung / Montageposition",
            "Einheit",
            "Menge",
            "Status Montage (DLC)",
            "Abnahme (Bauleiter Emil)",
            "Bemerkung / Mangel (Emil)",
        ]
        for c_idx, c_name in enumerate(l_cols, 1):
            c_cell = ws.cell(row_cur, c_idx, c_name)
            c_cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
            c_cell.fill = PatternFill("solid", fgColor=COLOR_NAVY_700)
            c_cell.alignment = Alignment(horizontal="center" if c_idx in (1, 3, 4, 5, 6) else "left", vertical="center")
            c_cell.border = thin_border()
        ws.row_dimensions[row_cur].height = 24
        row_cur += 1

        start_items_row = row_cur
        perf_rows = s_summary.position_rows

        for p_idx, prow in enumerate(perf_rows, 1):
            ws.cell(row_cur, 1, p_idx).alignment = Alignment(horizontal="center", vertical="top")
            ws.cell(row_cur, 2, prow.name).alignment = Alignment(horizontal="left", vertical="top")
            ws.cell(row_cur, 3, prow.unit or "Stk.").alignment = Alignment(horizontal="center", vertical="top")

            if prow.quantity is not None:
                qty_cell = ws.cell(row_cur, 4, float(prow.quantity))
                qty_cell.alignment = Alignment(horizontal="right", vertical="top")
                qty_cell.number_format = "0.##"
            else:
                qty_cell = ws.cell(row_cur, 4, "—")
                qty_cell.alignment = Alignment(horizontal="center", vertical="top")

            st_dlc = ws.cell(row_cur, 5, "In Ausführung" if is_in_prog else "✓ Erledigt")
            st_dlc.font = Font(
                name=FONT_FAMILY, size=9, bold=True, color=COLOR_AMBER_700 if is_in_prog else COLOR_GREEN_600
            )
            st_dlc.alignment = Alignment(horizontal="center", vertical="top")

            abn_cell = ws.cell(row_cur, 6, "—")
            abn_cell.alignment = Alignment(horizontal="center", vertical="top")
            abn_cell.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_600)

            bem_cell = ws.cell(row_cur, 7, "")
            bem_cell.alignment = Alignment(horizontal="left", vertical="top")
            bem_cell.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)

            row_fill = PatternFill("solid", fgColor=COLOR_WHITE if row_cur % 2 == 0 else COLOR_GRAY_100)
            for c in range(1, 8):
                cc = ws.cell(row_cur, c)
                cc.fill = row_fill
                cc.border = thin_border()
                if c not in (5, 6):
                    cc.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)
            ws.row_dimensions[row_cur].height = 20
            row_cur += 1

        end_items_row = row_cur - 1
        dv_item_status.add(f"F{start_items_row}:F{end_items_row}")

        # Условное форматирование строк приёмки позиции
        rule_io = CellIsRule(
            operator="equal",
            formula=['"✓ i.O."'],
            fill=PatternFill("solid", fgColor=COLOR_GREEN_50),
            font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GREEN_600),
        )
        rule_nio = CellIsRule(
            operator="equal",
            formula=['"✗ Nicht i.O."'],
            fill=PatternFill("solid", fgColor=COLOR_RED_100),
            font=Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_RED_700),
        )
        ws.conditional_formatting.add(f"F{start_items_row}:F{end_items_row}", rule_io)
        ws.conditional_formatting.add(f"F{start_items_row}:F{end_items_row}", rule_nio)

        row_cur += 1

        # ======================================================================
        # БЛОК 2: РЕЗОЛЮЦИЯ И ЗАМЕЧАНИЯ КУРАТОРА (ABNAHMEPROTOKOLL)
        # ======================================================================
        ws.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=7)
        ab_t = ws.cell(row_cur, 1, "2. ABNAHMEPROTOKOLL & MÄNGELRÜGE (Ausfüllen durch Bauleiter Emil)")
        ab_t.font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_WHITE)
        ab_t.fill = PatternFill("solid", fgColor=COLOR_NAVY_900)
        ab_t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row_cur].height = 24
        row_cur += 1

        # Поле общего статуса приёмки
        ws.cell(row_cur, 1, "Gesamturteil Abnahme:").font = Font(
            name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900
        )
        ws.cell(row_cur, 1).alignment = Alignment(horizontal="left", vertical="center")

        cur_status = s_summary.curator_review.status
        is_freigegeben = cur_status == "Freigegeben"
        st_val_cell = ws.cell(row_cur, 2, cur_status)
        st_val_cell.font = Font(
            name=FONT_FAMILY, size=10, bold=True, color=COLOR_GREEN_600 if is_freigegeben else COLOR_AMBER_700
        )
        st_val_cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50 if is_freigegeben else COLOR_AMBER_100)
        st_val_cell.alignment = Alignment(horizontal="center", vertical="center")
        st_val_cell.border = thin_border()

        dv_status.add(f"B{row_cur}")

        ws.cell(row_cur, 4, "Freigabe zur Auszahlung:").font = Font(
            name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY_900
        )
        ws.cell(row_cur, 4).alignment = Alignment(horizontal="right", vertical="center")

        cur_freigabe = s_summary.curator_review.freigabe
        is_freigabe_ja = cur_freigabe == "Ja"
        fr_val_cell = ws.cell(row_cur, 5, cur_freigabe)
        fr_val_cell.font = Font(
            name=FONT_FAMILY, size=10, bold=True, color=COLOR_GREEN_600 if is_freigabe_ja else COLOR_AMBER_700
        )
        fr_val_cell.fill = PatternFill("solid", fgColor=COLOR_GREEN_50 if is_freigabe_ja else COLOR_AMBER_100)
        fr_val_cell.alignment = Alignment(horizontal="center", vertical="center")
        fr_val_cell.border = thin_border()

        dv_yes_no.add(f"E{row_cur}")

        ws.row_dimensions[row_cur].height = 24
        row_cur += 2

        # Чек-лист проверки
        ws.cell(row_cur, 1, "Prüfkriterien:").font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
        ws.cell(row_cur, 2, "☐ Fotodokumentation vollständig eingereicht").font = Font(
            name=FONT_FAMILY, size=9, color=COLOR_GRAY_900
        )
        ws.cell(row_cur, 5, "☐ Messprotokoll / Schutzleiterprüfung erfolgt").font = Font(
            name=FONT_FAMILY, size=9, color=COLOR_GRAY_900
        )
        ws.row_dimensions[row_cur].height = 20
        row_cur += 1

        ws.cell(row_cur, 2, "☐ Kabelwege & Trassenführung nach Vorgabe").font = Font(
            name=FONT_FAMILY, size=9, color=COLOR_GRAY_900
        )
        ws.cell(row_cur, 5, "☐ Baustelle besenrein & Abfälle geräumt").font = Font(
            name=FONT_FAMILY, size=9, color=COLOR_GRAY_900
        )
        ws.row_dimensions[row_cur].height = 20
        row_cur += 2

        # Поле для замечаний и доделок (Korrekturbedarf)
        ws.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=7)
        ws.cell(row_cur, 1, "Mängelliste / Korrekturbedarf / Anmerkungen des Bauleiters:").font = Font(
            name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900
        )
        ws.row_dimensions[row_cur].height = 20
        row_cur += 1

        # 4 строки для ввода замечаний Эмилем
        for m_i in range(1, 5):
            ws.cell(row_cur, 1, f"{m_i}.").font = Font(name=FONT_FAMILY, size=9, bold=True, color=COLOR_GRAY_600)
            ws.cell(row_cur, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row_cur, 1).border = thin_border()

            ws.merge_cells(start_row=row_cur, start_column=2, end_row=row_cur, end_column=7)
            default_m = s_summary.curator_review.note if m_i == 1 else ""
            m_txt = ws.cell(row_cur, 2, default_m)
            m_txt.font = Font(name=FONT_FAMILY, size=9, color=COLOR_GRAY_900)
            m_txt.alignment = Alignment(horizontal="left", vertical="center")
            for col_c in range(2, 8):
                ws.cell(row_cur, col_c).border = thin_border()
            ws.row_dimensions[row_cur].height = 22
            row_cur += 1

        row_cur += 1

        # Подпись куратора
        ws.merge_cells(start_row=row_cur, start_column=1, end_row=row_cur, end_column=3)
        sig1 = ws.cell(row_cur, 1, "Geprüft durch Bauleiter: Emil (SPIE)")
        sig1.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)

        ws.merge_cells(start_row=row_cur, start_column=5, end_row=row_cur, end_column=7)
        cur_date = s_summary.curator_review.date
        sig2 = ws.cell(row_cur, 5, f"Datum der Abnahme: {cur_date}" if cur_date else "Datum der Abnahme: ")
        sig2.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY_900)
        sig2.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[row_cur].height = 24

        # Ширины колонок
        obj_widths = {
            "A": 6,
            "B": 46,
            "C": 12,
            "D": 10,
            "E": 22,
            "F": 24,
            "G": 38,
        }
        for cl, w in obj_widths.items():
            ws.column_dimensions[cl].width = w

        ws.freeze_panes = "A6"

    # Сохраняем файл
    os.makedirs(r"d:\Workspace\clients\dlc_montage\media\exports", exist_ok=True)
    os.makedirs(r"d:\Workspace\clients\dlc_montage\outputs", exist_ok=True)

    out1 = r"d:\Workspace\clients\dlc_montage\media\exports\DLC_Montage_AC_Abnahme_KW39-KW40_Emil.xlsx"
    out2 = r"d:\Workspace\clients\dlc_montage\outputs\DLC_Montage_AC_Abnahme_KW39-KW40_Emil.xlsx"

    wb.save(out1)
    wb.save(out2)
    print(f"Abnahme-Report saved to: {out1}")
    print(f"Backup saved to: {out2}")
    return out1


if __name__ == "__main__":
    generate_abnahme_report()
